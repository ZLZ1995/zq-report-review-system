"""Read-only .xls intake: tracked temp conversion; originals are never written.

The generation pipeline reads .xlsx via openpyxl, so legacy .xls sources are
converted inside the run work directory. The original file id and SHA256 stay
authoritative; every conversion is recorded with its tool and hashes.
"""
from pathlib import Path

import xlrd

from .skills import digest

TOOL = f'xlrd {xlrd.__version__} + openpyxl'


def _cell_value(book, sheet, row, col):
    cell = sheet.cell(row, col)
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
        except (ValueError, OverflowError):
            return None
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    return cell.value


def visible_sheets(book):
    return [sheet for sheet in book.sheets() if sheet.visibility == 0]


def convert_xls_to_xlsx(source, target):
    """Convert a .xls source into a run-local .xlsx copy; return (path, report entry)."""
    from openpyxl import Workbook
    source = Path(source)
    try:
        book = xlrd.open_workbook(str(source))
    except Exception as exc:  # xlrd raises several exception types for bad OLE2
        raise ValueError(f'无法读取 .xls 工作簿（xlrd：{exc}）') from exc
    out = Workbook()
    out.remove(out.active)
    skipped = []
    for sheet in book.sheets():
        if sheet.visibility != 0:
            skipped.append(sheet.name)
            continue
        target_sheet = out.create_sheet(sheet.name[:31] or 'Sheet')
        for row in range(sheet.nrows):
            for col in range(sheet.ncols):
                value = _cell_value(book, sheet, row, col)
                if value is not None:
                    target_sheet.cell(row=row + 1, column=col + 1, value=value)
    if not out.sheetnames:
        raise ValueError('.xls 工作簿没有可见工作表，不能作为填报来源')
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    out.save(target)
    entry = {'source_name': source.name, 'source_sha256': digest(source),
             'output_name': target.name, 'output_sha256': digest(target),
             'tool': TOOL, 'skipped_hidden_sheets': skipped,
             'visible_sheets': list(out.sheetnames)}
    return target, entry


def xls_header_values(path, scan_rows):
    """Visible-sheet header text for metadata extraction; None when unreadable."""
    try:
        book = xlrd.open_workbook(str(path))
    except Exception:  # unreadable workbook means unidentifiable
        return None
    sheets = visible_sheets(book)
    if not sheets:
        return ()
    sheet = next((item for item in sheets if item.name == '资产负债表'), sheets[0])
    values = []
    for row in range(min(sheet.nrows, scan_rows)):
        for col in range(sheet.ncols):
            cell = sheet.cell(row, col)
            if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                continue
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    moment = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                except (ValueError, OverflowError):
                    continue
                values.append(f'{moment.year}年{moment.month}月{moment.day}日')
                values.append(moment.date().isoformat())
                continue
            text = str(cell.value).strip()
            if text:
                values.append(text)
    return values
