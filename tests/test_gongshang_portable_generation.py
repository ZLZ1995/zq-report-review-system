import importlib.util
from pathlib import Path

from docx import Document
from openpyxl import Workbook


def module():
    path = Path(__file__).resolve().parents[1] / '.codex/skills/gongshang-change-history-docx/scripts/build_gongshang_docx.py'
    spec = importlib.util.spec_from_file_location('portable_gongshang', path)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


def test_all_dates_retain_capital_type_and_explicit_representative():
    m = module()
    rows = [{'item': name, 'before': '旧值', 'after': '新值'} for name in
            ['注册资本变更', '企业类型变更', '法定代表人变更']]
    assert len(m.normal_rows_for_human_logic('2025-01-02', rows, False)) == 3
    assert not m.event_has_legal_representative_change([
        {'item': '主要人员变更', 'before': '张三（法定代表人）', 'after': '李四'}])


def test_name_only_transfer_never_assigns_unproven_amount():
    m = module()
    row = {'before': '甲;乙', 'after': '丙;丁'}
    resolved = m.resolve_name_only_equity(row, {'甲': '40', '乙': '60'})
    assert resolved is None


def test_source_reader_preserves_inputs_and_groups_dates(tmp_path):
    m = module()
    source, template = tmp_path / 'source.xlsx', tmp_path / 'template.docx'
    wb = Workbook()
    wb.active.title = '变更信息'
    wb.active.append(['变更日期', '变更事项', '变更前', '变更后'])
    wb.active.append(['2025-01-02', '法定代表人变更', '张三', '李四'])
    wb.save(source)
    doc = Document()
    doc.add_paragraph('旧项目残留')
    doc.save(template)
    before = source.read_bytes(), template.read_bytes()
    events = m.load_events(source)
    assert events[0][0] == '2025-01-02'
    assert events[0][1][0]['after'] == '李四'
    assert before == (source.read_bytes(), template.read_bytes())


def test_source_reader_does_not_silently_drop_incomplete_event(tmp_path):
    import pytest
    m = module()
    source = tmp_path / 'source.xlsx'
    wb = Workbook()
    wb.active.title = '变更信息'
    wb.active.append(['变更日期', '变更事项', '变更前', '变更后'])
    wb.active.append([None, '法定代表人变更', '张三', '李四'])
    wb.save(source)
    with pytest.raises(ValueError, match='日期|事项'):
        m.load_events(source)
