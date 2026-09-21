"""Read-only extraction of reviewable text from local project files."""

from __future__ import annotations

import os
import re
import shutil
import zipfile
from pathlib import Path
from typing import Any, Protocol

from pydantic import BaseModel, ConfigDict, Field

from ..domain.enums import FileRole
from ..domain.models import IssueLocation, SourceFile
from .workbook_dependency_service import formula_references


class ExtractionModel(BaseModel):
    model_config = ConfigDict(extra="forbid")


class DocumentChunk(ExtractionModel):
    chunk_id: str
    source_file_id: str
    source_file_name: str
    role: FileRole
    text: str
    location: IssueLocation
    reference_only: bool = False
    summary_impact: bool = False
    story_type: str = "body"
    style_id: str | None = None
    style_name: str | None = None
    outline_level: int | None = None
    is_toc_entry: bool = False
    sequence: int | None = None
    cached_values: dict[str, Any] = Field(default_factory=dict)


class ExtractedDocument(ExtractionModel):
    source_file: SourceFile
    chunks: list[DocumentChunk] = Field(default_factory=list)
    extraction_mode: str
    warnings: list[str] = Field(default_factory=list)


class OcrEngine(Protocol):
    def extract_pages(self, path: Path) -> list[str]:
        ...


class TesseractOcrEngine:
    """Optional local OCR adapter; dependencies are loaded only when required."""

    def __init__(self, language: str = "chi_sim+eng") -> None:
        self.language = language

    def extract_pages(self, path: Path) -> list[str]:
        try:
            import pypdfium2 as pdfium
            import pytesseract
        except ImportError as exc:  # pragma: no cover - environment dependent
            raise RuntimeError(
                "Tesseract OCR dependencies are unavailable; install pytesseract and pypdfium2"
            ) from exc
        if shutil.which("tesseract") is None:
            program_files = Path(os.environ.get("PROGRAMFILES", r"C:\Program Files"))
            executable = program_files / "Tesseract-OCR" / "tesseract.exe"
            if executable.exists():
                pytesseract.pytesseract.tesseract_cmd = str(executable)
        tessdata_config = ""
        local_app_data = os.environ.get("LOCALAPPDATA")
        if local_app_data:
            tessdata_dir = Path(local_app_data) / "Tesseract-OCR" / "tessdata"
            required_languages = self.language.split("+")
            if all(
                (tessdata_dir / f"{language}.traineddata").exists()
                for language in required_languages
            ):
                tessdata_config = f"--tessdata-dir {tessdata_dir}"
        document = pdfium.PdfDocument(str(path))
        pages: list[str] = []
        try:
            for page in document:
                image = page.render(scale=2.0).to_pil()
                pages.append(
                    pytesseract.image_to_string(
                        image,
                        lang=self.language,
                        config=tessdata_config,
                    )
                )
        finally:
            document.close()
        return pages


class DocumentExtractionService:
    def __init__(self, ocr_engine: OcrEngine | None = None) -> None:
        self.ocr_engine = ocr_engine or TesseractOcrEngine()

    def extract(self, source_file: SourceFile) -> ExtractedDocument:
        path = Path(source_file.original_path)
        suffix = source_file.extension.lower()
        if suffix == ".docx":
            return self._extract_docx(source_file, path)
        if suffix in {".xlsx", ".xlsm"}:
            return self._extract_workbook(source_file, path)
        if suffix == ".pdf":
            return self._extract_pdf(source_file, path)
        raise ValueError(f"unsupported extraction type: {suffix}")

    @staticmethod
    def _extract_docx(source_file: SourceFile, path: Path) -> ExtractedDocument:
        from docx import Document

        document = Document(str(path))
        chunks: list[DocumentChunk] = []
        paragraph_index = 0
        table_index = 0
        for sequence, (block_type, block) in enumerate(
            _iter_document_blocks(document),
            start=1,
        ):
            if block_type == "paragraph":
                paragraph_index += 1
                text = block.text.strip()
                if text:
                    chunks.append(
                        _chunk(
                            source_file,
                            f"paragraph-{paragraph_index}",
                            text,
                            IssueLocation(paragraph=paragraph_index),
                            style_id=block.style.style_id,
                            style_name=block.style.name,
                            outline_level=_paragraph_outline_level(block),
                            is_toc_entry=_is_toc_style(block.style.name),
                            sequence=sequence,
                        )
                    )
                continue
            table_index += 1
            for row_index, row in enumerate(block.rows, start=1):
                values = [cell.text.strip() for cell in row.cells]
                text = " | ".join(value for value in values if value)
                if text:
                    chunks.append(
                        _chunk(
                            source_file,
                            f"table-{table_index}-row-{row_index}",
                            text,
                            IssueLocation(
                                table=str(table_index),
                                cell=f"row:{row_index}",
                            ),
                            sequence=sequence,
                        )
                    )
        seen_parts: set[str] = set()
        for section_index, section in enumerate(document.sections, start=1):
            for story_type, story in (
                ("header", section.header),
                ("footer", section.footer),
            ):
                part_name = str(story.part.partname)
                if part_name in seen_parts:
                    continue
                seen_parts.add(part_name)
                for paragraph_index, paragraph in enumerate(
                    story.paragraphs,
                    start=1,
                ):
                    text = paragraph.text.strip()
                    if not text:
                        continue
                    chunks.append(
                        _chunk(
                            source_file,
                            f"{story_type}-{section_index}-{paragraph_index}",
                            text,
                            IssueLocation(paragraph=paragraph_index),
                            story_type=story_type,
                            style_id=paragraph.style.style_id,
                            style_name=paragraph.style.name,
                        )
                    )
        chunks.extend(_extract_textbox_chunks(source_file, path))
        return ExtractedDocument(
            source_file=source_file,
            chunks=chunks,
            extraction_mode="python-docx-readonly",
        )

    @staticmethod
    def _extract_workbook(source_file: SourceFile, path: Path) -> ExtractedDocument:
        from openpyxl import load_workbook

        workbook = load_workbook(
            path,
            read_only=True,
            data_only=False,
            keep_vba=source_file.extension.lower() == ".xlsm",
        )
        cached_workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_vba=source_file.extension.lower() == ".xlsm",
        )
        chunks: list[DocumentChunk] = []
        try:
            hidden_dependency_cells = _hidden_dependency_cells(workbook)
            for worksheet in workbook.worksheets:
                if worksheet.sheet_state != "visible":
                    continue
                cached_worksheet = cached_workbook[worksheet.title]
                for row_index, row in enumerate(worksheet.iter_rows(), start=1):
                    if any(
                        (
                            worksheet.title,
                            getattr(cell, "coordinate", None),
                        )
                        in hidden_dependency_cells
                        for cell in row
                    ):
                        continue
                    selected_cells = [
                        cell
                        for cell in row
                        if cell.value not in (None, "")
                    ]
                    values = [
                        f"{cell.coordinate}={cell.value}"
                        for cell in selected_cells
                    ]
                    cached_values = {
                        cell.coordinate: cached_worksheet[cell.coordinate].value
                        for cell in selected_cells
                        if cached_worksheet[cell.coordinate].value is not None
                    }
                    if values:
                        chunks.append(
                            _chunk(
                                source_file,
                                f"sheet-{worksheet.title}-row-{row_index}",
                                " | ".join(values),
                                IssueLocation(
                                    table=worksheet.title,
                                    cell=f"row:{row_index}",
                                ),
                                cached_values=cached_values,
                            )
                        )
        finally:
            workbook.close()
            cached_workbook.close()
        return ExtractedDocument(
            source_file=source_file,
            chunks=chunks,
            extraction_mode="openpyxl-readonly",
        )

    def _extract_pdf(self, source_file: SourceFile, path: Path) -> ExtractedDocument:
        import pdfplumber

        page_texts: list[str] = []
        with pdfplumber.open(path) as document:
            for page in document.pages:
                page_texts.append((page.extract_text() or "").strip())
        extraction_mode = "pdfplumber"
        warnings: list[str] = []
        if sum(len(text) for text in page_texts) < 200:
            page_texts = self.ocr_engine.extract_pages(path)
            extraction_mode = "tesseract"
            warnings.append("native PDF text was insufficient; local OCR was used")
        chunks = [
            _chunk(
                source_file,
                f"page-{index}",
                text.strip(),
                IssueLocation(page=index),
                reference_only=True,
            )
            for index, text in enumerate(page_texts, start=1)
            if text.strip()
        ]
        return ExtractedDocument(
            source_file=source_file,
            chunks=chunks,
            extraction_mode=extraction_mode,
            warnings=warnings,
        )


_SUMMARY_SHEET_MARKERS = ("汇总", "结果", "资产负债表")
_FORMULA_PREFIX = r"(?:^|[=+\-*/^(,<>&;])"


def _is_summary_worksheet(title: str) -> bool:
    return any(marker in title for marker in _SUMMARY_SHEET_MARKERS)


def _hidden_dependency_cells(workbook: object) -> set[tuple[str, str]]:
    sheet_names = {worksheet.title for worksheet in workbook.worksheets}
    hidden_sheets = {
        worksheet.title
        for worksheet in workbook.worksheets
        if worksheet.sheet_state != "visible"
    }
    hidden_defined_names = _defined_names_backed_by_hidden_sheets(
        workbook,
        hidden_sheets,
    )
    dependencies: dict[tuple[str, str], set[tuple[str, str]]] = {}
    contaminated: set[tuple[str, str]] = set()
    for worksheet in workbook.worksheets:
        if worksheet.sheet_state != "visible":
            continue
        for row in worksheet.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str) or not value.startswith("="):
                    continue
                key = (worksheet.title, cell.coordinate)
                references = formula_references(
                    value,
                    worksheet.title,
                    sheet_names,
                )
                dependencies[key] = references
                referenced_sheets = _formula_sheet_references(
                    value,
                    list(sheet_names),
                )
                if (
                    hidden_sheets.intersection(referenced_sheets)
                    or _formula_defined_name_references(
                        value,
                        hidden_defined_names,
                    )
                ):
                    contaminated.add(key)

    changed = True
    while changed:
        changed = False
        for key, references in dependencies.items():
            if key in contaminated or not references.intersection(contaminated):
                continue
            contaminated.add(key)
            changed = True
    return contaminated


def _defined_names_backed_by_hidden_sheets(
    workbook: object,
    hidden_sheets: set[str],
) -> set[str]:
    names: set[str] = set()
    for defined_name in workbook.defined_names.values():
        try:
            destinations = list(defined_name.destinations)
        except (AttributeError, TypeError, ValueError):
            continue
        if any(sheet_name in hidden_sheets for sheet_name, _ in destinations):
            names.add(defined_name.name)
    return names


def _formula_defined_name_references(
    formula: str,
    defined_names: set[str],
) -> bool:
    if not defined_names:
        return False
    from openpyxl.formula import Tokenizer

    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return False
    return any(
        token.type == "OPERAND"
        and token.subtype == "RANGE"
        and token.value in defined_names
        for token in tokens
    )


def _worksheet_dependencies(
    worksheet: object,
    sheet_names: list[str],
) -> set[str]:
    dependencies: set[str] = set()
    for row in worksheet.iter_rows():
        for cell in row:
            value = cell.value
            if not isinstance(value, str) or not value.startswith("="):
                continue
            dependencies.update(_formula_sheet_references(value, sheet_names))
    dependencies.discard(worksheet.title)
    return dependencies


def _formula_sheet_references(formula: str, sheet_names: list[str]) -> set[str]:
    references: set[str] = set()
    for sheet_name in sheet_names:
        quoted_name = sheet_name.replace("'", "''")
        if f"'{quoted_name}'!" in formula:
            references.add(sheet_name)
            continue
        pattern = _FORMULA_PREFIX + re.escape(sheet_name) + r"!"
        if re.search(pattern, formula):
            references.add(sheet_name)
    return references


def _transitive_dependencies(
    roots: set[str],
    dependencies: dict[str, set[str]],
) -> set[str]:
    resolved: set[str] = set()
    pending = list(roots)
    while pending:
        sheet_name = pending.pop()
        for dependency in dependencies.get(sheet_name, set()):
            if dependency in resolved:
                continue
            resolved.add(dependency)
            pending.append(dependency)
    return resolved


def _chunk(
    source_file: SourceFile,
    suffix: str,
    text: str,
    location: IssueLocation,
    *,
    reference_only: bool = False,
    summary_impact: bool = False,
    story_type: str = "body",
    style_id: str | None = None,
    style_name: str | None = None,
    outline_level: int | None = None,
    is_toc_entry: bool = False,
    sequence: int | None = None,
    cached_values: dict[str, Any] | None = None,
) -> DocumentChunk:
    return DocumentChunk(
        chunk_id=f"{source_file.file_id}:{suffix}",
        source_file_id=source_file.file_id,
        source_file_name=source_file.original_name,
        role=source_file.role,
        text=text,
        location=location,
        reference_only=reference_only,
        summary_impact=summary_impact,
        story_type=story_type,
        style_id=style_id,
        style_name=style_name,
        outline_level=outline_level,
        is_toc_entry=is_toc_entry,
        sequence=sequence,
        cached_values=cached_values or {},
    )


def _paragraph_outline_level(paragraph: object) -> int | None:
    nodes = paragraph._p.xpath("./w:pPr/w:outlineLvl")
    if not nodes:
        return None
    value = nodes[0].get(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"
    )
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _iter_document_blocks(document: object):
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield "paragraph", Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield "table", Table(child, document)


def _is_toc_style(style_name: str | None) -> bool:
    compact = (style_name or "").replace(" ", "").lower()
    return compact.startswith("toc") or compact.startswith("目录")


def _extract_textbox_chunks(
    source_file: SourceFile,
    path: Path,
) -> list[DocumentChunk]:
    from lxml import etree

    namespace = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    }
    chunks: list[DocumentChunk] = []
    seen: set[tuple[str, str]] = set()
    with zipfile.ZipFile(path) as archive:
        names = [
            name
            for name in archive.namelist()
            if name == "word/document.xml"
            or re.fullmatch(r"word/(?:header|footer)\d+\.xml", name)
        ]
        for name in names:
            root = etree.fromstring(archive.read(name))
            story_type = (
                "header_textbox"
                if "/header" in name
                else "footer_textbox"
                if "/footer" in name
                else "textbox"
            )
            for index, textbox in enumerate(
                root.xpath(".//w:txbxContent", namespaces=namespace),
                start=1,
            ):
                text = "".join(
                    textbox.xpath(".//w:t/text()", namespaces=namespace)
                ).strip()
                identity = (name, text)
                if text and identity not in seen:
                    seen.add(identity)
                    chunks.append(
                        _chunk(
                            source_file,
                            f"{story_type}-{index}",
                            text,
                            IssueLocation(),
                            story_type=story_type,
                        )
                    )
    return chunks
