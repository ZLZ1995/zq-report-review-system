from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from summary_sheet_policy import SUMMARY_SHEET_NAMES, is_linked_summary_sheet


CHAIN_SHEETS = {
    "封面": "metadata",
    "封面页": "metadata",
    "索引目录": "metadata",
    "填表说明": "metadata",
    "资产负债表": "balance_sheet",
    "汇总表": "summary",
    "分类汇总": "summary",
    "流动汇总": "summary",
    "非流动资产汇总": "summary",
    "流动负债汇总": "summary",
    "银行存款": "detail",
    "现金": "detail",
    "其他货币资金": "detail",
    "应收票据": "detail",
    "应收账款": "detail",
    "预付账款": "detail",
    "应收利息": "detail",
    "应收股利（利润）": "detail",
    "其他应收款": "detail",
    "存货汇总": "detail",
    "一年到期非流动资产": "detail",
    "其他流动资产": "detail",
    "可供出售金融资产汇总": "detail",
    "持有到期投资": "detail",
    "长期应收": "detail",
    "股权投资": "detail",
    "4-5-1投资性房地产": "detail",
    "4-5-2投资性房地产": "detail",
    "4-5-3投资性地产": "detail",
    "4-5-4投资性地产": "detail",
    "固定资产汇总": "detail",
    "房屋建筑物": "detail",
    "构筑物": "detail",
    "井巷": "detail",
    "管道沟槽": "detail",
    "机器设备": "detail",
    "车辆": "detail",
    "电子设备": "detail",
    "土地": "detail",
    "在建工程汇总": "detail",
    "在建（土建）": "detail",
    "在建（设备）": "detail",
    "工程物资": "detail",
    "固定资产清理": "detail",
    "生产性生物资产": "detail",
    "油气资产": "detail",
    "使用权资产": "detail",
    "无形资产汇总": "detail",
    "无形-土地": "detail",
    "无形-矿业权": "detail",
    "无形-其他": "detail",
    "开发支出": "detail",
    "商誉": "detail",
    "长期待摊费用": "detail",
    "递延所得税资产": "detail",
    "其他非流动资产": "detail",
    "短期借款": "detail",
    "交易性金融负债": "detail",
    "应付票据": "detail",
    "应付账款": "detail",
    "预收账款": "detail",
    "职工薪酬": "detail",
    "应交税费": "detail",
    "应付利息": "detail",
    "应付股利（利润）": "detail",
    "其他应付款": "detail",
    "一年到期非流动负债": "detail",
    "其他流动负债": "detail",
    "非流动负债汇总 ": "summary",
    "长期借款": "detail",
    "应付债券": "detail",
    "长期应付款": "detail",
    "专项应付款": "detail",
    "预计负债": "detail",
    "租赁负债": "detail",
    "递延所得税负债": "detail",
    "其他非流动负债": "detail",
}

CHAIN_SHEETS.update({name: 'summary' for name in SUMMARY_SHEET_NAMES})


def chain_role_for_sheet(name: str) -> str | None:
    return 'summary' if is_linked_summary_sheet(name) else CHAIN_SHEETS.get(name)


DETAIL_COLUMN_RULES = {
    "银行存款": {"A", "B", "C", "D", "E", "F", "G", "H", "I"},
    "现金": {"A", "B", "C", "D", "E", "F", "G", "H"},
    "其他货币资金": {"A", "B", "C", "D", "E", "F", "G", "H"},
    "应收票据": {"A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O"},
    "应收账款": {"A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P"},
    "预付账款": {"A", "B", "C", "D", "E", "F", "G", "H"},
    "应收利息": {"A", "B", "C", "D", "E", "F", "G", "H", "I"},
    "应收股利（利润）": {"A", "B", "C", "D", "E", "F", "G"},
    "其他应收款": {"A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P"},
    "应付账款": {"A", "B", "C", "D", "E", "F", "G", "I"},
    "预收账款": {"A", "B", "C", "D", "E", "F", "G", "I"},
    "其他应付款": {"A", "B", "C", "D", "E", "F", "G", "I"},
    "应交税费": {"A", "B", "C", "D", "E", "F", "G"},
    "职工薪酬": {"A", "B", "C", "D", "E", "F"},
    "短期借款": {"A", "B", "C", "D", "E", "F", "G"},
    "交易性金融负债": {"A", "B", "C", "D", "E", "F", "G"},
    "应付票据": {"A", "B", "C", "D", "E", "F", "G"},
    "应付利息": {"A", "B", "C", "D", "E", "F", "G", "H", "I"},
    "应付股利（利润）": {"A", "B", "C", "D", "E", "F", "G"},
    "一年到期非流动负债": {"A", "B", "C", "D", "E", "F", "G"},
    "其他流动负债": {"A", "B", "C", "D", "E", "F", "G"},
}


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def coord_to_colrow(cell_ref: str) -> tuple[int, int]:
    row, col = coordinate_to_tuple(cell_ref)
    return col, row


def is_inside_zone(cell_ref: str, zone: dict[str, int]) -> bool:
    col, row = coord_to_colrow(cell_ref)
    return zone["start_row"] <= row <= zone["end_row"] and zone["start_col"] <= col <= zone["end_col"]


def filter_body_candidates(meta: dict[str, Any]) -> tuple[list[str], list[str], list[str]]:
    candidates = meta.get("writable_input_candidates", [])
    body = meta.get("writable_body", {})
    summary = meta.get("summary_zone", {})
    formula_cells = set(meta.get("formula_cells", []))

    detail_body_cells = []
    forbidden_non_formula_cells = []
    confirmed_input_cells = []

    for cell_ref in candidates:
        if cell_ref in formula_cells:
            continue
        if body and is_inside_zone(cell_ref, body):
            detail_body_cells.append(cell_ref)
        elif summary and is_inside_zone(cell_ref, summary):
            forbidden_non_formula_cells.append(cell_ref)
        else:
            forbidden_non_formula_cells.append(cell_ref)

    sheet_name = meta["sheet"]
    role = chain_role_for_sheet(sheet_name)
    if role == "detail":
        allowed_cols = DETAIL_COLUMN_RULES.get(sheet_name)
        if allowed_cols:
            confirmed_input_cells = [
                cell for cell in detail_body_cells
                if get_column_letter(coord_to_colrow(cell)[0]) in allowed_cols
            ]
            if sheet_name in {"应收账款", "其他应收款"}:
                confirmed_input_cells.extend(
                    [
                        cell
                        for cell in detail_body_cells
                        if get_column_letter(coord_to_colrow(cell)[0]) == "P"
                    ]
                )
            forbidden_non_formula_cells.extend([cell for cell in detail_body_cells if cell not in confirmed_input_cells])
        else:
            confirmed_input_cells = list(detail_body_cells)
    elif role == "balance_sheet":
        confirmed_input_cells = []
        asset_rows = [7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 22, 23, 24, 25, 26, 28, 31, 33, 34, 35, 36, 37, 38]
        liability_rows = [7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 21, 22, 23, 24, 25, 26, 27, 28, 29, 31, 32, 34, 35, 36, 38]
        for row in asset_rows:
            for col in ("C", "D"):
                confirmed_input_cells.append(f"{col}{row}")
        for row in liability_rows:
            for col in ("H", "I"):
                confirmed_input_cells.append(f"{col}{row}")
        confirmed_input_cells = [
            cell
            for cell in confirmed_input_cells
            if cell not in formula_cells
        ]
    elif role == "metadata":
        confirmed_input_cells = []
        detail_body_cells = []
        footer_start = meta.get("footer_start_row", meta["max_row"] + 1)
        for cell_ref in candidates:
            col, row = coord_to_colrow(cell_ref)
            if meta["header_row"] < row < footer_start:
                confirmed_input_cells.append(cell_ref)
        forbidden_non_formula_cells.extend([cell for cell in candidates if cell not in confirmed_input_cells])
    elif role == "summary":
        confirmed_input_cells = []
        detail_body_cells = []
        forbidden_non_formula_cells.extend(candidates)
    else:
        confirmed_input_cells = []
        forbidden_non_formula_cells.extend(detail_body_cells)
        detail_body_cells = []

    return sorted(set(confirmed_input_cells)), sorted(set(detail_body_cells)), sorted(set(forbidden_non_formula_cells))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--layout-map", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    layout_map = load_json(Path(args.layout_map))
    selected_sheets = {}
    for sheet_name, meta in layout_map["sheets"].items():
        chain_role = chain_role_for_sheet(sheet_name)
        if not chain_role:
            continue
        confirmed, detail_body, forbidden = filter_body_candidates(meta)
        selected_sheets[sheet_name] = {
            "chain_role": chain_role,
            "confirmed_input_cells": confirmed,
            "detail_body_cells": detail_body,
            "forbidden_non_formula_cells": forbidden,
            "role_rules": {
                "allow_title_inputs": chain_role == "metadata",
                "allow_detail_body_inputs": chain_role == "detail",
                "allow_balance_sheet_inputs": chain_role == "balance_sheet",
                "allow_summary_inputs": False,
            },
            "formula_cells": meta.get("formula_cells", []),
            "protected_text_ranges": meta.get("protected_text_ranges", []),
            "readonly_summary_zone": meta.get("summary_zone", {}),
        }

    payload = {
        "template": layout_map["template"],
        "selected_sheet_count": len(selected_sheets),
        "selected_sheets": selected_sheets,
    }
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(out.as_posix())


if __name__ == "__main__":
    main()
