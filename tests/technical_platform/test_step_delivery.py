import json
import threading
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from test_compound_task import prepared

from asset_based_agent.technical_platform.execution import execute_task
from asset_based_agent.technical_platform.permissions import PermissionService


def reviewed(tmp_path, monkeypatch, *, with_issues=False):
    from asset_based_agent.report_review_app.services import remote_review_llm
    store, session, snapshot = prepared(tmp_path, remote=True)

    class Provider:
        def __init__(self, client, *, model_id, skill_instructions):
            self.model_id, self.skill_instructions = model_id, skill_instructions

        def set_client_job_id(self, value):
            pass

        def review_batches(self, batches, progress_callback=None):
            if with_issues:
                from asset_based_agent.report_review_app.services.rule_registry import (
                    IssueCandidate,
                )
                source = next(tmp_path.glob('runs/*/steps/*/output/history_fragment.docx'))
                text = next(p.text for p in Document(source).paragraphs if p.text.strip())
                chunk = batches[0].chunks[0]
                return [IssueCandidate(source_file_id=chunk.source_file_id,
                    source_file_name=chunk.source_file_name, original_text=text,
                    description='合成测试意见', recommendation='核对来源', category='data_inconsistency',
                    risk_level='low', location={}, confidence=0.8, requires_verification=True)]
            return []

    monkeypatch.setattr(remote_review_llm, 'RemoteReviewLlm', Provider)
    run = store.start_run(session, snapshot)
    PermissionService(store).authorize(run, snapshot, confirmed=True)
    execute_task(store, run, threading.Event(), lambda _: None,
                 client=SimpleNamespace(access_token='synthetic'))
    return store, session, run


def test_step_report_export_uses_only_reviewed_output_and_preserves_checkpoints(tmp_path, monkeypatch):
    from asset_based_agent.technical_platform.report_export import export_review
    from asset_based_agent.technical_platform.review_delivery import step_review_store
    store, session, run = reviewed(tmp_path, monkeypatch)
    before = json.loads(store.run(run)['result'])['steps']
    scoped = step_review_store(store, session, run, 1)
    files = json.loads(scoped.run(run)['snapshot'])['files']
    assert [f['name'] for f in files] == ['history_fragment.docx']
    path = export_review(scoped, run, tmp_path / 'review.docx')
    text = '\n'.join(p.text for p in Document(path).paragraphs)
    assert '资产评估报告审核记录' in text
    assert json.loads(store.run(run)['result'])['steps'] == before
    assert json.loads(step_review_store(store, session, run, 1).run(run)['result'])['exported_report'] == str(path)
    with pytest.raises(ValueError, match='原文件'):
        export_review(scoped, run, files[0]['path'])


def test_step_delivery_rejects_scope_and_immutable_result_changes(tmp_path, monkeypatch):
    from asset_based_agent.technical_platform.review_delivery import step_review_store
    store, session, run = reviewed(tmp_path, monkeypatch)
    for selected_session, ordinal in [('other', 1), (session, -1), (session, 0)]:
        with pytest.raises((ValueError, PermissionError)):
            step_review_store(store, selected_session, run, ordinal)
    scoped = step_review_store(store, session, run, 1)
    changed = json.loads(scoped.run(run)['result'])
    changed['issues'] = [{'description': 'forged'}]
    with pytest.raises(PermissionError):
        scoped.save_result(run, changed)
    with pytest.raises(PermissionError):
        scoped.run('other')


def test_compound_review_has_chat_export_link(tmp_path, monkeypatch):
    from PySide6.QtCore import QUrl
    from PySide6.QtWidgets import QApplication, QFileDialog

    from asset_based_agent.technical_platform.app import PlatformWindow
    app = QApplication.instance() or QApplication([])
    store, session, run = reviewed(tmp_path, monkeypatch)
    window = PlatformWindow(store)
    window.reload_projects(store.session(session)['project'])
    try:
        assert 'zq-step-export:' in window.transcript.toHtml()
        target = tmp_path / 'selected.docx'
        monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *a, **kw: (str(target), ''))
        window.handle_report_link(QUrl(f'zq-step-export:{run}/1'))
        assert target.is_file()
        assert 'zq-step-report:' in window.transcript.toHtml()
    finally:
        window.close()
    assert app is not None


def test_compound_annotation_offer_once_and_decline_preserves_files(tmp_path, monkeypatch):
    from PySide6.QtWidgets import QApplication, QMessageBox

    from asset_based_agent.technical_platform.app import PlatformWindow
    from asset_based_agent.technical_platform.skills import digest
    app = QApplication.instance() or QApplication([])
    store, session, run = reviewed(tmp_path, monkeypatch, with_issues=True)
    before = {str(p): digest(p) for p in tmp_path.rglob('*.docx')}
    window = PlatformWindow(store)
    window.reload_projects(store.session(session)['project'])
    calls = []
    monkeypatch.setattr(QMessageBox, 'question', lambda *a: calls.append(True) or QMessageBox.StandardButton.No)
    try:
        window.offer_annotations(run)
        window.offer_annotations(run)
        assert calls == [True]
        assert 'zq-step-annotate:' in window.transcript.toHtml()
        assert {str(p): digest(p) for p in tmp_path.rglob('*.docx')} == before
        assert json.loads(store.run(run)['result'])['review_deliveries']['inspect']['annotation_prompted']
    finally:
        window.close()
    assert app is not None


def test_compound_annotations_only_copy_reviewed_generated_file(tmp_path, monkeypatch):
    from asset_based_agent.technical_platform.annotations import generate_annotations
    from asset_based_agent.technical_platform.review_delivery import step_review_store
    from asset_based_agent.technical_platform.skills import digest
    store, session, run = reviewed(tmp_path, monkeypatch, with_issues=True)
    scoped = step_review_store(store, session, run, 1)
    files = json.loads(scoped.run(run)['snapshot'])['files']
    before = digest(Path(files[0]['path']))
    checkpoints = json.loads(store.run(run)['result'])['steps']
    records, artifacts = generate_annotations(scoped, run, [1], tmp_path)
    assert len(artifacts) == 1 and records == [{'issue': 1, 'reason': '已标注'}]
    assert digest(Path(files[0]['path'])) == before
    assert json.loads(store.run(run)['result'])['steps'] == checkpoints
    assert json.loads(scoped.run(run)['result'])['annotations'][0]['files'] == artifacts


def test_compound_annotation_accept_runs_worker_and_delivers_chat_link(tmp_path, monkeypatch):
    import time

    from PySide6.QtWidgets import QApplication, QDialog, QFileDialog, QMessageBox

    from asset_based_agent.technical_platform.app import PlatformWindow
    app = QApplication.instance() or QApplication([])
    store, session, run = reviewed(tmp_path, monkeypatch, with_issues=True)
    window = PlatformWindow(store)
    window.reload_projects(store.session(session)['project'])
    monkeypatch.setattr(QMessageBox, 'question', lambda *a: QMessageBox.StandardButton.Yes)
    monkeypatch.setattr(QDialog, 'exec', lambda *a: QDialog.DialogCode.Accepted)
    monkeypatch.setattr(QFileDialog, 'getExistingDirectory', lambda *a: str(tmp_path))
    try:
        window.offer_annotations(run)
        deadline = time.monotonic() + 20
        while window.worker is not None and time.monotonic() < deadline:
            app.processEvents()
            time.sleep(0.01)
        assert window.worker is None
        assert 'zq-step-comment:' in window.transcript.toHtml()
        assert len(list(tmp_path.glob('*_标注版_*.docx'))) == 1
    finally:
        if window.worker is not None:
            window.worker.cancel.set()
            window.worker.wait(20000)
            app.processEvents()
        window.close()
