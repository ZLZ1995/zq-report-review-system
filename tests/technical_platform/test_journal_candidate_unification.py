"""K08: fill and review must select journal evidence from the same candidates.

The writer picked 发生日期/业务内容 with select_journal_entry over the
UNFILTERED journal, while post-generation review re-derived expectations
from a strict candidate set limited to the sheet's own account root.
For A8T this produced systematic disagreements:

- 其他应付款 rows reclassified from 内部往来 (tb_code 1124*) were filled
  from 1124 journal entries but the reviewer only accepted 2241*, so it
  flagged journal_unmatched against evidence that genuinely exists;
- for shared counterparties (管仁良, 平安保险) the writer matched later
  rows on unrelated accounts (6601 expenses), the reviewer matched the
  correct 2241/2202 row, and both reported source_mismatch.

The strict rule now lives in the pipeline and BOTH sides use it, with
one extension: a journal row is also accepted when its tb_code equals
the detail row's own reclassified source account.
"""
import importlib.util
import sys
import unittest
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src'))

from openpyxl import Workbook

SCRIPTS = ROOT / '.codex' / 'skills' / 'valuation-detail-workbook-fill' / 'scripts'


def load_pipeline():
    sys.path.insert(0, str(SCRIPTS))
    spec = importlib.util.spec_from_file_location(
        'run_detail_workbook_pipeline', SCRIPTS / 'run_detail_workbook_pipeline.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


TM_ROW = {
    'tb_code': '1124050000', 'account_name': '内部往来-集团内公司借款（本金）',
    'vendor_name': '浙江天猫技术有限公司', 'counterparty_desc': '浙江天猫技术有限公司',
    'debit': 0.0, 'credit': 209.05, 'gl_date': datetime(2026, 1, 1),
    'summary': 'TMI入账_阿里集团', 'line_desc': 'TMI入账',
}
DISTRACTOR_ROW = {
    'tb_code': '6601010000', 'account_name': '管理费用-差旅费',
    'vendor_name': '浙江天猫技术有限公司', 'counterparty_desc': '浙江天猫技术有限公司',
    'debit': 0.0, 'credit': 500.0, 'gl_date': datetime(2026, 6, 1),
    'summary': '费用报销', 'line_desc': '费用报销',
}


def make_payable_sheet():
    wb = Workbook()
    ws = wb.active
    ws.title = '其他应付款'
    for col in 'ABCDEFGHI':
        ws[f'{col}5'] = '表头'
    ws['A27'] = '合计'
    ws._locked_template_layout = True
    confirmed = [f'{col}{row}' for col in 'ABCDEFGI' for row in range(6, 27)]
    registry = {'selected_sheets': {'其他应付款': {
        'chain_role': 'detail',
        'confirmed_input_cells': confirmed,
        'detail_body_cells': list(confirmed),
        'forbidden_non_formula_cells': [],
        'role_rules': {'allow_title_inputs': False, 'allow_detail_body_inputs': True,
                       'allow_balance_sheet_inputs': False, 'allow_summary_inputs': False},
        'formula_cells': [],
    }}}
    protection = {'sheets': {'其他应付款': {'formula_cells': []}}}
    return ws, registry, protection


class StrictJournalCandidatesTest(unittest.TestCase):
    def test_reclassified_row_code_is_accepted(self):
        pipeline = load_pipeline()
        rows = pipeline.strict_journal_candidates(
            '其他应付款', '浙江天猫技术有限公司', [TM_ROW, DISTRACTOR_ROW],
            row_tb_code='1124050000')
        self.assertEqual([TM_ROW['tb_code']], [r['tb_code'] for r in rows])

    def test_without_row_code_old_strictness_kept(self):
        pipeline = load_pipeline()
        rows = pipeline.strict_journal_candidates(
            '其他应付款', '浙江天猫技术有限公司', [TM_ROW], row_tb_code='')
        self.assertEqual([], rows)

    def test_sheet_root_still_accepted(self):
        pipeline = load_pipeline()
        root_row = {**TM_ROW, 'tb_code': '2241020000', 'account_name': '其他应付款-中转'}
        rows = pipeline.strict_journal_candidates(
            '其他应付款', '浙江天猫技术有限公司', [root_row], row_tb_code='')
        self.assertEqual([root_row], rows)


class FillReviewConsistencyTest(unittest.TestCase):
    def test_fill_uses_strict_candidates_not_latest_global_row(self):
        pipeline = load_pipeline()
        ws, registry, protection = make_payable_sheet()
        settle_row = {
            'tb_code': '2241120000', 'account_name': '其他应付款-员工款',
            'vendor_name': '管仁良', 'counterparty_desc': '管仁良',
            'debit': 0.0, 'credit': 264.0, 'gl_date': datetime(2026, 5, 6),
            'summary': 'SETTLEMENT', 'line_desc': 'ER0211571440',
        }
        later_unrelated = {
            'tb_code': '6601010000', 'account_name': '管理费用-差旅费',
            'vendor_name': '管仁良', 'counterparty_desc': '管仁良',
            'debit': 0.0, 'credit': 480.0, 'gl_date': datetime(2026, 5, 7),
            'summary': 'PAYMENT', 'line_desc': 'EP0123507133',
        }
        rows = [{'counterparty': '管仁良', 'book_value': 480.0,
                 'tb_code': '2241120000', 'sub_name': '', 'remark': '',
                 'fill_mode': 'detail_fillable'}]
        pipeline.write_simple_detail_sheet(
            ws, rows, 'detail', protection, registry, [settle_row, later_unrelated])
        self.assertEqual(datetime(2026, 5, 6), ws['C6'].value,
                         '必须选同科目严格候选，而不是全序时账最新行')
        self.assertEqual('ER0211571440', ws['D6'].value)


if __name__ == '__main__':
    unittest.main()
