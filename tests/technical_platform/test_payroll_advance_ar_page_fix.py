"""P-round: 职工薪酬展开 / 预收账款字段保留 / 应收账款空壳行 的最小失败测试。

T1 职工薪酬：TB 2211 叶子科目应展开为明细行（当前 HEAD 合成单行汇总）。
T2 预收账款：无客商证据时仍须按科目从序时账取最后一笔真实贷方日期（当前 HEAD 不写日期）。
T3 应收账款：零余额空表不得遗留模板/旧项目序号空壳（当前发布件 A6=1 残留）。

修复前：各行为测试稳定失败；修复后全绿。夹具全部为合成数据。
"""
import importlib.util
import sys
import unittest
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src'))

from openpyxl import Workbook

SCRIPTS = ROOT / '.codex' / 'skills' / 'valuation-detail-workbook-fill' / 'scripts'


def load_pipeline():
    sys.path.insert(0, str(SCRIPTS))
    spec = importlib.util.spec_from_file_location(
        'run_detail_workbook_pipeline', SCRIPTS / 'run_detail_workbook_pipeline.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


# ---------------------------------------------------------------- fixtures

PAYROLL_LEAVES = [
    ('2211030100', '应付雇员成本-社会保障-子项甲', 22919.04),
    ('2211030300', '应付雇员成本-社会保障-子项乙', 14037.92),
    ('2211030400', '应付雇员成本-社会保障-子项丙', 716.24),
    ('2211030600', '应付雇员成本-社会保障-子项丁', 286.48),
    ('2211030810', '应付雇员成本-社会保障-个人交纳-子项甲', 11459.52),
    ('2211030830', '应付雇员成本-社会保障-个人交纳-子项乙', 2876.88),
    ('2211030840', '应付雇员成本-社会保障-个人交纳-子项丙', 716.24),
    ('2211040100', '应付雇员成本-奖金-子项戊', 1171883.43),
    ('2211060100', '应付雇员成本-长期子项己', 76610.22),
]
PAYROLL_TOTAL = round(sum(a for _, _, a in PAYROLL_LEAVES), 2)  # 1301505.97


def payroll_tb_rows():
    rows = [
        {'tb_code': '2211010000', 'sub_name': '应付雇员成本-工资', 'book_value': 0.0},
        {'tb_code': '2211030000', 'sub_name': '应付雇员成本-社会保障', 'book_value': 53012.32},
        {'tb_code': '1122010000', 'sub_name': '应收账款-干扰科目', 'book_value': 999.0},
    ]
    for code, name, amount in PAYROLL_LEAVES:
        rows.append({'tb_code': code, 'sub_name': name, 'book_value': amount})
    return rows


def payroll_journal_rows():
    """每个叶子科目两笔真实贷方（旧+新），加冲销噪声；期望取 2026-07-31。"""
    rows = []
    for code, name, amount in PAYROLL_LEAVES:
        rows.append({'tb_code': code, 'account_name': name, 'credit': amount,
                     'debit': 0.0, 'gl_date': datetime(2026, 6, 30),
                     'summary': '计提上月', 'line_desc': '计提', 'vendor_name': '', 'counterparty_desc': ''})
        rows.append({'tb_code': code, 'account_name': name, 'credit': amount,
                     'debit': 0.0, 'gl_date': datetime(2026, 7, 31),
                     'summary': '计提本月', 'line_desc': '计提', 'vendor_name': '', 'counterparty_desc': ''})
        rows.append({'tb_code': code, 'account_name': name, 'credit': 1.0,
                     'debit': 0.0, 'gl_date': datetime(2026, 7, 31),
                     'summary': '冲销调整', 'line_desc': '冲销', 'vendor_name': '', 'counterparty_desc': ''})
    return rows


def make_locked_sheet(title, headers, total_row=27):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for col, text in headers.items():
        ws[f'{col}5'] = text
    ws[f'A{total_row}'] = '合计'
    ws._locked_template_layout = True
    cols = list(headers.keys())
    confirmed = [f'{col}{r}' for col in cols for r in range(6, total_row)]
    registry = {'selected_sheets': {title: {
        'chain_role': 'detail',
        'confirmed_input_cells': list(confirmed),
        'detail_body_cells': list(confirmed),
        'forbidden_non_formula_cells': [],
        'role_rules': {'allow_title_inputs': False, 'allow_detail_body_inputs': True,
                       'allow_balance_sheet_inputs': False, 'allow_summary_inputs': False},
        'formula_cells': [],
    }}}
    protection = {'sheets': {title: {'formula_cells': []}}}
    return wb, ws, registry, protection


def run_stage2(pipeline, wb, sheet_plan, bs_values, journal_rows):
    registry = {'selected_sheets': {}}
    protection = {'sheets': {}}
    for title in {name for name, _, _ in sheet_plan}:
        ws = wb[title]
        total_row = 27
        cols = [c for c in 'ABCDEFGHIJKLMNOP' if ws[f'{c}5'].value is not None]
        confirmed = [f'{c}{r}' for c in cols for r in range(6, total_row)]
        registry['selected_sheets'][title] = {
            'chain_role': 'detail', 'confirmed_input_cells': confirmed,
            'detail_body_cells': list(confirmed), 'forbidden_non_formula_cells': [],
            'role_rules': {'allow_detail_body_inputs': True}, 'formula_cells': [],
        }
        protection['sheets'][title] = {'formula_cells': []}
    bs = {'values': bs_values, 'values_prior': {}}
    return pipeline.stage2_fill_detail_pages_from_trial_balance(
        wb, sheet_plan, bs, journal_rows, protection, registry)


# ------------------------------------------------------------------ tests

class ModuleIsolationTest(unittest.TestCase):
    def test_pipeline_module_loaded_from_this_repo(self):
        pipeline = load_pipeline()
        self.assertTrue(
            str(Path(pipeline.__file__).resolve()).startswith(str(ROOT)),
            f'管线模块必须来自本仓库: {pipeline.__file__}')


class T1PayrollExpandTest(unittest.TestCase):
    def test_stage2_expands_2211_leaf_rows(self):
        """TB 2211 叶子科目必须展开为明细行，父级/零值/干扰科目不得混入。"""
        pipeline = load_pipeline()
        headers = {'A': '序号', 'B': '结算内容', 'C': '发生日期', 'F': '账面价值', 'H': '备注'}
        wb, ws, _, _ = make_locked_sheet('职工薪酬', headers)
        sheet_plan = [('职工薪酬', payroll_tb_rows(), 'detail')]
        bs_values = {'应付职工薪酬': PAYROLL_TOTAL}
        _, _completed, _, written_rows, _ = run_stage2(
            pipeline, wb, sheet_plan, bs_values, payroll_journal_rows())
        payroll_written = [w for w in written_rows if w.get('_sheet') == '职工薪酬']
        self.assertEqual(9, len(payroll_written),
                         f'2211 九个非零叶子应展开为 9 行，实际 {len(payroll_written)} 行')
        names = [ws[f'B{r}'].value for r in range(6, 15)]
        for leaf in ('子项甲', '子项戊', '子项己'):
            self.assertTrue(any(leaf in (n or '') for n in names),
                            f'叶子名称 {leaf} 应可追溯地出现在明细中: {names}')
        self.assertNotIn('应付雇员成本-社会保障', [n for n in names if n],
                         '父级汇总行不得作为明细行重复出现')
        amounts = [ws[f'F{r}'].value for r in range(6, 15)]
        self.assertEqual(PAYROLL_TOTAL, round(sum(float(a) for a in amounts), 2))
        for r in range(6, 15):
            self.assertIsNotNone(ws[f'C{r}'].value, f'C{r} 发生日期不得为空')

    def test_build_payroll_rows_records_date_source_and_lineage(self):
        """新纯函数：日期取最后一笔真实贷方业务，且记录来源与策略。"""
        pipeline = load_pipeline()
        result = pipeline.build_payroll_rows(
            payroll_tb_rows(), payroll_journal_rows(),
            bs_total=PAYROLL_TOTAL, report_date=datetime(2026, 7, 31),
            date_policy='journal_last_real_credit')
        self.assertEqual('ok', result['status'])
        rows = result['rows']
        self.assertEqual(9, len(rows))
        for row in rows:
            self.assertEqual(datetime(2026, 7, 31), row['date_value'])
            self.assertTrue(row.get('date_source'), '必须记录 date_source')
            self.assertTrue(row.get('source_account_code'), '必须保留来源科目代码')
            self.assertEqual('detail_fillable', row['fill_mode'])

    def test_build_payroll_rows_blocks_on_total_mismatch(self):
        """叶子合计与 BS 不一致：阻断，不得静默回退单行汇总。"""
        pipeline = load_pipeline()
        result = pipeline.build_payroll_rows(
            payroll_tb_rows(), payroll_journal_rows(),
            bs_total=PAYROLL_TOTAL + 100.0, report_date=datetime(2026, 7, 31),
            date_policy='journal_last_real_credit')
        self.assertEqual('blocked', result['status'])
        self.assertNotEqual(1, len(result.get('rows', [])),
                            '不一致时不得回退为单行汇总伪装成功')
        self.assertTrue(result.get('unreconciled_reason'))

    def test_stage2_postfix_does_not_overwrite_payroll_detail(self):
        """stage2 postfix 不得把已展开的 9 行覆盖回单行汇总。"""
        pipeline = load_pipeline()
        wb = Workbook()
        ws = wb.active
        ws.title = '职工薪酬'
        for idx, (_, name, amount) in enumerate(PAYROLL_LEAVES, start=1):
            r = 5 + idx
            ws[f'A{r}'] = idx
            ws[f'B{r}'] = name
            ws[f'C{r}'] = datetime(2026, 7, 31)
            ws[f'F{r}'] = amount
            ws[f'G{r}'] = f'=F{r}'
        bs = {'values': {'应付职工薪酬': PAYROLL_TOTAL}, 'values_prior': {}}
        pipeline.stage2_postfix_key_sheets(wb, bs)
        self.assertIn('子项甲', ws['B6'].value or '',
                        f'postfix 不得覆盖已写入明细, B6={ws["B6"].value!r}')
        self.assertEqual(PAYROLL_LEAVES[0][2], ws['F6'].value)
        self.assertIn('子项己', ws['B14'].value or '')


class T2AdvanceReceiptsTest(unittest.TestCase):
    def _stage2_advance(self, pipeline, journal_rows, bs_amount=-0.32):
        headers = {'A': '序号', 'B': '户名', 'C': '发生日期', 'D': '业务内容', 'G': '账面价值'}
        wb, ws, _, _ = make_locked_sheet('预收账款', headers)
        source_rows = [{
            'counterparty': '', 'sub_name': '递延收入-销项税暂估',
            'book_value': bs_amount, 'tb_code': '2203010000',
            'fill_mode': 'detail_fillable', 'fill_reason': '',
        }]
        sheet_plan = [('预收账款', source_rows, 'detail')]
        _, _, _, written_rows, _ = run_stage2(
            pipeline, wb, sheet_plan, {'预收账款': bs_amount}, journal_rows)
        return wb, ws, [w for w in written_rows if w.get('_sheet') == '预收账款']

    def test_date_filled_from_account_journal_without_counterparty(self):
        """无客商证据时，仍须按 2203 科目取最后一笔真实贷方业务日期。"""
        pipeline = load_pipeline()
        journal = [
            {'tb_code': '2203010000', 'account_name': '预收账款-递延收入',
             'credit': 0.32, 'debit': 0.0, 'gl_date': datetime(2025, 12, 31),
             'summary': '计提利息', 'line_desc': '销项税暂估', 'vendor_name': '', 'counterparty_desc': ''},
            {'tb_code': '2203010000', 'account_name': '预收账款-递延收入',
             'credit': 5.0, 'debit': 0.0, 'gl_date': datetime(2026, 1, 5),
             'summary': '冲销暂估', 'line_desc': '冲销', 'vendor_name': '', 'counterparty_desc': ''},
        ]
        _, ws, written = self._stage2_advance(pipeline, journal)
        self.assertEqual(1, len(written))
        self.assertEqual(datetime(2025, 12, 31), ws['C6'].value,
                         f'发生日期应来自序时账最后一笔真实贷方业务, 实际 C6={ws["C6"].value!r}')

    def test_business_label_not_written_to_counterparty(self):
        """销项税暂估是业务/税务标签，不得进入户名列。"""
        pipeline = load_pipeline()
        _, ws, _written = self._stage2_advance(pipeline, [])
        self.assertIn(ws['B6'].value, (None, ''),
                      f'户名列不得写入业务/税务标签: {ws["B6"].value!r}')

    def test_no_synthetic_row_when_source_differs_from_bs(self):
        """来源合计与 BS 有差额：保留真实行并报告，不得生成调节/凑数行。"""
        pipeline = load_pipeline()
        _, ws, written = self._stage2_advance(pipeline, [], bs_amount=-0.32)
        self.assertEqual(1, len(written), '不得追加合成调节行')
        self.assertEqual(-0.32, ws['G6'].value)


class T3EmptyReceivablePageTest(unittest.TestCase):
    def _make_ar_sheet(self, pipeline, with_legacy_a6=True):
        headers = {'A': '序号', 'B': '欠款单位名称', 'C': '业务内容', 'D': '发生日期', 'F': '审计前账面值'}
        wb, ws, _, _ = make_locked_sheet('应收账款', headers)
        if with_legacy_a6:
            ws['A6'] = 1  # 模板/旧项目遗留序号
        return wb, ws

    def test_empty_ar_page_has_no_orphan_row_index(self):
        """零余额、空来源：数据体不得残留孤立序号 A6。"""
        pipeline = load_pipeline()
        wb, ws = self._make_ar_sheet(pipeline)
        sheet_plan = [('应收账款', [], 'detail')]
        run_stage2(pipeline, wb, sheet_plan, {'应收账款': 0.0}, [])
        self.assertIsNone(ws['A6'].value,
                          f'空表不得残留孤立序号, A6={ws["A6"].value!r}')

    def test_empty_ar_page_preserves_footer_and_formulas(self):
        """空表清理不得破坏合计行、页脚结构。"""
        pipeline = load_pipeline()
        wb, ws = self._make_ar_sheet(pipeline)
        ws['F27'] = '=SUM(F6:F26)'
        sheet_plan = [('应收账款', [], 'detail')]
        run_stage2(pipeline, wb, sheet_plan, {'应收账款': 0.0}, [])
        self.assertEqual('合计', ws['A27'].value)
        self.assertEqual('=SUM(F6:F26)', ws['F27'].value)

    def test_page_state_zero_balance_cleanup(self):
        pipeline = load_pipeline()
        state = pipeline.resolve_detail_page_state('应收账款', 0.0, [], {'in_scope': True})
        self.assertEqual('zero_balance_cleanup', state)

    def test_page_state_nonzero_bs_without_rows_is_not_empty_cleanup(self):
        """BS 非零但无明细来源：placeholder 或阻断，不得按普通空表处理。"""
        pipeline = load_pipeline()
        state = pipeline.resolve_detail_page_state('应收账款', 5000.0, [], {'in_scope': True})
        self.assertIn(state, ('placeholder_only', 'blocked'),
                      f'BS 非零但无明细不得视为普通空表: {state}')

    def test_page_state_detail_fillable_with_rows(self):
        pipeline = load_pipeline()
        rows = [{'counterparty': '某公司', 'book_value': 10.0}]
        state = pipeline.resolve_detail_page_state('应收账款', 10.0, rows, {'in_scope': True})
        self.assertEqual('detail_fillable', state)


if __name__ == '__main__':
    unittest.main()


class T4PayrollReviewClassificationTest(unittest.TestCase):
    """P09 复测发现：职工薪酬写入行 counterparty 刻意为空（B 列为费用项目），
    与 应交税费/银行存款 同属非客商明细表。审查层必须按非客商表分类，
    不得把空 counterparty 当语义异常、不得按客商表要求逐格原始资料映射。"""

    def _make_payroll_workbook(self, tmp_dir):
        wb = Workbook()
        ws = wb.active
        ws.title = '职工薪酬'
        wb_path = Path(tmp_dir) / 'payroll_review.xlsx'
        wb.save(wb_path)
        wb.close()
        return wb_path

    def test_payroll_sheet_is_semantic_exempt(self):
        pipeline = load_pipeline()
        self.assertIn('职工薪酬', pipeline.SEMANTIC_EXEMPT_SHEETS,
                      '职工薪酬是非客商明细表，必须进入 SEMANTIC_EXEMPT_SHEETS')

    def test_empty_payroll_counterparty_is_not_anomaly(self):
        import tempfile
        pipeline = load_pipeline()
        with tempfile.TemporaryDirectory() as tmp_dir:
            wb_path = self._make_payroll_workbook(tmp_dir)
            written_rows = [
                {'_sheet': '职工薪酬', '_written_row': 6, 'counterparty': '',
                 'sub_name': '工资', 'book_value': 100.0, 'tb_code': '2211010000',
                 'fill_mode': 'detail_fillable'},
                {'_sheet': '职工薪酬', '_written_row': 7, 'counterparty': '',
                 'sub_name': '社会保障-子项甲', 'book_value': 50.0, 'tb_code': '2211030100',
                 'fill_mode': 'detail_fillable'},
            ]
            report = pipeline.build_counterparty_anomaly_report(wb_path, [], written_rows)
        self.assertEqual([], report['anomalies'],
                         f'职工薪酬空 counterparty 不得计为语义异常: {report["anomalies"]}')

    def test_other_sheet_empty_counterparty_still_flagged(self):
        """豁免不得扩大：六往来表空客商仍必须报异常（守卫既有强度）。"""
        import tempfile
        pipeline = load_pipeline()
        with tempfile.TemporaryDirectory() as tmp_dir:
            wb_path = self._make_payroll_workbook(tmp_dir)
            written_rows = [
                {'_sheet': '其他应付款', '_written_row': 6, 'counterparty': '',
                 'sub_name': '某业务', 'book_value': 10.0, 'tb_code': '2241010000',
                 'fill_mode': 'detail_fillable'},
            ]
            report = pipeline.build_counterparty_anomaly_report(wb_path, [], written_rows)
        self.assertEqual(1, len(report['anomalies']))
        self.assertEqual('empty', report['anomalies'][0]['reason'])
