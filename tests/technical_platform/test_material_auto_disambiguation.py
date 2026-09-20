"""K04: deterministic same-entity multi-period disambiguation for TB/TBD.

Earlier trial balances / journals are historical reference material, never
comparison statements (which the pipeline treats as --financial-statement).
Only genuine conflicts may ask; file names never override in-table evidence.
"""
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src'))

from openpyxl import Workbook

ENTITY = 'A8T'
FULL_NAME = '阿里云飞天（北京）云计算有限公司'


def make_tb(path, period, entity=ENTITY):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet0'
    ws['A1'] = '科目汇总试算表'
    if entity is not None:
        ws['A2'] = f'公司: {entity} - {FULL_NAME}       期间: {period}    PRCGAAP'
    else:
        ws['A2'] = f'期间: {period}'
    ws['A4'] = '科目编码'
    ws['B4'] = '科目名称'
    wb.save(path)


def make_tbd(path, periods, entity=ENTITY):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet0'
    ws['A1'] = '总帐明细帐追溯报表'
    if entity is not None:
        ws['A2'] = f'公司：{entity}-{entity}帐套：PRCGAAP(CNY)'
    ws['A3'] = '期间'
    ws['B3'] = '凭证号'
    for row, period in enumerate(periods, start=4):
        ws.cell(row=row, column=1, value=period)
    wb.save(path)


def make_bs(path, period, entity=ENTITY):
    wb = Workbook()
    ws = wb.active
    ws.title = '销项 1 (A8T)'
    ws['A1'] = 'PRC-资产负债表'
    ws['A3'] = f'本期：{period}'
    ws['A7'] = f'公司={entity} ({FULL_NAME})'
    ws['A9'] = '资产'
    ws['E9'] = '期末余额'
    wb.save(path)


def files_arg(*paths):
    return [{'id': f'f{i}', 'name': p.name, 'path': str(p)} for i, p in enumerate(paths)]


def plan_for(role_by_index):
    return {'assignments': [{'file_id': f'f{i}', 'role': role, 'reason': '表头'}
                            for i, role in role_by_index.items()]}


class TrialBalanceDisambiguationTest(unittest.TestCase):
    def setUp(self):
        import tempfile
        self.tmp = tempfile.TemporaryDirectory(dir=r'D:\ZQ-Acceptance\tmp-pytest-basetemp')
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def resolve(self, plan, *paths):
        from asset_based_agent.technical_platform.material_analysis import resolve_materials
        return resolve_materials(plan, files_arg(*paths))

    def test_multi_period_trial_balance_selects_latest_as_reference_not_comparison(self):
        paths = []
        for i, period in enumerate(('2024-12', '2025-12', '2026-07')):
            path = self.dir / f'{period}_TB.xlsx'
            make_tb(path, period)
            paths.append(path)
        resolution = self.resolve(plan_for({0: 'trial_balance', 1: 'trial_balance',
                                            2: 'trial_balance'}), *paths)
        self.assertEqual('resolved', resolution.status)
        self.assertEqual('f2', resolution.selected['trial_balance'])
        self.assertEqual((), resolution.comparison_artifact_ids)
        self.assertEqual({'f0', 'f1'}, set(resolution.reference_artifact_ids))
        self.assertTrue(any('历史参考' in reason for reason in resolution.reasons))

    def test_multi_period_journal_selects_latest_by_max_entry_period(self):
        paths = []
        for i, periods in enumerate((('2024-01', '2024-12'), ('2026-01', '2026-07'))):
            path = self.dir / f'tbd_{i}.xlsx'
            make_tbd(path, periods)
            paths.append(path)
        resolution = self.resolve(plan_for({0: 'journal', 1: 'journal'}), *paths)
        self.assertEqual('resolved', resolution.status)
        self.assertEqual('f1', resolution.selected['journal'])
        self.assertEqual(('f0',), resolution.reference_artifact_ids)
        self.assertEqual((), resolution.comparison_artifact_ids)

    def test_a8t_set_resolves_without_questions_and_pairs_tb_tbd(self):
        bs = self.dir / 'A8T-BS202607.xlsx'
        make_bs(bs, '2026-07')
        tb_old = self.dir / '2025_TB.xlsx'
        make_tb(tb_old, '2025-12')
        tb_new = self.dir / '2026.07_TB.xlsx'
        make_tb(tb_new, '2026-07')
        tbd_old = self.dir / '2025_TBD.xlsx'
        make_tbd(tbd_old, ('2025-01', '2025-12'))
        tbd_new = self.dir / '2026.1-7_TBD.xlsx'
        make_tbd(tbd_new, ('2026-01', '2026-07'))
        resolution = self.resolve(
            plan_for({0: 'balance_sheet', 1: 'trial_balance', 2: 'trial_balance',
                      3: 'journal', 4: 'journal'}),
            bs, tb_old, tb_new, tbd_old, tbd_new)
        self.assertEqual('resolved', resolution.status, resolution.questions)
        self.assertEqual((), resolution.questions)
        self.assertEqual('f0', resolution.selected['balance_sheet'])
        self.assertEqual('f2', resolution.selected['trial_balance'])
        self.assertEqual('f4', resolution.selected['journal'])
        self.assertEqual({'f1', 'f3'}, set(resolution.reference_artifact_ids))
        self.assertTrue(any('配对' in reason for reason in resolution.reasons),
                        resolution.reasons)

    def test_filename_does_not_override_in_table_period(self):
        older_named = self.dir / '2023_TB.xlsx'
        make_tb(older_named, '2026-07')  # table says 2026-07 despite the name
        newer_named = self.dir / '2026_TB.xlsx'
        make_tb(newer_named, '2024-12')
        resolution = self.resolve(plan_for({0: 'trial_balance', 1: 'trial_balance'}),
                                  older_named, newer_named)
        self.assertEqual('resolved', resolution.status)
        self.assertEqual('f0', resolution.selected['trial_balance'])

    def test_tb_missing_entity_waits_user_with_minimal_question(self):
        paths = []
        for i, entity in enumerate((ENTITY, None)):
            path = self.dir / f'tb_{i}.xlsx'
            make_tb(path, '2026-07', entity=entity)
            paths.append(path)
        # different periods so only the missing entity blocks resolution
        make_tb(paths[1], '2025-12', entity=None)
        resolution = self.resolve(plan_for({0: 'trial_balance', 1: 'trial_balance'}), *paths)
        self.assertEqual('waiting_user', resolution.status)
        self.assertEqual(1, len(resolution.questions))
        self.assertIn('主体', resolution.questions[0])
        self.assertIn('tb_1.xlsx', resolution.questions[0])

    def test_tb_missing_period_waits_user(self):
        first = self.dir / 'tb_a.xlsx'
        make_tb(first, '2026-07')
        second = self.dir / 'tb_b.xlsx'
        wb = Workbook()
        ws = wb.active
        ws['A1'] = '科目汇总试算表'
        ws['A2'] = f'公司: {ENTITY} - {FULL_NAME}'
        wb.save(second)
        resolution = self.resolve(plan_for({0: 'trial_balance', 1: 'trial_balance'}),
                                  first, second)
        self.assertEqual('waiting_user', resolution.status)
        self.assertTrue(any('期间' in question for question in resolution.questions))

    def test_same_period_duplicate_versions_wait_user(self):
        paths = []
        for i in range(2):
            path = self.dir / f'tb_v{i}.xlsx'
            make_tb(path, '2026-07')
            paths.append(path)
        resolution = self.resolve(plan_for({0: 'trial_balance', 1: 'trial_balance'}), *paths)
        self.assertEqual('waiting_user', resolution.status)
        self.assertTrue(any('版本' in question or '为准' in question
                            for question in resolution.questions))

    def test_tb_tbd_period_mismatch_waits_user(self):
        tb = self.dir / 'tb.xlsx'
        make_tb(tb, '2026-07')
        tbd_a = self.dir / 'tbd_a.xlsx'
        make_tbd(tbd_a, ('2025-01', '2025-12'))
        tbd_b = self.dir / 'tbd_b.xlsx'
        make_tbd(tbd_b, ('2024-01', '2024-12'))
        resolution = self.resolve(plan_for({0: 'trial_balance', 1: 'journal', 2: 'journal'}),
                                  tb, tbd_a, tbd_b)
        self.assertEqual('waiting_user', resolution.status)
        self.assertTrue(any('期间' in question for question in resolution.questions))

    def test_cross_role_entity_conflict_waits_user(self):
        tb_a = self.dir / 'tb_a.xlsx'
        make_tb(tb_a, '2026-07', entity='A8T')
        tb_b = self.dir / 'tb_b.xlsx'
        make_tb(tb_b, '2025-12', entity='A8T')
        tbd = self.dir / 'tbd.xlsx'
        make_tbd(tbd, ('2026-01', '2026-07'), entity='OTHER')
        resolution = self.resolve(plan_for({0: 'trial_balance', 1: 'trial_balance',
                                            2: 'journal'}), tb_a, tb_b, tbd)
        self.assertEqual('waiting_user', resolution.status)
        self.assertTrue(any('主体' in question for question in resolution.questions))


if __name__ == '__main__':
    unittest.main()
