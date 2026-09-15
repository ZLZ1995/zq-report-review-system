import importlib
import json

import pytest
from openpyxl import Workbook

from test_detail_workbook_pipeline_guards import load_pipeline_module


def review_module():
    load_pipeline_module()
    return importlib.import_module('post_generation_review')


def test_final_value_checked_against_reopened_source(tmp_path):
    m = review_module()
    source = tmp_path / 'source.xlsx'
    final = tmp_path / 'final.xlsx'
    wb = Workbook()
    wb.active['B5'] = 100
    wb.save(source)
    wb.active['B5'] = 90
    wb.save(final)
    checks = [{'sheet': 'Sheet', 'cell': 'B5', 'source': str(source),
               'source_sheet': 'Sheet', 'source_cell': 'B5', 'kind': 'amount'}]
    report = m.review_source_cells(final, checks)
    assert report['status'] == 'fail'
    assert report['issues'][0]['expected'] == 100
    assert report['issues'][0]['actual'] == 90
    assert report['issues'][0]['difference'] == -10
    assert report['issues'][0]['status'] == 'unresolved'
    wb.active['B5'] = 90
    wb.save(source)
    assert m.review_source_cells(final, checks)['status'] == 'pass'


def test_failed_release_preserves_previous_delivery_and_reports(tmp_path):
    m = review_module()
    staging = tmp_path / 'staging.xlsx'
    target = tmp_path / 'published.xlsx'
    staging.write_bytes(b'new')
    target.write_bytes(b'previous approved file')
    with pytest.raises(m.ReviewBlocked):
        m.publish_after_review(staging, target, tmp_path,
            [{'status': 'fail', 'issues': [{'sheet': '应付账款', 'cell': 'G6',
               'reason': 'source_mismatch', 'expected': 100, 'actual': 90}]}])
    assert target.read_bytes() == b'previous approved file'
    assert '应付账款' in (tmp_path / 'user_feedback.md').read_text(encoding='utf-8')
    assert json.loads((tmp_path / 'completion_status.json').read_text())['status'] == 'blocked'


def test_success_requires_explicit_passing_reports(tmp_path):
    m = review_module()
    wb = Workbook()
    staging = tmp_path / 'staging.xlsx'
    target = tmp_path / 'published.xlsx'
    wb.save(staging)
    with pytest.raises(m.ReviewBlocked):
        m.publish_after_review(staging, target, tmp_path, [])
    assert not target.exists()
    m.publish_after_review(staging, target, tmp_path, [{'status': 'pass', 'issues': []}])
    assert target.read_bytes() == staging.read_bytes()


def test_uncalculated_never_promoted_to_ok(tmp_path):
    m = load_pipeline_module()
    wb = Workbook()
    wb.active.title = '资产负债表'
    wb.active['D38'] = wb.active['I38'] = 100
    wb.create_sheet('分类汇总')['J4'] = '=IF(1=1,"OK","出错")'
    p = tmp_path / 'uncalculated.xlsx'
    wb.save(p)
    result = m.stage4_validate_and_self_check(p, [], [], [], {})[0]
    assert result['classification_j4'] != 'OK'


def test_stage1_calculates_before_reading_balance_gate(tmp_path, monkeypatch):
    m = load_pipeline_module()
    events = []
    monkeypatch.setattr(m, 'force_excel_recalc', lambda path: events.append('calculate'))
    def read(path):
        events.append('read')
        return {'assets_equal_liabilities_equity': events == ['calculate', 'read']}
    monkeypatch.setattr(m, 'build_validation_report', read)
    assert m.validate_saved_stage1(tmp_path / 'staging.xlsx', allow_recalc=True)['assets_equal_liabilities_equity']
    assert events == ['calculate', 'read']


def test_locked_template_guard_covers_detail_and_balance_formulas():
    m = load_pipeline_module()
    wb = Workbook()
    wb.active.title = '资产负债表'
    wb.active['D38'] = '=SUM(D7:D37)'
    wb.create_sheet('应付账款')['G28'] = '=SUM(G6:G27)'
    baseline = m.capture_template_formulas(wb)
    wb['应付账款']['G28'] = 1000
    with pytest.raises(m.ProtectionViolation, match='template_formula_changed'):
        m.assert_template_formulas_preserved(wb, baseline)


def test_postfix_preserves_balance_sheet_formulas():
    m = load_pipeline_module()
    wb = Workbook()
    wb.active.title = '资产负债表'
    wb.active['D38'] = '=SUM(D18,D37)'
    m.stage2_postfix_key_sheets(wb, {'values': {'资产总计': 1000}})
    assert wb.active['D38'].value == '=SUM(D18,D37)'


def test_difference_does_not_create_accounting_adjustment():
    m = load_pipeline_module()
    rows = [{'counterparty': '甲公司', 'book_value': 100}]
    assert m.append_net_reconciliation_row('应付账款', rows, {'应付账款': 120}) == rows


def test_review_journal_candidates_require_party_account_and_direction():
    from datetime import datetime
    m, p = review_module(), load_pipeline_module()
    good = dict(vendor_name='甲公司', account_name='应付账款', tb_code='2202',
        debit=0, credit=100, gl_date=datetime(2024, 10, 1), summary='采购服务')
    rows = [good, {**good, 'vendor_name': '乙公司', 'summary': '转付甲公司'},
        {**good, 'account_name': '其他应付款', 'tb_code': '2241'},
        {**good, 'credit': 0, 'debit': 100}, {**good, 'credit': -100}]
    assert m.strict_journal_candidates('应付账款', '甲公司', rows, p) == [good]


def test_repair_only_unique_input_and_rechecks(tmp_path):
    m = review_module()
    wb = Workbook()
    wb.active['B5'] = 100
    source, final = tmp_path / 'source.xlsx', tmp_path / 'final.xlsx'
    wb.save(source)
    wb.active['B5'] = 90
    wb.active['C5'] = '=B5'
    wb.save(final)
    checks = [{'sheet': 'Sheet', 'cell': c, 'source': str(source),
               'source_sheet': 'Sheet', 'source_cell': 'B5', 'kind': 'amount'} for c in ['B5','C5']]
    registry = {'selected_sheets': {'Sheet': {'confirmed_input_cells': ['B5','C5']}}}
    result = m.repair_unique_source_cells(final, checks, registry)
    assert len(result['repairs']) == 1
    assert result['repairs'][0]['cell'] == 'B5'
    assert result['recheck']['status'] == 'fail'  # Formula not calculated, not overwritten.


def test_excel_recalc_keeps_formula_and_cache(tmp_path):
    pytest.importorskip('win32com.client')
    review_module()
    m = importlib.import_module('recalculate_readonly')
    from openpyxl import load_workbook
    wb = Workbook()
    wb.active['A1'] = 100
    wb.active['A2'] = '=A1+37'
    wb.active['A3'] = '=IF(A2=137,"OK","ERROR")'
    path = tmp_path / 'calculated.xlsx'
    wb.save(path)
    m.recalculate_formula_caches(path)
    result = load_workbook(path, data_only=True)
    assert result.active['A2'].value == 137
    assert result.active['A3'].value == 'OK'
    formula = load_workbook(path, data_only=False)
    assert formula.active['A2'].value == '=A1+37'


def test_bank_source_review_reads_original_not_writer_record(tmp_path):
    p, m = load_pipeline_module(), review_module()
    source = tmp_path / 'bank.xlsx'
    wb = Workbook()
    wb.active.append(['账号', '账户名称', '开户行', '交易日期', '账户余额'])
    wb.active.append(['1234567890', '测试公司', '测试银行', '2026/03/31', 100])
    wb.save(source)
    output = Workbook()
    output.active.title = '银行存款'
    output.active['B6'], output.active['C6'], output.active['I6'] = '测试银行', '1234567890', 100
    record = [{'_sheet': '银行存款', '_written_row': 6, 'sub_name': '1234567890', 'book_value': 999}]
    checks, problems = m.review_bank_sources(output, [str(source)], {'货币资金': 100}, record, p)
    assert not problems and len(checks) == 3
    output.active['I6'] = 99
    _, problems = m.review_bank_sources(output, [str(source)], {'货币资金': 100}, record, p)
    assert any(x['reason'] == 'source_mismatch' for x in problems)


def test_cli_failure_always_emits_user_feedback(tmp_path):
    import subprocess
    import sys
    from pathlib import Path
    script = Path('.codex/skills/valuation-detail-workbook-fill/scripts/run_detail_workbook_pipeline.py')
    flags = ['trial-balance','balance-sheet','template','project-mapping','formula-chain',
             'sheet-layout','formula-protection','input-cell-registry','summary-chain-input-registry']
    result = subprocess.run([sys.executable, str(script), '--output-dir', str(tmp_path),
        '--published-workbook', str(tmp_path / 'delivery.xlsx'),
        *[token for flag in flags for token in ['--' + flag, str(tmp_path / 'missing')]]], capture_output=True)
    assert result.returncode != 0
    assert not (tmp_path / 'delivery.xlsx').exists()
    assert (tmp_path / 'user_feedback.md').exists()
    assert json.loads((tmp_path / 'completion_status.json').read_text())['status'] == 'blocked'


@pytest.mark.parametrize('source_account', ['应付账款', '无法识别的往来'])
def test_pipeline_review_reopens_actual_counterparty_source(tmp_path, source_account):
    from types import SimpleNamespace
    from test_detail_cover_metadata import statement
    from openpyxl import load_workbook
    p = load_pipeline_module()
    m = review_module()
    bs = statement(tmp_path / 'bs.xlsx', '2026年3月31日')
    raw = load_workbook(bs)
    raw['资产负债表']['B8'] = '资产总计'
    raw['资产负债表']['C8'] = 100
    raw['资产负债表']['D8'] = 80
    raw.save(bs)
    cp, tb, final = [tmp_path / name for name in ['cp.xlsx', 'tb.xlsx', 'final.xlsx']]
    wb = Workbook()
    wb.active['B5'], wb.active['C5'], wb.active['D5'] = '甲公司', '2202', source_account
    wb.active['F5'], wb.active['G5'] = 0, 100
    wb.save(cp)
    Workbook().save(tb)
    wb = Workbook()
    wb.active.title = '资产负债表'
    wb.active['B8'], wb.active['D8'], wb.active['C8'] = '资产总计', 100, 80
    ws = wb.create_sheet('应付账款')
    ws['B6'], ws['G6'], ws['A7'] = '甲公司', 90, '合计'
    wb.save(final)
    args = SimpleNamespace(source_checks=None, financial_statement=[str(bs)],
        balance_sheet=str(bs), journal=None, counterparty_balance=str(cp), trial_balance=str(tb))
    hashes = m.source_fingerprints([bs, cp, tb])
    report = m.review_pipeline_sources(final, args, p, [], hashes)
    assert report['status'] == 'fail'
    if source_account == '应付账款':
        problem = next(x for x in report['issues'] if x.get('counterparty') == '甲公司')
        assert (problem['expected'], problem['actual'], problem['difference']) == (100, 90, -10)
        assert problem['source_evidence'][0]['source_row'] == 5
        raw = load_workbook(cp)
        raw.active['G5'] = 90
        raw.save(cp)
        changed = m.review_pipeline_sources(final, args, p, [], hashes)
        assert any(x['reason'] == 'source_changed' for x in changed['issues'])
        assert not any(x.get('counterparty') == '甲公司' for x in changed['issues'])
    else:
        assert any(x['reason'] == 'source_unverified' and x.get('source_row') == 5 for x in report['issues'])
