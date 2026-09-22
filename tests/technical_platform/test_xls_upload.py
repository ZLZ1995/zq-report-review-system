"""XLS intake: picker, drag-drop, capability contract, hash fidelity."""
import os

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

import pytest

from asset_based_agent.technical_platform.skills import digest
from asset_based_agent.technical_platform.store import PlatformStore


def make_xls(path):
    import xlwt
    book = xlwt.Workbook()
    sheet = book.add_sheet('资产负债表')
    sheet.write(0, 0, '资产负债表')
    sheet.write(1, 0, '编制单位：测试公司 2026年7月31日 单位：元')
    book.save(str(path))


def make_xlsx(path):
    from openpyxl import Workbook
    book = Workbook()
    book.active.title = '资产负债表'
    book.save(path)


def build_window(tmp_path):
    from PySide6.QtWidgets import QApplication

    from asset_based_agent.technical_platform.app import PlatformWindow
    assert QApplication.instance() or QApplication([])
    store = PlatformStore(tmp_path / 'state.db', 'test')
    project = store.create_project('xls接入')
    store.create_session(project)
    window = PlatformWindow(store)
    window.reload_projects(project)
    return window, store, project


def drop_urls(window, paths):
    from PySide6.QtCore import QMimeData, QPointF, Qt, QUrl
    from PySide6.QtGui import QDropEvent
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile(str(path)) for path in paths])
    event = QDropEvent(QPointF(1, 1), Qt.DropAction.CopyAction, mime,
                       Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
    window.eventFilter(window.composer, event)


def test_picker_path_accepts_xls(tmp_path):
    window, store, project = build_window(tmp_path)
    try:
        source = tmp_path / 'A8T-BS202607.xls'
        make_xls(source)
        window.import_files([source])
        names = [f['name'] for f in store.files(project)]
        assert 'A8T-BS202607.xls' in names
    finally:
        window.close()


def test_drag_drop_accepts_xls(tmp_path):
    window, store, project = build_window(tmp_path)
    try:
        source = tmp_path / 'A8T-PL202607.xls'
        make_xls(source)
        drop_urls(window, [source])
        assert 'A8T-PL202607.xls' in [f['name'] for f in store.files(project)]
    finally:
        window.close()


def test_mixed_xls_xlsx_batch_loses_nothing(tmp_path):
    window, store, project = build_window(tmp_path)
    try:
        xls = tmp_path / 'report.xls'
        xlsx = tmp_path / 'report.xlsx'
        make_xls(xls)
        make_xlsx(xlsx)
        window.import_files([xls, xlsx])
        names = sorted(f['name'] for f in store.files(project))
        assert names == ['report.xls', 'report.xlsx']
    finally:
        window.close()


def test_detail_workbook_contract_declares_xls_and_xlsx():
    from asset_based_agent.technical_platform.skill_contracts import builtin_contracts
    contract = builtin_contracts()['valuation-detail-workbook-fill']
    assert '.xls' in contract.source_extensions
    assert '.xlsx' in contract.source_extensions


def test_unsupported_extension_still_rejected_without_blocking_others(tmp_path):
    window, store, project = build_window(tmp_path)
    try:
        bad = tmp_path / 'notes.txt'
        bad.write_text('plain text')
        good = tmp_path / 'report.xlsx'
        make_xlsx(good)
        window.import_files([bad, good])
        names = [f['name'] for f in store.files(project)]
        assert names == ['report.xlsx']
        assert '不支持的文件类型' in window.status.text()
    finally:
        window.close()


def test_picker_and_drag_paths_behave_identically(tmp_path):
    window, store, project = build_window(tmp_path)
    try:
        source = tmp_path / 'same.xls'
        make_xls(source)
        window.import_files([source])
        first = store.files(project)
        assert len(first) == 1
        drop_urls(window, [source])  # same name+hash via drag: dedupe, no double registration
        second = store.files(project)
        assert [f['id'] for f in second] == [f['id'] for f in first]
        # a changed copy through the picker registers as a new version, same as drag would
        source.write_bytes(source.read_bytes() + b' ')
        window.import_files([source])
        assert len(store.files(project)) == 2
    finally:
        window.close()


def test_registered_xls_sha256_matches_source(tmp_path):
    window, store, project = build_window(tmp_path)
    try:
        source = tmp_path / 'A8T-BS202607.xls'
        make_xls(source)
        before = digest(source)
        window.import_files([source])
        (entry,) = store.files(project)
        assert entry['sha256'] == before
        assert digest(source) == before  # original untouched
    finally:
        window.close()


# --- K02/K03：.xls 临时转换与表头元数据读取 ------------------------------------

def test_convert_xls_to_xlsx_preserves_visible_values(tmp_path):
    from openpyxl import load_workbook

    from asset_based_agent.technical_platform.xls_support import convert_xls_to_xlsx
    source = tmp_path / 'A8T-BS202607.xls'
    make_xls(source)
    target, entry = convert_xls_to_xlsx(source, tmp_path / 'out' / 'balance_sheet.converted.xlsx')
    book = load_workbook(target)
    assert '资产负债表' in book.sheetnames
    sheet = book['资产负债表']
    assert sheet['A1'].value == '资产负债表'
    assert '测试公司' in str(sheet['A2'].value)
    book.close()
    assert entry['tool'].startswith('xlrd')
    assert entry['source_sha256'] == digest(source)
    assert entry['output_sha256'] == digest(target)
    assert digest(source) == entry['source_sha256']  # 原件未动


def test_convert_xls_skips_hidden_sheets_and_reports(tmp_path):
    import xlwt

    from asset_based_agent.technical_platform.xls_support import convert_xls_to_xlsx
    book = xlwt.Workbook()
    book.add_sheet('资产负债表')
    hidden = book.add_sheet('隐藏底稿')
    hidden.visibility = 1
    very = book.add_sheet('深度隐藏')
    very.visibility = 2
    source = tmp_path / 'hidden.xls'
    book.save(str(source))
    target, entry = convert_xls_to_xlsx(source, tmp_path / 'converted.xlsx')
    from openpyxl import load_workbook
    converted = load_workbook(target)
    assert converted.sheetnames == ['资产负债表']
    converted.close()
    assert sorted(entry['skipped_hidden_sheets']) == ['深度隐藏', '隐藏底稿']


def test_convert_xls_failure_reports_real_reason(tmp_path):
    from asset_based_agent.technical_platform.xls_support import convert_xls_to_xlsx
    broken = tmp_path / 'broken.xls'
    broken.write_bytes(b'not an ole2 workbook')
    with pytest.raises(ValueError, match='无法读取|不是有效的|xlrd'):
        convert_xls_to_xlsx(broken, tmp_path / 'out.xlsx')


def make_statement_xls(path):
    import xlwt
    book = xlwt.Workbook()
    sheet = book.add_sheet('资产负债表')
    sheet.write(0, 0, '资产负债表')
    sheet.write(1, 0, '编制单位：测试公司')
    sheet.write(2, 0, '2026年7月31日')
    book.save(str(path))


def test_statement_metadata_reads_xls_headers(tmp_path):
    from asset_based_agent.technical_platform.material_analysis import (
        statement_metadata,
    )
    source = tmp_path / 'A8T-BS202607.xls'
    make_statement_xls(source)
    entity, period, evidence = statement_metadata(source)
    assert entity == '测试公司'
    assert str(period) == '2026-07-31'
    assert any('编制单位' in item for item in evidence)
