"""Local read-only material pre-identification for task understanding.

Produces a bounded, structured summary per source file so the understanding
request can carry real evidence (document type / entity / period) instead of
bare filenames. Never reads hidden sheets, never emits local paths, raw
binary, or formulas; derived values are always flagged in ``warnings``.
"""
import calendar
import json
import re
from pathlib import Path

from . import xls_support

SUPPORTED_EXTENSIONS = {'.xls', '.xlsx', '.xlsm'}

_DOC_TYPE_RULES = (
    ('cash_flow_statement', ('现金流量表',)),
    ('balance_sheet', ('资产负债表',)),
    ('income_statement', ('利润表',)),
    ('trial_balance', ('科目汇总试算表', '科目余额表', '试算平衡表')),
    ('journal', ('总帐明细帐', '总账明细账', '明细帐追溯', '明细账追溯', '序时账')),
)

_EVIDENCE_KEYWORDS = (
    '资产负债表', '利润表', '现金流量表', '科目汇总试算表', '总帐明细帐',
    '期末余额', '期初余额', '本期金额', '科目编码', '期间',
)

_ENTITY_PATTERNS = (
    re.compile(r'公司\s*[=:：]\s*([A-Za-z0-9]+)'),
    re.compile(r'编制单位\s*[:：]\s*(\S+)'),
    re.compile(r'单位名称\s*[:：]\s*(\S+)'),
    re.compile(r'(?<![\u4e00-\u9fff])名称\s*[:：]\s*(\S+)'),
)

_MARKER_PERIOD = (
    re.compile(r'(?:本期|期间)\s*[:：]?\s*(20\d{2})\s*[-年/.]\s*(\d{1,2})(?:\s*[-月/.]\s*(\d{1,2}))?'),
    re.compile(r'(20\d{2})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日'),
)
_GENERIC_PERIOD = re.compile(r'(20\d{2})\s*[-年/.]\s*(\d{1,2})(?!\s*[-月/.]\s*\d)')

_FIRST_SHEET_ROWS = 200
_OTHER_SHEET_ROWS = 15
_MAX_SHEETS = 3
_MAX_COLS = 50
_TITLE_ROWS = 8


def _month_end(year, month):
    return f'{year:04d}-{month:02d}-{calendar.monthrange(year, month)[1]:02d}'


def _xls_texts(path, hidden_counter):
    import xlrd
    book = xlrd.open_workbook(str(path))
    texts = []
    sheets = []
    for sheet in book.sheets():
        if sheet.visibility != 0:
            hidden_counter.append(sheet.name)
            continue
        sheets.append(sheet.name)
        if len(sheets) > _MAX_SHEETS:
            continue
        rows = _FIRST_SHEET_ROWS if len(sheets) == 1 else _OTHER_SHEET_ROWS
        for row in range(min(sheet.nrows, rows)):
            for col in range(min(sheet.ncols, _MAX_COLS)):
                value = xls_support._cell_value(book, sheet, row, col)
                if value is None:
                    continue
                texts.append(str(value).strip())
    return sheets, texts


def _xlsx_texts(path, hidden_counter):
    from openpyxl import load_workbook
    book = load_workbook(str(path), read_only=True, data_only=True)
    try:
        texts = []
        sheets = []
        for sheet in book.worksheets:
            if sheet.sheet_state != 'visible':
                hidden_counter.append(sheet.title)
                continue
            sheets.append(sheet.title)
            if len(sheets) > _MAX_SHEETS:
                continue
            rows = _FIRST_SHEET_ROWS if len(sheets) == 1 else _OTHER_SHEET_ROWS
            for row in sheet.iter_rows(min_row=1, max_row=rows, max_col=_MAX_COLS,
                                       values_only=True):
                for value in row:
                    if value is None:
                        continue
                    texts.append(str(value).strip())
        return sheets, texts
    finally:
        book.close()


def _detect_document_type(title_text):
    for doc_type, keywords in _DOC_TYPE_RULES:
        if any(keyword in title_text for keyword in keywords):
            return doc_type
    return 'other'


def _detect_entity(texts):
    for text in texts:
        for pattern in _ENTITY_PATTERNS:
            match = pattern.search(text)
            if match:
                return match.group(1).strip()
    return None


def _detect_period(texts, doc_type, warnings):
    for text in texts:
        for pattern in _MARKER_PERIOD:
            match = pattern.search(text)
            if match:
                year, month, day = int(match.group(1)), int(match.group(2)), match.group(3)
                if not 1 <= month <= 12:
                    continue
                if day:
                    return None, f'{year:04d}-{month:02d}-{int(day):02d}'
                end = _month_end(year, month)
                warnings.append(
                    f'期间仅识别到年月 {year:04d}-{month:02d}，period_end 按月末规则推导为 {end}')
                return None, end
    if doc_type in ('journal', 'trial_balance'):
        candidates = set()
        for text in texts:
            for match in _GENERIC_PERIOD.finditer(text):
                year, month = int(match.group(1)), int(match.group(2))
                if 1 <= month <= 12:
                    candidates.add((year, month))
        if candidates:
            year, month = max(candidates)
            end = _month_end(year, month)
            warnings.append(
                f'期间仅识别到年月 {year:04d}-{month:02d}，period_end 按月末规则推导为 {end}')
            return None, end
    return None, None


def summarize_file(path, artifact_id, name=None):
    """Build a bounded structured summary for one source file."""
    path = Path(path)
    warnings = []
    summary = {
        'artifact_id': artifact_id,
        'name': name or path.name,
        'format': path.suffix.lower().lstrip('.'),
        'readable': False,
        'document_type': 'other',
        'entity_name': None,
        'period_start': None,
        'period_end': None,
        'sheet_names': [],
        'header_evidence': [],
        'confidence': 0.1,
        'warnings': warnings,
    }
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        warnings.append(f'不支持的扩展名 {path.suffix.lower() or "(无)"}，无法预识别')
        return summary
    hidden = []
    try:
        if path.suffix.lower() == '.xls':
            sheets, texts = _xls_texts(path, hidden)
        else:
            sheets, texts = _xlsx_texts(path, hidden)
    except Exception:
        warnings.append('无法读取工作簿（文件损坏或格式不受支持）')
        return summary
    summary['readable'] = True
    summary['sheet_names'] = sheets[:_MAX_SHEETS]
    if hidden:
        warnings.append(f'已跳过 {len(hidden)} 个隐藏工作表（不进入摘要）')
    title_text = '\n'.join(texts[:_TITLE_ROWS * _MAX_COLS])
    doc_type = _detect_document_type(title_text)
    summary['document_type'] = doc_type
    summary['header_evidence'] = [kw for kw in _EVIDENCE_KEYWORDS
                                  if any(kw in text for text in texts)]
    entity = _detect_entity(texts)
    summary['entity_name'] = entity
    if entity is None:
        warnings.append('未能从表内识别主体')
    period_start, period_end = _detect_period(texts, doc_type, warnings)
    summary['period_start'] = period_start
    summary['period_end'] = period_end
    if period_end is None:
        warnings.append('未能从表内识别期间')
    confidence = 0.2
    if doc_type != 'other':
        confidence += 0.4
    if entity:
        confidence += 0.2
    if period_end:
        confidence += 0.2
    summary['confidence'] = round(min(confidence, 0.98), 2)
    return summary


def summarize_to_json(path, artifact_id, name=None):
    return json.dumps(summarize_file(path, artifact_id, name), ensure_ascii=False)
