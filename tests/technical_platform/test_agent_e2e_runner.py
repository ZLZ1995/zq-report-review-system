import importlib.util
import json
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]


def load_script():
    path = ROOT / 'scripts' / 'run_agent_e2e.py'
    spec = importlib.util.spec_from_file_location('run_agent_e2e', path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_local_e2e_covers_understanding_confirmation_execution_and_artifact(tmp_path):
    module = load_script()
    source = tmp_path / 'synthetic.xlsx'
    book = Workbook()
    book.active['A1'] = 'visible synthetic evidence'
    hidden = book.create_sheet('hidden')
    hidden.sheet_state = 'hidden'
    hidden['A1'] = 'must-not-be-counted'
    book.save(source)
    original = source.read_bytes()
    output = tmp_path / 'evidence.json'

    evidence = module.run_e2e(
        workspace=tmp_path / 'workspace',
        source=source,
        message='只预检本轮附件，不调用模型，也不要修改原文件。',
        output=output,
        confirmed=True,
    )

    assert source.read_bytes() == original
    assert evidence['status'] == 'passed'
    assert evidence['stages'] == {
        'message': 'passed',
        'understanding': 'passed',
        'confirmation': 'passed',
        'execution': 'passed',
        'verification': 'passed',
        'artifact': 'passed',
    }
    assert evidence['result']['kind'] == 'preflight'
    assert evidence['result']['model_called'] is False
    assert evidence['result']['file_count'] == 1
    assert evidence['source']['name'] == source.name
    assert 'path' not in json.dumps(evidence, ensure_ascii=False).lower()
    assert json.loads(output.read_text(encoding='utf-8')) == evidence


def test_local_e2e_requires_confirmation_before_execution(tmp_path):
    module = load_script()
    source = tmp_path / 'synthetic.xlsx'
    Workbook().save(source)

    try:
        module.run_e2e(
            workspace=tmp_path / 'workspace',
            source=source,
            message='预检这个文件',
            output=tmp_path / 'evidence.json',
            confirmed=False,
        )
    except PermissionError as exc:
        assert '确认' in str(exc)
    else:
        raise AssertionError('unconfirmed e2e run executed')
    assert not (tmp_path / 'evidence.json').exists()


def test_local_e2e_never_overwrites_evidence(tmp_path):
    module = load_script()
    source = tmp_path / 'synthetic.xlsx'
    Workbook().save(source)
    output = tmp_path / 'evidence.json'
    output.write_text('{}', encoding='utf-8')

    try:
        module.run_e2e(
            workspace=tmp_path / 'workspace',
            source=source,
            message='预检这个文件',
            output=output,
            confirmed=True,
        )
    except FileExistsError:
        pass
    else:
        raise AssertionError('existing evidence was overwritten')
    assert output.read_text(encoding='utf-8') == '{}'
