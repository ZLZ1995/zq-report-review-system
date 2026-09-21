import threading

from openpyxl import Workbook

from asset_based_agent.technical_platform.skills import digest, preflight
from asset_based_agent.technical_platform.store import PlatformStore


def test_preflight_hidden_content_excluded_and_original_unchanged(tmp_path):
    path = tmp_path / "detail.xlsx"
    book = Workbook()
    book.active.title = "Visible"
    book.active["A1"] = "safe"
    book.create_sheet("Secret").sheet_state = "veryHidden"
    book["Secret"]["A1"] = "NEVER_UPLOAD_SECRET"
    book.active["A2"] = "=Secret!A1"
    book.save(path)
    original = digest(path)
    store = PlatformStore(tmp_path / "p.sqlite", "alice")
    project = store.create_project("A")
    session = store.create_session(project)
    store.add_file(project, path, original)
    run = store.start_run(session, {"files": store.files(project)})
    result = preflight(store, run, threading.Event(), lambda _: None)
    assert result["files"][0]["chunks"] == 1
    assert result["model_called"] is False
    assert digest(path) == original
    assert store.run(run)["state"] == "succeeded"
    assert "NEVER_UPLOAD_SECRET" not in store.run(run)["result"]


def test_cancel_before_read_does_not_extract(tmp_path):
    store = PlatformStore(tmp_path / "p.sqlite", "alice")
    session = store.create_session(store.create_project("A"))
    run = store.start_run(session, {"files": [{"path": "not-present"}]})
    cancel = threading.Event()
    cancel.set()
    assert preflight(store, run, cancel, lambda _: None) == {"kind": "cancelled"}
    assert store.run(run)["state"] == "cancelled"


def test_remote_adapter_receives_only_filtered_chunks(tmp_path):
    path = tmp_path / "detail.xlsx"
    book = Workbook()
    book.active.title = "Visible"
    book.active["A1"] = "visible evidence"
    book.create_sheet("Hidden").sheet_state = "hidden"
    book["Hidden"]["A1"] = "DO_NOT_SEND"
    book.active["A2"] = "=Hidden!A1"
    book.save(path)
    store = PlatformStore(tmp_path / "p.sqlite", "alice")
    project = store.create_project("A")
    session = store.create_session(project)
    store.add_file(project, path, digest(path))
    run = store.start_run(session, {"files": store.files(project)})

    class Provider:
        def set_client_job_id(self, key):
            assert key == f"PLATFORM-{run}"

        def review_batches(self, batches, progress_callback=None):
            assert progress_callback is not None
            progress_callback(
                {
                    "state": "waiting",
                    "completed_batches": 1,
                    "batch_total": 3,
                    "elapsed_seconds": 20,
                }
            )
            assert len(batches) == 1
            text = " ".join(chunk.text for batch in batches for chunk in batch.chunks)
            assert "visible evidence" in text
            assert "DO_NOT_SEND" not in text
            assert "Hidden!" not in text
            return []

    events = []
    result = preflight(
        store, run, threading.Event(), events.append, provider=Provider()
    )
    assert any("1/3" in event and "20" in event for event in events)
    assert result["kind"] == "review"
    assert result["model_called"] is True
    assert store.run(run)["state"] == "succeeded"
