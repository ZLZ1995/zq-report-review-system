"""Reopen sources and output, report differences, and release only after review."""

import hashlib
import json
import os
import shutil
import tempfile
import zipfile
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


class ReviewBlocked(RuntimeError):
    pass


def dump(path, data):
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding='utf-8')


def issue(reason, **details):
    actions = {
        'source_mismatch': '核对列映射、借贷方向、单位和原始行；有唯一证据后修正明细并重新核对。',
        'source_unverified': '补充原始来源文件、页签及行号，不能以生成记录代替原始证据。',
        'journal_unmatched': '补充同客商、同科目的原始分录；不得串用其他客商摘要。',
        'source_changed': '来源在本轮执行中发生变化，需要使用更新后的来源重新生成并核验。',
        'uncalculated': '使用可靠计算引擎重新计算并读取结果；未计算不得按零或通过处理。',
    }
    return {'reason': reason, 'status': 'unresolved', 'severity': 'blocking',
            'resolution': actions.get(reason, '定位原因并重新生成、复核；不能强行补差或忽略问题。'), **details}


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def equal(a, b, kind='text'):
    if kind == 'amount':
        return numeric(a) and numeric(b) and abs(a - b) < 0.005
    if kind == 'date':
        def day(v):
            if isinstance(v, (date, datetime)):
                return v.strftime('%Y/%m/%d')
            return str(v or '').replace('-', '/')
        return day(a) == day(b)
    return str(a or '').strip() == str(b or '').strip()


def compare(check, expected, actual):
    if equal(actual, expected, check.get('kind', 'text')):
        return None
    return issue('source_mismatch', **check, expected=expected, actual=actual,
                 difference=round(actual - expected, 2) if numeric(actual) and numeric(expected) else None)


def review_source_cells(workbook, checks):
    books = {}
    output = load_workbook(workbook, data_only=True)
    problems = []
    try:
        for check in checks:
            try:
                path = check['source']
                if path not in books:
                    books[path] = load_workbook(path, data_only=True)
                expected = books[path][check['source_sheet']][check['source_cell']].value
                if expected is None:
                    problems.append(issue('source_unverified', **check, detail='来源单元格为空或公式未计算'))
                    continue
                actual = output[check['sheet']][check['cell']].value
                problem = compare(check, expected, actual)
                if problem:
                    problems.append(problem)
            except (KeyError, OSError, ValueError) as exc:
                problems.append(issue('source_unverified', **check, error=str(exc)))
    finally:
        output.close()
        for book in books.values():
            book.close()
    return {'status': 'fail' if problems or not checks else 'pass', 'checked_count': len(checks), 'issues': problems}


def source_fingerprints(paths):
    result = {}
    for path in paths:
        if path:
            p = Path(path)
            with p.open('rb') as stream:
                result[str(p)] = hashlib.file_digest(stream, 'sha256').hexdigest()
    return result


def strict_journal_candidates(sheet, party, journals, pipeline):
    roots = {'应收账款': '1122', '预付账款': '1123', '其他应收款': '1221',
             '应付账款': '2202', '预收账款': '2203', '其他应付款': '2241'}
    direction = 'credit' if sheet in {'应付账款', '预收账款', '其他应付款'} else 'debit'
    accepted = []
    for row in journals:
        parties = {pipeline.normalize_counterparty_display(pipeline.extract_entity_from_text(pipeline.clean(row.get(k))))
                   for k in ('vendor_name', 'counterparty_desc') if row.get(k)}
        name = pipeline.clean(row.get('account_name')).replace('预收款项', '预收账款')
        same_name = name == sheet or name.startswith(sheet + '_') or name.startswith(sheet + '-')
        code = pipeline.clean(row.get('tb_code'))
        same_code = code.startswith(roots[sheet])
        if party not in parties or not (same_name or same_code):
            continue
        if (name and any(name.startswith(other) for other in roots if other != sheet)):
            continue
        if not numeric(row.get(direction)) or row[direction] < .005 or not row.get('gl_date'):
            continue
        accepted.append(row)
    return accepted


def review_pipeline_sources(workbook, args, pipeline, written_rows, initial_hashes):
    """Fresh reads, never expected totals taken solely from writer bookkeeping."""
    output = load_workbook(workbook, data_only=True)
    formulas = load_workbook(workbook, data_only=False)
    problems, checks, notices = [], [], []
    explicit = json.loads(Path(args.source_checks).read_text(encoding='utf-8')) if getattr(args, 'source_checks', None) else []
    verified_cells = set()
    if explicit:
        explicit_report = review_source_cells(workbook, explicit)
        problems.extend(explicit_report['issues'])
        failed = {(item.get('sheet'), item.get('cell')) for item in explicit_report['issues']}
        verified_cells = {(item['sheet'], item['cell']) for item in explicit} - failed
        checks.extend(explicit)
    def check(sheet, cell, expected, kind, source):
        entry = {'sheet': sheet, 'cell': cell, 'kind': kind, **source}
        actual = output[sheet][cell].value if sheet in output.sheetnames else None
        checks.append({**entry, 'expected': expected, 'actual': actual})
        problem = compare(entry, expected, actual)
        if problem:
            problems.append(problem)
    try:
        current = source_fingerprints(initial_hashes)
        for path, digest in initial_hashes.items():
            if current.get(path) != digest:
                problems.append(issue('source_changed', source=path))
        bs_path = args.financial_statement or [args.balance_sheet]
        metadata = pipeline.select_latest_statement(bs_path)
        # Prefer the original financial statement over an intermediate normalization.
        bs = pipeline.parse_balance_sheet(Path(metadata['source']))
        journals = pipeline.load_journal_rows(Path(args.journal) if args.journal else None)
        cp = pipeline.load_counterparty_balance_rows(Path(args.counterparty_balance) if args.counterparty_balance else None)
        tb = pipeline.load_trial_balance_rows(Path(args.trial_balance))
        source = {'source': metadata['source'], 'source_sheet': metadata['sheet']}
        covered_labels = set()
        if '资产负债表' not in output.sheetnames:
            problems.append(issue('source_unverified', sheet='资产负债表'))
        else:
            ws = output['资产负债表']
            for row in range(6, min(ws.max_row, 45) + 1):
                for label_cols, current_col, prior_col in [('AB', 'D', 'C'), ('FG', 'I', 'H')]:
                    label = next((pipeline.clean(ws[f'{c}{row}'].value).replace(' ', '')
                                  for c in label_cols if pipeline.clean(ws[f'{c}{row}'].value).replace(' ', '') in bs['values']), None)
                    if label:
                        covered_labels.add(label)
                        check(ws.title, f'{current_col}{row}', bs['values'][label], 'amount', {**source, 'source_label': label})
                        if label in bs.get('values_prior', {}):
                            check(ws.title, f'{prior_col}{row}', bs['values_prior'][label], 'amount', {**source, 'source_label': label, 'period': 'prior'})
            if not covered_labels:
                problems.append(issue('source_unverified', sheet=ws.title, detail='未识别资产负债表标签映射'))

        six = {'应收账款','预付账款','其他应收款','应付账款','预收账款','其他应付款'}
        payable = {'应付账款','预收账款','其他应付款'}
        expected, actual = defaultdict(float), defaultdict(float)
        evidence = defaultdict(list)
        if cp:
            for item in cp:
                name = str(item.get('account_name', '')).replace('预收款项', '预收账款')
                sheet = next((s for s in six if name == s or name.startswith(s + '_') or name.startswith(s + '-')), None)
                if not sheet and abs(item['credit'] - item['debit']) >= .005:
                    problems.append(issue('source_unverified', source=args.counterparty_balance,
                        source_sheet=item.get('source_sheet'), source_row=item.get('source_row'),
                        account_name=name, detail='非零往来来源行未能明确路由，不能忽略。'))
                if sheet:
                    amount = item['credit'] - item['debit'] if sheet in payable else item['debit'] - item['credit']
                    if abs(amount) >= .005:
                        key = (sheet, pipeline.normalize_counterparty_display(item['counterparty']))
                        expected[key] += amount
                        evidence[key].append({'source': args.counterparty_balance, 'source_sheet': item.get('source_sheet'), 'source_row': item.get('source_row'), 'tb_code': item['tb_code']})
        else:
            mapping = json.loads(Path(args.project_mapping).read_text(encoding='utf-8'))
            routes = {str(x['tb_code']): x['target_sheet'] for x in mapping.get('account_mappings', [])}
            codes = {item['tb_code'] for item in tb}
            for item in tb:
                code = item['tb_code']
                sheet = routes.get(code)
                if sheet in six and not any(c.startswith(code + '.') for c in codes):
                    key = (sheet, pipeline.normalize_counterparty_display(item['aux_name']))
                    expected[key] += item['credit_end'] - item['debit_end'] if sheet in payable else item['debit_end'] - item['credit_end']
                    evidence[key].append({'source': args.trial_balance, 'source_row': item.get('source_row'), 'tb_code': code})

        for sheet in six:
            if sheet not in output.sheetnames:
                continue
            ws = output[sheet]
            last = pipeline.find_effective_total_row(formulas[sheet], 6)
            book_col = 'G' if sheet in payable else ('H' if sheet == '预付账款' else 'P')
            for row in range(6, last):
                amount = ws[f'{book_col}{row}'].value
                party = pipeline.normalize_counterparty_display(ws[f'B{row}'].value)
                if numeric(amount) and abs(amount) >= .005:
                    actual[(sheet, party)] += amount
                if not party or not numeric(amount) or abs(amount) < .005:
                    continue
                candidates = strict_journal_candidates(sheet, party, journals, pipeline)
                selected = pipeline.select_journal_entry(party, amount, 'payable' if sheet in payable else 'receivable', candidates)
                date_col, desc_col = ('C', 'D') if sheet in payable else ('D', 'C')
                if selected:
                    ref = {'source': args.journal, 'source_row': selected.get('row'), 'tb_code': selected.get('tb_code'), 'counterparty': party}
                    check(sheet, f'{date_col}{row}', selected.get('gl_date'), 'date', ref)
                    check(sheet, f'{desc_col}{row}', pipeline.normalize_business_desc(pipeline.clean(selected.get('line_desc') or selected.get('summary'))), 'text', ref)
                elif ws[f'{date_col}{row}'].value or ws[f'{desc_col}{row}'].value:
                    problems.append(issue('journal_unmatched', sheet=sheet, cell=f'{date_col}{row}', counterparty=party))
                else:
                    notices.append({'sheet': sheet, 'row': row, 'counterparty': party, 'reason': '序时账未检出匹配分录；业务内容和日期留空', 'status': 'needs_materials'})
        for key in expected.keys() | actual.keys():
            exp, got = round(expected.get(key, 0), 2), round(actual.get(key, 0), 2)
            entry = {'sheet': key[0], 'counterparty': key[1], 'source_evidence': evidence.get(key, []), 'expected': exp, 'actual': got}
            checks.append(entry)
            if abs(exp - got) >= .005:
                problems.append(issue('source_mismatch', **entry, difference=round(got - exp, 2)))

        for item in written_rows:
            sheet, row = item['_sheet'], item['_written_row']
            for col, field in pipeline.DETAIL_WRITE_TEMPLATES.get(sheet, []):
                if field not in {'book_value', 'counterparty', 'sub_name', 'date_value'} or col == 'age_bucket_col':
                    continue
                entry = {'sheet': sheet, 'cell': f'{col}{row}', 'kind': 'amount' if field == 'book_value' else ('date' if field == 'date_value' else 'text'), 'verification_layer': 'saved_file_vs_write_record'}
                problem = None if (sheet, f'{col}{row}') in verified_cells else compare(entry, item.get(field), output[sheet][f'{col}{row}'].value)
                if problem:
                    problems.append(problem)
            if sheet not in six:
                for col, field in pipeline.DETAIL_WRITE_TEMPLATES.get(sheet, []):
                    if field not in {'book_value', 'counterparty', 'sub_name', 'date_value'} or item.get(field) in (None, ''):
                        continue
                    if (sheet, f'{col}{row}') not in verified_cells:
                        problems.append(issue('source_unverified', sheet=sheet, cell=f'{col}{row}', field=field,
                            detail='该字段没有原始资料证据映射；需要补充 source_checks，不能仅以合计相等通过。'))
        if not checks:
            problems.append(issue('source_unverified', detail='独立核对无有效检查项'))
    finally:
        output.close()
        formulas.close()
    return {'status': 'fail' if problems else 'pass', 'checked_count': len(checks),
            'issues': problems, 'checks': checks, 'notices': notices,
            'source_hashes': initial_hashes, 'automatic_repair': 'none_without_unique_source_evidence'}


def repair_unique_source_cells(workbook, checks, registry):
    """One bounded repair pass for explicit, unambiguous, writable input cells."""
    from collections import Counter
    from summary_sheet_policy import is_linked_summary_sheet
    original = review_source_cells(workbook, checks) if checks else {'issues': []}
    counts = Counter((x['sheet'], x['cell']) for x in checks)
    wb = load_workbook(workbook, data_only=False)
    repaired = []
    try:
        for item in original['issues']:
            sheet, address = item.get('sheet'), item.get('cell')
            if item['reason'] != 'source_mismatch' or counts[(sheet, address)] != 1 or is_linked_summary_sheet(sheet):
                continue
            meta = registry.get('selected_sheets', {}).get(sheet, {})
            if address not in meta.get('confirmed_input_cells', []) or address in meta.get('forbidden_non_formula_cells', []):
                continue
            cell = wb[sheet][address]
            if cell.data_type == 'f':
                continue
            cell.value = item['expected']
            repaired.append({**item, 'status': 'repaired_pending_recheck', 'before': item['actual'], 'after': item['expected']})
        if repaired:
            wb.save(workbook)
    finally:
        wb.close()
    final = review_source_cells(workbook, checks) if checks else {'status': 'pass', 'issues': []}
    return {'attempt_limit': 1, 'repairs': repaired, 'recheck': final}


def write_feedback(folder, issues, notices=(), status='blocked'):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    numbered = [{'id': f'D{i:04d}', **item} for i, item in enumerate(issues, 1)]
    dump(folder / 'difference_resolution_report.json', {'status': status, 'issues': numbered, 'notices': list(notices)})
    dump(folder / 'completion_status.json', {'status': status, 'unresolved_count': len(numbered), 'feedback_required': True, 'feedback': str(folder / 'user_feedback.md')})
    lines = ['# 生成结果复核反馈', '', '状态：' + ('未通过验收，未发布本轮成果。' if status != 'complete' else '验收通过。'), '']
    for item in numbered:
        lines += [f"- {item['id']}：{item.get('sheet', '')} {item.get('cell', '')} {item.get('counterparty', '')}；{item.get('reason', '')}",
                  f"  原始值：{item.get('expected', '待核实')}；生成值：{item.get('actual', '待核实')}；差异：{item.get('difference', '待核实')}",
                  f"  来源：{item.get('source', item.get('source_evidence', '详见来源记录'))}；处理要求：{item.get('resolution', '核实来源并重新复核。')}"]
    if notices:
        lines += ['', '## 待补资料或证据边界', ''] + ['- ' + json.dumps(n, ensure_ascii=False, default=str) for n in notices]
    (folder / 'user_feedback.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    return folder / 'user_feedback.md'


def publish_after_review(staging, target, folder, reports, notices=()):
    problems = []
    if not reports:
        problems.append(issue('review_reports_missing'))
    for report in reports:
        problems.extend(report.get('issues', []))
        if report.get('status') not in {'pass', 'ok'} and not report.get('issues'):
            problems.append(issue('review_failed', report=report))
    if problems:
        feedback = write_feedback(folder, problems, notices)
        print('未通过验收，未发布本轮成果。必须向使用者说明未解决差异及处理要求：' + str(feedback))
        raise ReviewBlocked(str(feedback))
    staging, target = Path(staging), Path(target)
    temporary = None
    try:
        with zipfile.ZipFile(staging) as archive:
            if archive.testzip() is not None:
                raise ValueError('invalid_workbook_zip')
        target.parent.mkdir(parents=True, exist_ok=True)
        fd, name = tempfile.mkstemp(prefix='validated_', suffix='.xlsx', dir=target.parent)
        os.close(fd)
        temporary = Path(name)
        shutil.copy2(staging, temporary)
        os.replace(temporary, target)
    except Exception as exc:
        write_feedback(folder, [issue('publish_failed', error=str(exc))], notices)
        raise
    finally:
        if temporary and temporary.exists():
            temporary.unlink()
    feedback = write_feedback(folder, [], notices, status='complete')
    print('复核反馈（包含待补资料）：' + str(feedback))
