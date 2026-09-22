"""K08: pipeline must parse single-sided A8T-export balance sheets.

ERP exports put labels in column A with 年初余额/期末余额 in columns B/C
under a late header row. The parser must detect this layout instead of
misreading it as a two-sided statement.
"""
import importlib.util
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src'))

from openpyxl import Workbook

SCRIPTS = ROOT / '.codex' / 'skills' / 'valuation-detail-workbook-fill' / 'scripts'


def load_pipeline():
    sys.path.insert(0, str(SCRIPTS))
    spec = importlib.util.spec_from_file_location(
        'run_detail_workbook_pipeline', SCRIPTS / 'run_detail_workbook_pipeline.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def make_single_sided(path):
    wb = Workbook()
    ws = wb.active
    ws.title = '销项 1 (A8T)'
    ws['B1'] = 'PRC-资产负债表'
    ws['B3'] = '本期：2026-07'
    ws['A7'] = '公司=A8T (阿里云飞天（北京）云计算有限公司)'
    ws['B9'] = '年初余额'
    ws['C9'] = '期末余额'
    rows = [('流动资产', None, None), ('货币资金', 5360.66, 726.52),
            ('流动资产合计', 5729.64, 11934.09), ('资产总计', 5729.64, 11934.09),
            ('应付账款', 2013205.76, 1301505.97), ('负债合计', 11293533.33, 15587687.03),
            ('实收资本', 0, 0), ('所有者权益合计', -11287803.69, -15575752.94),
            ('负债和所有者权益合计', 5729.64, 11934.09)]
    for index, (label, prior, current) in enumerate(rows, start=10):
        ws.cell(row=index, column=1, value=label)
        if prior is not None:
            ws.cell(row=index, column=2, value=prior)
        if current is not None:
            ws.cell(row=index, column=3, value=current)
    wb.save(path)


class SingleSidedBalanceSheetTest(unittest.TestCase):
    def setUp(self):
        import tempfile
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_single_sided_layout_parses_labels_and_amounts(self):
        pipeline = load_pipeline()
        path = self.dir / 'A8T-BS202607.xlsx'
        make_single_sided(path)
        result = pipeline.parse_balance_sheet(path)
        values = result['values']
        self.assertEqual(726.52, values['货币资金'])
        self.assertEqual(11934.09, values['资产总计'])
        self.assertEqual(11934.09, values['负债和所有者权益合计'])
        self.assertEqual(5360.66, result['values_prior']['货币资金'])
        self.assertEqual(5729.64, result['values_prior']['资产总计'])

    def test_section_headers_are_not_treated_as_amounts(self):
        pipeline = load_pipeline()
        path = self.dir / 'A8T-BS202607.xlsx'
        make_single_sided(path)
        values = pipeline.parse_balance_sheet(path)['values']
        self.assertNotIn('流动资产', values)
        self.assertNotIn('PRC-资产负债表', values)

    def test_accrued_expenses_merge_into_other_current_liabilities(self):
        """A8T exports carry both 预提费用 and an explicit-zero 其他流动负债
        row; the alias fallback alone would zero out the accrued amount and
        unbalance the statement."""
        pipeline = load_pipeline()
        path = self.dir / 'A8T-BS202607.xlsx'
        wb = Workbook()
        ws = wb.active
        ws['B1'] = 'PRC-资产负债表'
        ws['B3'] = '本期：2026-07'
        ws['A7'] = '公司=A8T (阿里云飞天（北京）云计算有限公司)'
        ws['B9'] = '年初余额'
        ws['C9'] = '期末余额'
        rows = [('货币资金', 100.0, 200.0), ('资产总计', 100.0, 200.0),
                ('预提费用', 40.0, 60.0), ('其他流动负债', 0.0, 0.0),
                ('负债合计', 140.0, 260.0), ('所有者权益合计', -40.0, -60.0),
                ('负债和所有者权益合计', 100.0, 200.0)]
        for index, (label, prior, current) in enumerate(rows, start=10):
            ws.cell(row=index, column=1, value=label)
            ws.cell(row=index, column=2, value=prior)
            ws.cell(row=index, column=3, value=current)
        wb.save(path)
        result = pipeline.parse_balance_sheet(path)
        self.assertEqual(60.0, result['values']['其他流动负债'])
        self.assertEqual(40.0, result['values_prior']['其他流动负债'])


if __name__ == '__main__':
    unittest.main()
