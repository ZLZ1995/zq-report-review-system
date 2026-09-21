"""Compound DETAIL steps accept planning-level reference files.

A single generator task whose understanding designates reference-only files
must not die at compile time with 'Generator reference roles are not
supported': DETAIL's automatic material analysis reads every provided file
read-only and classifies it by content, so a reference file is simply an
extra input annotated with user_role='reference' and a visible hint in the
excerpt sent to the model.  Positional generators (HISTORY etc.) keep the
strict contract.
"""
import threading

import pytest
from openpyxl import Workbook

from asset_based_agent.technical_platform.execution import execute_task
from asset_based_agent.technical_platform.permissions import PermissionService
from asset_based_agent.technical_platform.skills import DETAIL, HISTORY, digest
from asset_based_agent.technical_platform.store import PlatformStore


def make_statement(path, entity, year, month, day):
    wb = Workbook()
    balance = wb.active
    balance.title = '资产负债表'
    balance['A1'] = '资产负债表'
    balance['A2'] = f'编制单位：{entity} {year}年{month}月{day}日 单位：元'
    balance['A4'] = '项目'
    balance['B4'] = '期末余额'
    for row, (label, value) in enumerate((('资产总计', 300000), ('负债合计', 100000),
                                          ('所有者权益（或股东权益）合计', 200000)), 5):
        balance.cell(row, 1, label)
        balance.cell(row, 2, value)
    wb.create_sheet('利润表')
    wb.save(path)


def prepared(tmp_path, skill, *, with_reference=True):
    from asset_based_agent.technical_platform.compound_task import (
        build_compound_task_spec,
    )
    store = PlatformStore(tmp_path / 'db.sqlite', 'alice')
    project = store.create_project('synthetic')
    session = store.create_session(project)
    main = tmp_path / 'main.xlsx'
    make_statement(main, '北京甲示例科技有限公司', 2026, 6, 30)
    store.add_file(project, main, digest(main))
    if with_reference:
        ref_path = tmp_path / 'ref.xlsx'
        make_statement(ref_path, '北京甲示例科技有限公司', 2025, 12, 31)
        store.add_file(project, ref_path, digest(ref_path))
    by_name = {f['name']: f for f in store.files(project)}
    target = by_name['main.xlsx']
    reference = by_name.get('ref.xlsx')
    chosen = [target] + ([reference] if reference else [])
    request = {'request_id': 'r', 'model_id': 'm', 'message_id': 'msg',
               'prompt': '生成评估明细表',
               'files': [{k: f[k] for k in ('id', 'name', 'sha256')} for f in chosen],
               'skills': [{'id': skill.id, 'adapter': skill.id, 'name': skill.name,
                           'description': ''}]}
    understanding = {'message_intent': 'execute', 'goal': request['prompt'],
                     'targets': [target['id']],
                     'references': [reference['id']] if reference else [],
                     'excluded': [], 'constraints': [], 'deliverables': ['Excel'],
                     'missing_inputs': [], 'evidence_message_ids': ['msg'],
                     'skill_ids': [skill.id], 'next_action': 'plan', 'reply': '准备生成'}
    inputs = [{'kind': 'file', 'ref': target['id'], 'role': 'target'}]
    if reference:
        inputs.append({'kind': 'file', 'ref': reference['id'], 'role': 'reference'})
    proposal = {'request_id': 'r', 'steps': [
        {'step_id': 'generate', 'skill_id': skill.id, 'goal': '生成明细表',
         'dependencies': [], 'inputs': inputs}]}
    snapshot = build_compound_task_spec(store, session,
                                        {'request': request, 'understanding': understanding},
                                        proposal, chosen, revision=2).to_snapshot()
    return store, session, snapshot, target, reference


def test_detail_compound_accepts_reference_inputs(tmp_path):
    _, _, snapshot, target, reference = prepared(tmp_path, DETAIL)
    step = snapshot['execution_plan']['steps'][0]
    assert step['target_inputs'] == [target['id']]
    assert step['reference_inputs'] == [reference['id']]
    assert snapshot['file_scope']['references'][0]['id'] == reference['id']
    assert snapshot['permissions']['call_model'] is True


def test_history_generator_still_rejects_reference_inputs(tmp_path):
    with pytest.raises(ValueError, match='reference roles'):
        prepared(tmp_path, HISTORY)


def test_compound_detail_step_snapshot_marks_reference_files(tmp_path, monkeypatch):
    from asset_based_agent.technical_platform import generation as generation_module
    from asset_based_agent.technical_platform.generation_paths import (
        generation_work_directory,
    )
    store, session, snapshot, target, reference = prepared(tmp_path, DETAIL)
    run = store.start_run(session, snapshot)
    PermissionService(store).authorize(run, snapshot, confirmed=True)
    captured = {}

    def fake_execute(store_, run_id, step_snapshot, cancel, progress, **kwargs):
        captured['snapshot'] = step_snapshot
        output = generation_work_directory(store_.path.parent, run_id,
                                           kwargs.get('step_id')) / 'output'
        output.mkdir(parents=True, exist_ok=True)
        artifact = output / 'detail_workbook.xlsx'
        Workbook().save(artifact)
        return {'kind': 'generation', 'ok': True, 'artifacts': [
            {'name': 'detail_workbook.xlsx', 'path': str(artifact),
             'sha256': digest(artifact)}]}

    monkeypatch.setattr(generation_module, 'execute_generation', fake_execute)

    class Client:
        access_token = 'token'

    result = execute_task(store, run, threading.Event(), lambda _: None,
                          client=Client())
    assert store.run(run)['state'] == 'succeeded', result
    files = {f['id']: f for f in captured['snapshot']['files']}
    assert files[reference['id']].get('user_role') == 'reference'
    assert 'user_role' not in files[target['id']]


def test_analysis_provider_prepends_reference_hint(tmp_path):
    from threading import Event

    from asset_based_agent.technical_platform.material_analysis import (
        MaterialAnalysisProvider,
    )
    path = tmp_path / 'ref.xlsx'
    make_statement(path, '北京甲示例科技有限公司', 2025, 12, 31)

    class Client:
        payload = None

        def analyze_materials(self, payload):
            Client.payload = payload
            return {'assignments': [
                {'file_id': 'ref', 'role': 'other', 'reason': '仅参考'}]}

    MaterialAnalysisProvider(Client(), 'model', 'rules').analyze(
        [{'id': 'ref', 'name': path.name, 'path': str(path), 'sha256': digest(path),
          'user_role': 'reference'}],
        'run', Event(), lambda _: None)
    text = Client.payload['files'][0]['text']
    assert text.startswith('（用户指定：本文件仅作参考资料，不作为填报依据。）')
