"""Explicitly authorized built-in generators; no external script execution."""

import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import time
from contextlib import nullcontext
from pathlib import Path

from ..report_review_app.services.resource_locks import CLIENT_RESOURCES
from .generation_paths import generation_work_directory
from .material_analysis import resolution_snapshot
from .project_catalog import validate_business_directory
from .skills import (
    DETAIL,
    FINANCIAL_BRIEF,
    HISTORY,
    WORKFLOW_TO_SKILL,
    SourceValidationError,
    digest,
)

# role, user-facing label, accepted extensions, required
INPUT_ROLES = {
    HISTORY.id: [('source_excel', '工商变更信息', ('.xlsx',), True)],
    DETAIL.id: [('trial_balance', '科目余额表', ('.xlsx', '.xls'), True),
                ('balance_sheet', '原始资产负债表（含单位、期间）', ('.xlsx', '.xls'), True),
                ('journal', '序时账（可选）', ('.xlsx', '.xls'), False),
                ('bank_statement', '银行对账单（可选，标准表头 XLSX）', ('.xlsx',), False)],
    FINANCIAL_BRIEF.id: [
        ('period_one', '较早完整年度财务报表', ('.xlsx',), True),
        ('period_two', '较近完整年度财务报表', ('.xlsx',), True),
        ('basis_date', '评估基准日财务报表', ('.xlsx',), True),
    ],
    WORKFLOW_TO_SKILL.id: [
        ('workflow_contract', '已确认的工作流契约 JSON', ('.json',), True),
    ],
}


def bundle_directory(skill_id):
    if skill_id not in INPUT_ROLES:
        raise PermissionError('不支持的本地生成 Skill')
    if getattr(sys, 'frozen', False):
        return Path(sys._MEIPASS) / 'builtin_skills' / skill_id
    return Path(__file__).resolve().parents[3] / '.codex/skills' / skill_id


def bundle_fingerprint(skill_id):
    root = bundle_directory(skill_id)
    if skill_id != WORKFLOW_TO_SKILL.id:
        locked_template(skill_id)
    if not (root / 'SKILL.md').is_file():
        raise ValueError('内置 Skill 文件缺失，请修复安装')
    result = hashlib.sha256()
    for path in sorted(root.rglob('*')):
        if path.is_file() and path.suffix in {'.py', '.md', '.json', '.yaml'}:
            result.update(path.relative_to(root).as_posix().encode('utf-8'))
            result.update(path.read_bytes())
    return result.hexdigest()


def locked_template(skill_id):
    """Only a release-bound, hash-pinned template can authorize generation."""
    root = bundle_directory(skill_id).resolve()
    lock = root / 'template.lock.json'
    if not lock.is_file():
        raise ValueError('尚未绑定经确认的锁定模板，禁止执行；请先由维护者完成模板绑定')
    data = json.loads(lock.read_text(encoding='utf-8'))
    if data.get('schema_version') != 1 or data.get('skill_id') != skill_id:
        raise ValueError('模板绑定信息无效')
    path = (root / data['path']).resolve()
    suffix = '.docx' if skill_id in {HISTORY.id, FINANCIAL_BRIEF.id} else '.xlsx'
    if not path.is_relative_to(root / 'assets') or path.suffix.lower() != suffix:
        raise PermissionError('锁定模板必须位于内置资源目录')
    if not path.is_file() or digest(path) != data.get('sha256'):
        raise ValueError('锁定模板缺失或被修改，禁止执行')
    return path


def validate_roles(skill_id, files, roles):
    definitions = INPUT_ROLES[skill_id]
    if not isinstance(roles, dict) or set(roles) - {r[0] for r in definitions}:
        raise ValueError('输入角色不符合所选 Skill')
    by_id = {f['id']: f for f in files}
    if len(by_id) != len(files) or len(set(roles.values())) != len(roles):
        raise ValueError('同一资料不能重复充当多个输入角色')
    for role, label, suffixes, required in definitions:
        if role not in roles and not required:
            continue
        item = by_id.get(roles.get(role))
        if item is None or Path(item['name']).suffix.lower() not in suffixes:
            raise ValueError(f'请指定正确的{label}')
    if set(roles.values()) != set(by_id):
        raise ValueError('每个选定文件都必须分配角色；请取消勾选本轮无关资料')


def infer_financial_roles(files):
    """Bind three financial workbooks by their balance-sheet header dates."""
    from datetime import date

    from openpyxl import load_workbook  # type: ignore[import-untyped]

    if len(files) != 3 or any(Path(item['name']).suffix.lower() != '.xlsx' for item in files):
        raise ValueError('财务简报需要且仅需要三个期间的 XLSX 财务报表')
    dated = []
    for item in files:
        book = load_workbook(item['path'], read_only=True, data_only=True)
        try:
            if '资产负债表' not in book.sheetnames or '利润表' not in book.sheetnames:
                raise ValueError(f"{item['name']} 缺少资产负债表或利润表")
            sheet = book['资产负债表']
            text = ' '.join(str(cell.value) for row in sheet.iter_rows(max_row=4)
                            for cell in row if cell.value not in (None, ''))
        finally:
            book.close()
        match = re.search(r'(\d{4})年(\d{1,2})月(\d{1,2})日', text)
        if match is None:
            raise ValueError(f"{item['name']} 的资产负债表表头缺少完整日期")
        dated.append((date(*map(int, match.groups())), item['id']))
    dated.sort()
    if len({item[0] for item in dated}) != 3:
        raise ValueError('三份财务报表的期间不能重复')
    return dict(zip(('period_one', 'period_two', 'basis_date'),
                    (item[1] for item in dated)))


def financial_report_date(item):
    """Balance-sheet header date of a financial workbook, or None when not a report."""
    from datetime import date

    from openpyxl import load_workbook  # type: ignore[import-untyped]

    if Path(item['name']).suffix.lower() != '.xlsx':
        return None
    try:
        book = load_workbook(item['path'], read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - unreadable workbooks simply do not qualify
        return None
    try:
        if '资产负债表' not in book.sheetnames or '利润表' not in book.sheetnames:
            return None
        sheet = book['资产负债表']
        text = ' '.join(str(cell.value) for row in sheet.iter_rows(max_row=4)
                        for cell in row if cell.value not in (None, ''))
    finally:
        book.close()
    match = re.search(r'(\d{4})年(\d{1,2})月(\d{1,2})日', text)
    return date(*map(int, match.groups())) if match else None


def select_financial_reports(files):
    """Dialog-free binding: pick the three period reports out of the current selection."""
    dated = [(stamp, item) for item in files
             if (stamp := financial_report_date(item)) is not None]
    if len(dated) < 3:
        raise ValueError('财务简报需要三个期间的 XLSX 财务报表（含资产负债表、利润表及表头日期）；'
                         f'本轮选定资料中只识别到 {len(dated)} 份，请补充或调整勾选后重试。')
    if len(dated) > 3:
        raise ValueError(f'本轮识别到 {len(dated)} 份可用财务报表，无法确定两年一期；'
                         '请只勾选三份报表后重试。')
    if len({stamp for stamp, _item in dated}) != 3:
        raise ValueError('三份财务报表的期间不能重复')
    dated.sort(key=lambda pair: pair[0])
    reports = [item for _stamp, item in dated]
    roles = dict(zip(('period_one', 'period_two', 'basis_date'),
                     (item['id'] for item in reports)))
    return roles, reports


def auto_generation_roles(skill_id, files):
    """Bind generator inputs without a dialog under a standing permission grant."""
    if skill_id == DETAIL.id:
        return None, files
    if skill_id == FINANCIAL_BRIEF.id:
        return select_financial_reports(files)
    labels = {HISTORY.id: ('工商变更信息 Excel', '.xlsx'),
              WORKFLOW_TO_SKILL.id: ('工作流契约 JSON', '.json')}
    if skill_id not in labels:
        raise PermissionError('不支持的本地生成 Skill')
    label, suffix = labels[skill_id]
    matches = [item for item in files if Path(item['name']).suffix.lower() == suffix]
    if len(matches) != 1:
        raise ValueError(f'该生成需要且仅需要一份{label}；本轮匹配到 {len(matches)} 份，'
                         '请调整勾选后重试。')
    return {INPUT_ROLES[skill_id][0][0]: matches[0]['id']}, matches


def run_feedback(work, succeeded):
    """User-facing summary; a missing worker note must not misreport success."""
    path = work / 'output/user_feedback.md'
    if path.is_file():
        return path.read_text(encoding='utf-8')
    return ('生成完成，成果已通过来源与交付校验。' if succeeded
            else '生成未完成，请查看本轮日志；没有发布正式成果。')


def artifact_path(store, session_id, run_id, index):
    run = store.run(run_id)
    if run['session'] != session_id:
        raise PermissionError('成果不属于当前会话')
    result = json.loads(run['result'] or '{}')
    items = result.get('artifacts', [])
    if result.get('kind') != 'generation' or not 0 <= index < len(items):
        raise ValueError('成果链接无效')
    item = items[index]
    path = Path(item['path']).resolve()
    root = (store.path.parent / 'runs' / run_id / 'output').resolve()
    if not path.is_relative_to(root) or path.suffix.lower() not in {'.docx', '.xlsx', '.pdf', '.png', '.md', '.json'}:
        raise PermissionError('成果路径超出任务范围')
    if not path.is_file() or digest(path) != item['sha256']:
        raise ValueError('成果已修改、移动或删除，请在项目目录核对')
    return path


def register_artifacts(skill_id, work, names, succeeded):
    """Stamp each registered artifact with its delivery role; evidence stays internal.

    Staging files are never registered; on failure or cancellation only the
    internal feedback note is listed, and no business file is published.
    """
    from .artifact_contract import detail_display_name, stamp_artifact
    artifacts = []
    for name in ([*names, 'user_feedback.md'] if succeeded else ['user_feedback.md']):
        path = work / 'output' / name
        if not (path.is_file() and path.resolve().is_relative_to(work)):
            continue
        display = (detail_display_name(work) if succeeded and skill_id == DETAIL.id
                   and name == 'detail_workbook.xlsx' else None)
        artifacts.append(stamp_artifact(skill_id, name, path, digest(path), display_name=display))
    return artifacts


def execute_generation(store, run_id, snapshot, cancel, progress, *, provider=None, manage_run=True,
                       step_id=None):
    root = validate_business_directory(store.path.parent)
    work = generation_work_directory(root, run_id, step_id)
    with CLIENT_RESOURCES.lease(('output:' + os.path.normcase(str(work)),), cancel):
        return _execute_generation(store, run_id, snapshot, cancel, progress, provider=provider,
                                   manage_run=manage_run, step_id=step_id)


def _execute_generation(store, run_id, snapshot, cancel, progress, *, provider=None, manage_run=True,
                        step_id=None):
    from ..report_review_app.services.task_cancellation import TaskCancelled
    skill_id = snapshot['skill_id']
    roles, files = snapshot.get('input_roles', {}), snapshot['files']
    automatic = snapshot.get('automatic_materials') is True
    if not automatic:
        validate_roles(skill_id, files, roles)
    if (snapshot.get('mode') != 'local_generation'
            or snapshot.get('skill_instructions') != bundle_fingerprint(skill_id)
            or snapshot.get('capabilities') != ['generate_artifacts', 'read_selected_files']):
        raise PermissionError('生成规则或任务能力已变化，请重新确认任务')
    root = validate_business_directory(store.path.parent)
    resuming = automatic and snapshot.get('material_resolution_override') is not None
    # A clarification resume reuses its own work directory (it holds the
    # persisted analysis); a fresh run must still fail on directory collision.
    work = generation_work_directory(root, run_id, step_id, create=not resuming)
    if resuming and not work.is_dir():
        raise ValueError('澄清恢复所需的已落盘分析目录缺失，请重新提交任务')
    if automatic:
        if provider is None or not snapshot['permissions'].get('call_model'):
            raise PermissionError('资料识别缺少模型授权')
        override = snapshot.get('material_resolution_override')
        if override is not None:
            # Resume after user clarification: reuse the persisted analysis,
            # never re-bill the model for a step that already succeeded.
            from .material_analysis import apply_resolution_override
            plan = json.loads((work / 'material_analysis.json').read_text(encoding='utf-8'))
            stored = json.loads((work / 'material_resolution.json').read_text(encoding='utf-8'))
            resolution = apply_resolution_override(override, files, stored)
            progress('已按用户澄清恢复资料范围；沿用已完成的模型识别结果，不重复调用模型。')
        else:
            request_id = run_id if step_id is None else run_id + '-' + hashlib.sha256(step_id.encode()).hexdigest()
            resolution, plan = provider.analyze(files, request_id, cancel, progress)
            if cancel.is_set():
                raise TaskCancelled('资料识别已取消，未开始生成')
            if any(digest(Path(f['path'])) != f['sha256'] for f in files):
                raise ValueError('识别过程中资料已变化，请重新添加')
            # Persist the raw model plan before disambiguation so failed or
            # pending resolutions keep their evidence on disk.
            (work / 'material_analysis.json').write_text(json.dumps(plan, ensure_ascii=False), encoding='utf-8')
            (work / 'material_resolution.json').write_text(
                json.dumps(resolution_snapshot(resolution), ensure_ascii=False), encoding='utf-8')
        identified = '\n'.join(f"{next(f['name'] for f in files if f['id'] == a['file_id'])}：{a['reason']}"
                               for a in plan['assignments'])
        progress('资料识别完成：' + identified)
        for reason in resolution.reasons:
            progress(reason)
        if resolution.status == 'waiting_user':
            feedback = ('资料自动识别结果：\n' + identified + '\n\n'
                        + '\n'.join(resolution.questions)
                        + '\n请直接回复说明；已完成的识别结果已保存，不会重复调用模型。')
            result = {'kind': 'generation', 'model_called': True, 'ok': False,
                      'status': 'waiting_user', 'questions': list(resolution.questions),
                      'pending': resolution_snapshot(resolution),
                      'artifacts': [], 'feedback': feedback}
            if manage_run:
                store.transition(run_id, 'validating', '资料识别及消歧完成，等待用户澄清')
                store.save_result(run_id, result)
                store.transition(run_id, 'waiting_user', '等待用户补充主体或期间')
            return result
        roles = dict(resolution.selected)
        if not roles.get('balance_sheet'):
            feedback = ('资料自动识别结果：\n' + identified + '\n\n'
                        '尚未识别到可确定主体、期间和范围的资产负债表。'
                        '请补充该资料或说明现有文件中哪一部分提供这些信息；没有生成正式工作簿。')
            result = {'kind': 'generation', 'model_called': True, 'ok': False,
                      'artifacts': [], 'feedback': feedback}
            if manage_run:
                store.transition(run_id, 'validating', '资料识别及生成范围检查')
                store.save_result(run_id, result)
                store.transition(run_id, 'failed', '资料已识别，尚需确认范围依据')
            return result
        bound_files = [f for f in files if f['id'] in roles.values()
                       or f['id'] in resolution.comparison_artifact_ids]
        if 'trial_balance' in roles:
            validate_roles(skill_id, bound_files, roles)
        elif any(Path(f['name']).suffix.lower() not in {'.xlsx', '.xls'} for f in bound_files):
            raise ValueError('资料已识别，但当前填报适配器需要可读取的 Excel 来源，不能仅凭识别文本生成')
    sources = work / 'inputs'
    sources.mkdir()
    selected = {}
    template = locked_template(skill_id) if skill_id != WORKFLOW_TO_SKILL.id else None
    template_digest = digest(template) if template is not None else None
    if template is not None:
        copied_template = sources / ('template' + template.suffix)
        shutil.copyfile(template, copied_template)
        if digest(copied_template) != template_digest:
            raise ValueError('锁定模板复制校验失败')
        selected['template'] = str(copied_template)
    by_id = {f['id']: f for f in files}
    for role, identity in roles.items():
        if cancel.is_set():
            raise TaskCancelled('任务已取消，未开始生成')
        item = by_id[identity]
        source = Path(item['path'])
        if digest(source) != item['sha256']:
            raise ValueError('资料已变化，请重新添加')
        target = sources / (role + source.suffix.lower())
        shutil.copyfile(source, target)
        if digest(target) != item['sha256']:
            raise ValueError('资料复制校验失败')
        selected[role] = str(target)
    # Multi-period statements: earlier periods are copied as comparison input
    # and listed for the pipeline's --financial-statement (latest first wins
    # the cover/valuation date via select_latest_statement).
    if automatic:
        statements = [selected['balance_sheet']] if 'balance_sheet' in selected else []
        for index, identity in enumerate(resolution.comparison_artifact_ids):
            item = by_id[identity]
            source = Path(item['path'])
            if digest(source) != item['sha256']:
                raise ValueError('资料已变化，请重新添加')
            target = sources / f'financial_statement_{index}{source.suffix.lower()}'
            shutil.copyfile(source, target)
            if digest(target) != item['sha256']:
                raise ValueError('资料复制校验失败')
            statements.append(str(target))
        if len(statements) > 1:
            selected['financial_statements'] = statements
    # Legacy .xls sources are converted inside this run's work directory; the
    # original file ids and hashes remain authoritative, and every conversion
    # is recorded with its tool. Hidden sheets never enter the converted copy.
    conversions = []
    for key, value in list(selected.items()):
        paths = value if isinstance(value, list) else [value]
        converted_paths = []
        for item_path in paths:
            path = Path(item_path)
            if path.suffix.lower() != '.xls':
                converted_paths.append(item_path)
                continue
            from .xls_support import convert_xls_to_xlsx
            converted, entry = convert_xls_to_xlsx(path, work / 'converted' / (path.stem + '.xlsx'))
            conversions.append({'input': key, **entry})
            converted_paths.append(str(converted))
        selected[key] = converted_paths if isinstance(value, list) else converted_paths[0]
    if conversions:
        (work / 'xls_conversion_report.json').write_text(
            json.dumps({'conversions': conversions}, ensure_ascii=False), encoding='utf-8')
    job = work / 'job.json'
    job.write_text(json.dumps({'skill_id': skill_id, 'inputs': selected}, ensure_ascii=False), encoding='utf-8')
    command = ([sys.executable, '--builtin-skill-worker', str(job)] if getattr(sys, 'frozen', False)
               else [sys.executable, '-X', 'utf8', '-m',
                     'asset_based_agent.technical_platform.generation_worker', str(job)])
    environment = dict(os.environ)
    environment['PYTHONPATH'] = str(Path(__file__).resolve().parents[2])
    environment['TEMP'] = environment['TMP'] = str(work)
    environment['PYTHONIOENCODING'] = 'utf-8'
    progress('正在本地生成副本；不上传资料，不调用模型。')
    if cancel.is_set():
        raise TaskCancelled('任务已取消，未启动生成进程')
    progress('等待本地生成资源；可取消排队。')
    office_lease = CLIENT_RESOURCES.lease(('office',), cancel) if skill_id == DETAIL.id else nullcontext()
    with office_lease, (work / 'worker.log').open('w', encoding='utf-8') as log:
        started = time.monotonic()
        process = subprocess.Popen(command, cwd=work, env=environment, stdout=log, stderr=log,
                                   creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        try:
            while process.poll() is None:
                if cancel.wait(0.25):
                    process.terminate()
                    process.wait(timeout=10)
                    break
                if time.monotonic() - started > 1800:
                    raise TimeoutError('本地生成超过30分钟，已停止')
                progress(f'本地生成中，已用 {int(time.monotonic() - started)} 秒；正在执行来源与成果校验。')
        finally:
            if process.poll() is None:
                process.kill()
                process.wait()
    if manage_run:
        store.transition(run_id, 'validating', '重新校验选定原件与生成状态')
    if ((template is not None and digest(template) != template_digest)
            or any(digest(Path(f['path'])) != f['sha256'] for f in files)):
        raise SourceValidationError('来源文件已变化，本轮成果不得发布')
    status_path = work / 'status.json'
    status = json.loads(status_path.read_text(encoding='utf-8')) if status_path.exists() else {}
    succeeded = not cancel.is_set() and process.returncode == 0 and status.get('ok') is True
    feedback = run_feedback(work, succeeded)
    if automatic:
        feedback = '资料自动识别结果：\n' + identified + '\n\n' + feedback
    if skill_id == HISTORY.id:
        names = ['history_fragment.docx', 'history_events.json', 'history_validation.json']
    elif skill_id == DETAIL.id:
        names = ['detail_workbook.xlsx', 'completion_status.json',
                 'delivery_check_report.json', 'execution_scope.json']
    elif skill_id == FINANCIAL_BRIEF.id:
        names = ['financial_brief.docx', 'financial_brief.pdf', 'financial_brief.png',
                 'extraction.json', 'timing.json']
    else:
        names = ['office_workflow_contract_validation.json']
    artifacts = register_artifacts(skill_id, work, names, succeeded)
    result = {'kind': 'generation', 'model_called': automatic, 'artifacts': artifacts,
              'feedback': feedback, 'ok': succeeded}
    if manage_run:
        store.save_result(run_id, result)
        store.transition(run_id, 'cancelled' if cancel.is_set() else 'succeeded' if succeeded else 'failed',
                         '生成与校验通过' if succeeded else '生成被阻断，未发布正式成果')
    return result
