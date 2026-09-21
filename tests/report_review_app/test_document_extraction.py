from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from asset_based_agent.report_review_app.domain.enums import FileRole
from asset_based_agent.report_review_app.domain.models import SourceFile
from asset_based_agent.report_review_app.services.document_extraction_service import (
    DocumentExtractionService,
)
from asset_based_agent.report_review_app.services.file_service import sha256_file


def source_file(path: Path, role: FileRole) -> SourceFile:
    return SourceFile(
        file_id=f"FILE-{path.stem}",
        original_name=path.name,
        extension=path.suffix,
        sha256=sha256_file(path),
        size_bytes=path.stat().st_size,
        role=role,
        round_number=1,
        original_path=str(path),
    )


def test_extract_docx_is_read_only_and_preserves_source_hash(tmp_path: Path) -> None:
    path = tmp_path / "报告.docx"
    document = Document()
    document.add_paragraph("第一段")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "金额"
    document.save(path)
    digest = sha256_file(path)

    extracted = DocumentExtractionService().extract(
        source_file(path, FileRole.MAIN_REPORT)
    )

    assert [chunk.text for chunk in extracted.chunks] == ["第一段", "项目 | 金额"]
    assert extracted.chunks[0].style_name == "Normal"
    assert sha256_file(path) == digest


def test_extract_docx_includes_header_and_footer_stories(tmp_path: Path) -> None:
    path = tmp_path / "报告.docx"
    document = Document()
    document.add_paragraph("正文")
    document.sections[0].header.paragraphs[0].text = "页眉报告编号XXXXX"
    document.sections[0].footer.paragraphs[0].text = "页脚项目名称"
    document.save(path)

    extracted = DocumentExtractionService().extract(
        source_file(path, FileRole.MAIN_REPORT)
    )
    stories = {(chunk.story_type, chunk.text) for chunk in extracted.chunks}

    assert ("header", "页眉报告编号XXXXX") in stories
    assert ("footer", "页脚项目名称") in stories


def test_extract_workbook_includes_sheet_and_cell_coordinates(tmp_path: Path) -> None:
    path = tmp_path / "测算表.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "结果"
    worksheet["A1"] = "评估值"
    worksheet["B1"] = 100
    workbook.save(path)

    extracted = DocumentExtractionService().extract(
        source_file(path, FileRole.CALCULATION_WORKBOOK)
    )

    assert extracted.chunks[0].text == "A1=评估值 | B1=100"
    assert extracted.chunks[0].location.table == "结果"
    assert extracted.extraction_mode == "openpyxl-readonly"


def test_extract_workbook_skips_unrelated_hidden_sheets(tmp_path: Path) -> None:
    path = tmp_path / "评估明细表.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "评估结果汇总表"
    summary["A1"] = "最终评估值"
    summary["B1"] = 100
    hidden = workbook.create_sheet("隐藏辅助表")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "不参与汇总的数据"
    workbook.save(path)

    extracted = DocumentExtractionService().extract(
        source_file(path, FileRole.CALCULATION_WORKBOOK)
    )

    assert {chunk.location.table for chunk in extracted.chunks} == {
        "评估结果汇总表"
    }


def test_extract_workbook_excludes_hidden_sheet_referenced_by_summary(
    tmp_path: Path,
) -> None:
    path = tmp_path / "评估明细表.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "评估结果汇总表"
    summary["A1"] = "最终评估值"
    summary["B1"] = "='隐藏数据'!B1"
    hidden = workbook.create_sheet("隐藏数据")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "评估值"
    hidden["B1"] = 100
    workbook.save(path)

    extracted = DocumentExtractionService().extract(
        source_file(path, FileRole.CALCULATION_WORKBOOK)
    )
    assert extracted.chunks == []


def test_extract_workbook_excludes_hidden_sheet_and_visible_rows_depending_on_it(
    tmp_path: Path,
) -> None:
    path = tmp_path / "评估明细表.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "评估结果汇总表"
    summary["A1"] = "不得审核的隐藏派生结果"
    summary["B1"] = "='隐藏数据'!B1"
    summary["A2"] = "可审核的可见结果"
    summary["B2"] = 200
    hidden = workbook.create_sheet("隐藏数据")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "机密计算参数"
    hidden["B1"] = 100
    workbook.save(path)

    extracted = DocumentExtractionService().extract(
        source_file(path, FileRole.CALCULATION_WORKBOOK)
    )

    assert [chunk.text for chunk in extracted.chunks] == [
        "A2=可审核的可见结果 | B2=200"
    ]
    assert all("隐藏数据" not in chunk.text for chunk in extracted.chunks)
    assert all("机密计算参数" not in chunk.text for chunk in extracted.chunks)


def test_extract_workbook_excludes_indirect_formula_referencing_hidden_sheet(
    tmp_path: Path,
) -> None:
    path = tmp_path / "评估明细表.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "评估结果汇总表"
    summary["A1"] = "不得审核的间接引用"
    summary["B1"] = '=INDIRECT("\'隐藏数据\'!B1")'
    summary["A2"] = "正常可见数据"
    summary["B2"] = 200
    hidden = workbook.create_sheet("隐藏数据")
    hidden.sheet_state = "hidden"
    hidden["B1"] = 100
    workbook.save(path)

    extracted = DocumentExtractionService().extract(
        source_file(path, FileRole.CALCULATION_WORKBOOK)
    )

    assert [chunk.text for chunk in extracted.chunks] == [
        "A2=正常可见数据 | B2=200"
    ]


def test_extract_workbook_excludes_defined_name_backed_by_hidden_sheet(
    tmp_path: Path,
) -> None:
    path = tmp_path / "评估明细表.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "评估结果汇总表"
    summary["A1"] = "不得审核的命名引用"
    summary["B1"] = "=HiddenParameter"
    summary["A2"] = "正常可见数据"
    summary["B2"] = 200
    hidden = workbook.create_sheet("隐藏数据")
    hidden.sheet_state = "hidden"
    hidden["B1"] = 100
    workbook.defined_names.add(
        DefinedName(
            "HiddenParameter",
            attr_text="'隐藏数据'!$B$1",
        )
    )
    workbook.save(path)

    extracted = DocumentExtractionService().extract(
        source_file(path, FileRole.CALCULATION_WORKBOOK)
    )

    assert [chunk.text for chunk in extracted.chunks] == [
        "A2=正常可见数据 | B2=200"
    ]


def test_extract_workbook_excludes_indirect_hidden_summary_dependencies(
    tmp_path: Path,
) -> None:
    path = tmp_path / "评估明细表.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "评估结果汇总表"
    summary["A1"] = "='隐藏一级'!A1"
    hidden_one = workbook.create_sheet("隐藏一级")
    hidden_one.sheet_state = "hidden"
    hidden_one["A1"] = "='隐藏二级'!A1"
    hidden_two = workbook.create_sheet("隐藏二级")
    hidden_two.sheet_state = "veryHidden"
    hidden_two["A1"] = 100
    workbook.save(path)

    extracted = DocumentExtractionService().extract(
        source_file(path, FileRole.CALCULATION_WORKBOOK)
    )

    assert extracted.chunks == []


def test_extract_workbook_does_not_use_hidden_template_summary_as_final_root(
    tmp_path: Path,
) -> None:
    path = tmp_path / "评估明细表.xlsx"
    workbook = Workbook()
    visible = workbook.active
    visible.title = "电子设备"
    visible["A1"] = "正常可见明细"
    summary = workbook.create_sheet("汇总表")
    summary.sheet_state = "hidden"
    summary["A1"] = "最终账面值"
    summary["B1"] = "='隐藏数据'!B1"
    hidden = workbook.create_sheet("隐藏数据")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "中间值"
    hidden["B1"] = "=#REF!-#REF!"
    workbook.save(path)

    extracted = DocumentExtractionService().extract(
        source_file(path, FileRole.CALCULATION_WORKBOOK)
    )

    extracted_sheets = {
        chunk.location.table
        for chunk in extracted.chunks
    }
    assert extracted_sheets == {"电子设备"}


def test_extract_workbook_excludes_unrelated_cells_on_hidden_dependency_sheet(
    tmp_path: Path,
) -> None:
    path = tmp_path / "评估明细表.xlsx"
    workbook = Workbook()
    visible = workbook.active
    visible.title = "电子设备"
    visible["A1"] = "正常可见明细"
    summary = workbook.create_sheet("汇总表")
    summary["B1"] = "='隐藏数据'!B1"
    hidden = workbook.create_sheet("隐藏数据")
    hidden.sheet_state = "hidden"
    hidden["B1"] = "=#REF!"
    hidden["B100"] = "=#REF!"
    workbook.save(path)

    extracted = DocumentExtractionService().extract(
        source_file(path, FileRole.CALCULATION_WORKBOOK)
    )
    assert {
        chunk.location.table
        for chunk in extracted.chunks
    } == {"电子设备"}
    assert all("隐藏数据" not in chunk.text for chunk in extracted.chunks)
