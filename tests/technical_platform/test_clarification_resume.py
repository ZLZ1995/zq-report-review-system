"""K05: reproduce the post-supplement understanding failure with diagnostics.

Scenario A drives the real two-round flow (insufficient files -> ask ->
supplement BS/PL -> resume) through AgentController + UnderstandingWorker.
Scenario B pins the misclassification: a server-side response-schema failure
must not surface as a network/connection problem.

Diagnostics (task id, revision, file id sets, question context, schema
results, exception class/stage) are written to the K05 artifact directory;
never tokens, cookies, absolute paths, or business content.
"""
import json
import sys
import time
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src'))

ARTIFACT_DIR = Path(r'D:\KimiData\kimi\Workspaces\Agent开发\k05_repro')

from openpyxl import Workbook

from asset_based_agent.technical_platform.agent_controller import AgentController
from asset_based_agent.technical_platform.skills import DETAIL, digest
from asset_based_agent.technical_platform.store import PlatformStore

ENTITY = 'A8T'
FULL_NAME = '阿里云飞天（北京）云计算有限公司'


def make_tb(path, period):
    wb = Workbook()
    ws = wb.active
    ws['A1'] = '科目汇总试算表'
    ws['A2'] = f'公司: {ENTITY} - {FULL_NAME}       期间: {period}    PRCGAAP'
    wb.save(path)


def make_tbd(path, periods):
    wb = Workbook()
    ws = wb.active
    ws['A1'] = '总帐明细帐追溯报表'
    ws['A2'] = f'公司：{ENTITY}-{ENTITY}帐套：PRCGAAP(CNY)'
    ws['A3'] = '期间'
    for row, period in enumerate(periods, start=4):
        ws.cell(row=row, column=1, value=period)
    wb.save(path)


def make_cf(path):
    wb = Workbook()
    ws = wb.active
    ws['A1'] = '现金流量表（未经审计）'
    wb.save(path)


def make_bs(path):
    wb = Workbook()
    ws = wb.active
    ws.title = '销项 1 (A8T)'
    ws['A1'] = 'PRC-资产负债表'
    ws['A3'] = '本期：2026-07'
    ws['A7'] = f'公司={ENTITY} ({FULL_NAME})'
    ws['E9'] = '期末余额'
    wb.save(path)


def make_pl(path):
    wb = Workbook()
    ws = wb.active
    ws.title = '销项 1 (A8T)'
    ws['A1'] = '利润表-PRC-PL'
    ws['A3'] = '本期：2026-07'
    ws['A7'] = f'公司={ENTITY} ({FULL_NAME})'
    wb.save(path)


class TwoRoundClient:
    """Conforming fake server: round 1 asks, round 2 plans."""

    def __init__(self):
        self.rounds = []

    def understand_task(self, payload, *, cancel=None):
        from asset_based_agent.agent_contracts import UnderstandingRequest
        request = UnderstandingRequest.model_validate(payload)  # request schema check
        self.rounds.append(request)
        if len(self.rounds) == 1:
            return {'schema_version': 1, 'message_intent': 'execute', 'goal': '',
                    'targets': [], 'references': [], 'excluded': [], 'constraints': [],
                    'deliverables': [],
                    'missing_inputs': [{'field': 'targets',
                                        'question': '缺少可作为填报依据的资产负债表，请补充。'}],
                    'evidence_message_ids': [request.message_id], 'skill_ids': [],
                    'next_action': 'ask',
                    'reply': '还缺少基准日资产负债表和利润表，请补充后回复。'}
        return {'schema_version': 1, 'message_intent': 'execute',
                'goal': '生成评估明细表',
                'targets': [f.id for f in request.files],
                'references': [], 'excluded': [], 'constraints': [],
                'deliverables': ['评估明细表'], 'missing_inputs': [],
                'evidence_message_ids': [request.message_id],
                'skill_ids': [DETAIL.id], 'next_action': 'plan',
                'reply': '资料已齐全，将生成评估明细表。'}


class ClarificationResumeReproTest(unittest.TestCase):
    def setUp(self):
        import tempfile
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        from PySide6.QtWidgets import QApplication
        self.app = QApplication.instance() or QApplication([])
        self.store = PlatformStore(self.dir / 'state.db', 'test')
        self.project = self.store.create_project('A8T')
        self.session = self.store.create_session(self.project)
        self.diagnostics = {'schema_version': 1, 'rounds': []}

    def tearDown(self):
        ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
        (ARTIFACT_DIR / 'diagnostics.json').write_text(
            json.dumps(self.diagnostics, ensure_ascii=False, indent=2), encoding='utf-8')
        self.tmp.cleanup()

    def add(self, path):
        return self.store.add_file(self.project, path, digest(path))

    def test_resume_after_supplement_reaches_plan(self):
        ids = []
        for period in ('2023-12', '2024-12', '2025-12', '2026-07'):
            path = self.dir / f'{period}_TB.xlsx'
            make_tb(path, period)
            ids.append(self.add(path))
        for index, periods in enumerate((('2023-01', '2023-12'), ('2024-01', '2024-12'),
                                         ('2025-01', '2025-12'), ('2026-01', '2026-07'))):
            path = self.dir / f'tbd_{index}.xlsx'
            make_tbd(path, periods)
            ids.append(self.add(path))
        cf = self.dir / 'A8T-CF202607.xlsx'
        make_cf(cf)
        ids.append(self.add(cf))

        client = TwoRoundClient()
        controller = AgentController(self.store)
        pending1 = controller.prepare(self.session, '根据上传的文件，帮我生成评估明细表',
                                      model_id='m', selected_ids=ids)
        from asset_based_agent.technical_platform.routing import UnderstandingWorker
        worker1 = UnderstandingWorker(client, pending1)
        worker1.run()
        self.assertIsNone(worker1.error, '第一轮理解不应失败')
        first = controller.complete(pending1, worker1.plan)
        self.assertEqual('ask', first.next_action)
        state = self.store  # question persisted; readable after client restart
        reopened = AgentController(state)
        self.diagnostics['rounds'].append({
            'round': 1, 'task_id': pending1.task_id, 'revision': pending1.revision,
            'file_ids': sorted(f.id for f in pending1.request.files),
            'decision': first.next_action,
            'question': first.reply[:200], 'exception': None, 'stage': 'complete'})

        # User supplements BS/PL converted to .xlsx and replies in the same session.
        bs = self.dir / 'A8T-BS202607.xlsx'
        make_bs(bs)
        pl = self.dir / 'A8T-PL202607.xlsx'
        make_pl(pl)
        ids += [self.add(bs), self.add(pl)]

        pending2 = reopened.prepare(self.session, '本轮已补充资产负债表和利润表',
                                    model_id='m', selected_ids=ids)
        worker2 = UnderstandingWorker(client, pending2)
        worker2.run()
        round2 = {'round': 2, 'task_id': pending2.task_id, 'revision': pending2.revision,
                  'file_ids': sorted(f.id for f in pending2.request.files),
                  'context_roles': [m.role for m in pending2.request.context],
                  'context_texts': [m.text[:80] for m in pending2.request.context],
                  'worker_error': worker2.error, 'exception': None, 'stage': 'understand'}
        self.diagnostics['rounds'].append(round2)
        self.assertIsNone(worker2.error, f'第二轮理解失败：{worker2.error}')
        self.assertEqual(2, len(client.rounds))
        second_request = client.rounds[1]
        # The resumed request must carry the original goal, the question and the answer.
        self.assertEqual('根据上传的文件，帮我生成评估明细表',
                         second_request.context[0].text)
        self.assertIn('补充', second_request.context[-1].text
                      if second_request.context[-1].role == 'assistant' else
                      second_request.prompt)
        self.assertEqual(sorted(ids), sorted(f.id for f in second_request.files))
        second = reopened.complete(pending2, worker2.plan)
        self.assertEqual('plan', second.next_action)
        self.assertEqual([DETAIL.id], second.skill_ids)
        # Clarification closed: a third turn starts a fresh understanding task.
        third = reopened.prepare(self.session, '再问一个问题', model_id='m', selected_ids=ids)
        self.assertNotEqual(pending2.task_id, third.task_id)

    def test_server_schema_failure_is_not_reported_as_network_problem(self):
        path = self.dir / 'bs.xlsx'
        make_bs(path)
        identity = self.add(path)
        from asset_based_agent.report_review_app.services.remote_auth_service import (
            RemoteAuthenticationError,
        )
        from asset_based_agent.technical_platform.routing import UnderstandingWorker

        class BrokenServer:
            def understand_task(self, payload, *, cancel=None):
                raise RemoteAuthenticationError(
                    '任务理解结果未通过校验，未开始执行，请补充本轮要求。')

        controller = AgentController(self.store)
        pending = controller.prepare(self.session, '生成评估明细表', model_id='m',
                                     selected_ids=[identity])
        worker = UnderstandingWorker(BrokenServer(), pending)
        worker.run()
        self.diagnostics['rounds'].append({
            'round': 'schema-failure', 'task_id': pending.task_id,
            'revision': pending.revision, 'worker_error': worker.error,
            'exception': 'RemoteAuthenticationError', 'stage': 'response-schema'})
        self.assertIsNotNone(worker.error)
        # A server-side response schema failure must say so; it is not a network issue.
        self.assertNotIn('连接', worker.error or '')
        self.assertNotIn('网络', worker.error or '')
        self.assertIn('校验', worker.error or '')
        self.assertIn('Skill尚未启动', worker.error or '')


if __name__ == '__main__':
    unittest.main()
