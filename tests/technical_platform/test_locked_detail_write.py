"""K08: locked-template detail writes must skip non-confirmed cells.

On locked templates the audit column (second book_value column H) is not a
confirmed input cell. The writer must leave it untouched instead of raising
a protection violation for the whole run.
"""
import importlib.util
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src'))

from openpyxl import Workbook

SCRIPTS = ROOT / '.codex' / 'skills' / 'valuation-detail-workbook-fill' / 'scripts'


def load_pipeline():
    sys.path.insert(0, str(SCRIPTS))
    spec = importlib.util.spec_from_file_location(
        'run_detail_workbook_pipeline', SCRIPTS / 'run_detail_workbook_pipeline.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def make_sheet():
    wb = Workbook()
    ws = wb.active
    ws.title = '预收账款'
    for col in 'ABCDEFGHI':
        ws[f'{col}5'] = '表头'
    ws['A27'] = '合计'
    ws._locked_template_layout = True
    confirmed = [f'{col}{row}' for col in 'ABCDEFGI' for row in range(6, 27)]
    registry = {'selected_sheets': {'预收账款': {
        'chain_role': 'detail',
        'confirmed_input_cells': confirmed,
        'detail_body_cells': list(confirmed),
        'forbidden_non_formula_cells': [],
        'role_rules': {'allow_title_inputs': False, 'allow_detail_body_inputs': True,
                       'allow_balance_sheet_inputs': False, 'allow_summary_inputs': False},
        'formula_cells': [],
    }}}
    protection = {'sheets': {'预收账款': {'formula_cells': []}}}
    return ws, registry, protection


class LockedDetailWriteTest(unittest.TestCase):
    def test_non_confirmed_audit_column_is_skipped_not_blocking(self):
        pipeline = load_pipeline()
        ws, registry, protection = make_sheet()
        rows = [{'counterparty': '某客商', 'book_value': -0.32, 'sub_name': '', 'remark': ''}]
        written = pipeline.write_simple_detail_sheet(ws, rows, 'detail', protection, registry, [])
        self.assertEqual(1, len(written))
        self.assertEqual('某客商', ws['B6'].value)
        self.assertEqual(-0.32, ws['G6'].value)
        self.assertIsNone(ws['H6'].value, '非确认输入列不得写入')

    def test_postfix_key_sheets_never_overwrites_formula_cells(self):
        pipeline = load_pipeline()
        wb = Workbook()
        ws = wb.active
        ws.title = '职工薪酬'
        ws['G6'] = '=F6'
        bs = {'values': {'应付职工薪酬': 1301505.97}, 'values_prior': {}}
        writes = pipeline.stage2_postfix_key_sheets(wb, bs)
        self.assertEqual(1301505.97, ws['F6'].value)
        self.assertEqual('=F6', ws['G6'].value, '公式单元格不得被静态值覆盖')
        self.assertNotIn('G6', {w['cell'] for w in writes})


if __name__ == '__main__':
    unittest.main()
