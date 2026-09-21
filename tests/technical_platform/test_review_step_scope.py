import threading
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from asset_based_agent.technical_platform.skills import REVIEW, digest
from asset_based_agent.technical_platform.store import PlatformStore
from asset_based_agent.technical_platform.task_spec import build_task_spec


def test_reference_workbook_is_filtered_before_role_conversion_and_never_reviewed(tmp_path):
    from asset_based_agent.report_review_app.domain.enums import FileRole
    from asset_based_agent.report_review_app.services.rule_registry import (
        IssueCandidate,
    )
    from asset_based_agent.technical_platform.adapters.review import ReviewAdapter
    from asset_based_agent.technical_platform.execution_plan import ExecutionPlan
    from asset_based_agent.technical_platform.harness import execute_plan
    from asset_based_agent.technical_platform.permissions import PermissionService
    from asset_based_agent.technical_platform.step_results import StepResults
    from asset_based_agent.technical_platform.tool_dispatcher import ToolDispatcher
    store = PlatformStore(tmp_path / 'db.sqlite', 'alice')
    project = store.create_project('synthetic')
    session = store.create_session(project)
    doc = Document()
    doc.add_paragraph('review target')
    path = tmp_path / 'report.docx'
    doc.save(path)
    store.add_file(project, path, digest(path))
    book = Workbook()
    book.active['A1'] = 'visible reference'
    book.create_sheet('secret').sheet_state = 'hidden'
    book['secret']['A1'] = 'DO_NOT_UPLOAD'
    book.active['A2'] = '=secret!A1'
    path = tmp_path / 'reference.xlsx'
    book.save(path)
    store.add_file(project, path, digest(path))
    files = store.files(project)
    target = next(f for f in files if f['name'] == 'report.docx')
    reference = next(f for f in files if f['name'] == 'reference.xlsx')
    snapshot = build_task_spec(store, session, 'Read reference, review report only', REVIEW, files,
                               model='test', instructions='rules').to_snapshot()
    step = snapshot['execution_plan']['steps'][0]
    step.update(target_inputs=[target['id']], reference_inputs=[reference['id']],
                goal='ONLY_THIS_STEP', constraints=['NO_ORIGINAL_EDITS'])
    run = store.start_run(session, snapshot)
    PermissionService(store).authorize(run, snapshot, confirmed=True)
    def issue(file):
        return IssueCandidate(source_file_id=file['id'], source_file_name=file['name'], category='test',
                              risk_level='low', location={}, description='synthetic issue', confidence=1)
    all_issues = [issue(target), issue(reference), issue({'id': 'old-file', 'name': 'old.docx'})]
    class Provider:
        model_id = 'test'
        skill_instructions = 'rules'
        def set_client_job_id(self, _):
            pass
        def review_batches(self, batches, progress_callback=None):
            chunks = [c for b in batches for c in b.chunks]
            assert 'DO_NOT_UPLOAD' not in str(chunks)
            assert not any(c.location.cell == 'A2' for c in chunks)
            refs = [c for c in chunks if c.source_file_id == reference['id']]
            assert refs and all(c.reference_only and c.role == FileRole.REFERENCE_DOCUMENT for c in refs)
            assert 'ONLY_THIS_STEP' in self.user_request and 'NO_ORIGINAL_EDITS' in self.user_request
            progress_callback({'state': 'output', 'issues': [i.model_dump(mode='json') for i in all_issues]})
            return all_issues
    outputs = []
    adapter = ReviewAdapter(store, run, provider=Provider(), progress=lambda _: None, output=outputs.extend)
    assert execute_plan(store, run, ExecutionPlan.model_validate(snapshot['execution_plan']),
                        ToolDispatcher({'review.execute': adapter}), threading.Event()) == 'succeeded'
    assert [i['source_file_id'] for i in outputs] == [target['id']]
    with store.connect() as db:
        result_id = db.execute('SELECT id FROM execution_results WHERE run=?', (run,)).fetchone()[0]
    result = StepResults(store).read(run, 'execute', result_id)
    assert [i['source_file_id'] for i in result['issues']] == [target['id']]
    assert all(digest(Path(f['path'])) == f['sha256'] for f in files)


@pytest.mark.parametrize('outside', [False, True])
def test_preflight_rejects_reference_only_review_and_unknown_reference_before_model(tmp_path, outside):
    import json

    from test_execution import make_run

    from asset_based_agent.technical_platform.skills import preflight
    store, run, _ = make_run(tmp_path, remote=True)
    source_id = json.loads(store.run(run)['snapshot'])['files'][0]['id']
    store.claim_run(run)
    with pytest.raises(ValueError, match='参考'):
        preflight(store, run, threading.Event(), lambda _: None, claimed=True, manage_run=False,
                  provider=object(), reference_file_ids={'outside' if outside else source_id})
    assert store.run(run)['result'] is None
