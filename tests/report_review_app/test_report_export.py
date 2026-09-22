from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from asset_based_agent.report_review_app.domain.enums import (
    IssueStatus,
    RiskLevel,
    RoundStatus,
)
from asset_based_agent.report_review_app.domain.models import (
    AuditRound,
    IssueHistoryEntry,
    IssueLocation,
    ReviewIssue,
)
from asset_based_agent.report_review_app.repositories.issue_repository import (
    IssueRepository,
)
from asset_based_agent.report_review_app.repositories.project_repository import (
    ProjectRepository,
)
from asset_based_agent.report_review_app.services.report_export_service import (
    ReportExportService,
)


def make_issue(issue_id: str, status: IssueStatus) -> ReviewIssue:
    return ReviewIssue(
        issue_id=issue_id,
        fingerprint=issue_id,
        source_file_id="FILE-1",
        source_file_name="评估报告.docx",
        category="data",
        risk_level=RiskLevel.HIGH,
        status=status,
        location=IssueLocation(chapter="评估结论", paragraph=1),
        original_text="100",
        description=f"问题-{issue_id}",
        recommendation="核对并修改",
        confidence=0.9,
        first_seen_round=1,
        last_seen_round=2,
        ignored=status == IssueStatus.IGNORED,
        history=[
            IssueHistoryEntry(round_number=1, status=IssueStatus.NEW),
            IssueHistoryEntry(round_number=2, status=status),
        ],
    )


def test_export_generates_deduplicated_summary_and_word_report(tmp_path: Path) -> None:
    repository = ProjectRepository(tmp_path / "projects")
    project = repository.create("项目A")
    round_dir = Path(project.project_path) / "rounds" / "round-002"
    issues_path = round_dir / "issues.json"
    issues = [
        make_issue("ISSUE-1", IssueStatus.FIXED),
        make_issue("ISSUE-1", IssueStatus.FIXED),
        make_issue("ISSUE-2", IssueStatus.IGNORED),
        make_issue("ISSUE-3", IssueStatus.UNMODIFIED),
    ]
    IssueRepository().save_issues(
        issues_path,
        project_id=project.project_id,
        round_number=2,
        issues=issues,
    )
    project.current_round = 2
    project.rounds.append(
        AuditRound(
            round_id="ROUND-2",
            round_number=2,
            status=RoundStatus.COMPLETED,
            progress_path=str(round_dir / "progress.json"),
            issues_path=str(issues_path),
            error_snapshot_path=str(round_dir / "error_snapshot.json"),
        )
    )

    selected_report_path = tmp_path / "user-selected" / "自选审核报告.docx"
    summary_path, report_path = ReportExportService().export(
        project,
        selected_report_path,
    )
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    document = Document(report_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert report_path == selected_report_path
    assert summary_path == selected_report_path.with_suffix(".summary.json")
    assert summary["issue_count"] == 3
    assert summary["status_counts"]["ignored"] == 1
    assert "六、已忽略问题汇总" in text
    assert "问题-ISSUE-2" in text
    assert text.count("问题-ISSUE-1") >= 1


def test_export_does_not_write_explicit_number_into_numbered_paragraph(
    tmp_path: Path,
) -> None:
    repository = ProjectRepository(tmp_path / "projects")
    project = repository.create("项目A")
    round_dir = Path(project.project_path) / "rounds" / "round-001"
    issues_path = round_dir / "issues.json"
    IssueRepository().save_issues(
        issues_path,
        project_id=project.project_id,
        round_number=1,
        issues=[make_issue("ISSUE-1", IssueStatus.NEW)],
    )
    project.current_round = 1
    project.rounds.append(
        AuditRound(
            round_id="ROUND-1",
            round_number=1,
            status=RoundStatus.COMPLETED,
            progress_path=str(round_dir / "progress.json"),
            issues_path=str(issues_path),
            error_snapshot_path=str(round_dir / "error_snapshot.json"),
        )
    )

    _, report_path = ReportExportService().export(
        project,
        tmp_path / "审核报告.docx",
    )
    document = Document(report_path)
    numbered = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "List Number"
    ]

    assert numbered
    assert numbered[0].text == "问题-ISSUE-1"


def test_export_never_includes_issue_conversation_content(
    tmp_path: Path,
) -> None:
    repository = ProjectRepository(tmp_path / "projects")
    project = repository.create("项目A")
    round_dir = Path(project.project_path) / "rounds" / "round-001"
    issues_path = round_dir / "issues.json"
    issue_repository = IssueRepository()
    issue_repository.save_issues(
        issues_path,
        project_id=project.project_id,
        round_number=1,
        issues=[make_issue("ISSUE-1", IssueStatus.NEW)],
    )
    issue_repository.save_conversations(
        round_dir / "conversations.json",
        {
            "schema_version": "1.0",
            "items": {
                "ISSUE-1": {
                    "messages": [
                        {
                            "role": "assistant",
                            "content": "仅供对话的最终修改方案-不可导出",
                            "created_at": "2026-07-31T00:00:00+00:00",
                        }
                    ]
                }
            },
        },
    )
    project.current_round = 1
    project.rounds.append(
        AuditRound(
            round_id="ROUND-1",
            round_number=1,
            status=RoundStatus.COMPLETED,
            progress_path=str(round_dir / "progress.json"),
            issues_path=str(issues_path),
            error_snapshot_path=str(round_dir / "error_snapshot.json"),
        )
    )

    summary_path, report_path = ReportExportService().export(
        project,
        tmp_path / "审核报告.docx",
    )
    report_text = "\n".join(
        paragraph.text for paragraph in Document(report_path).paragraphs
    )
    summary_text = summary_path.read_text(encoding="utf-8")

    assert "仅供对话的最终修改方案-不可导出" not in report_text
    assert "仅供对话的最终修改方案-不可导出" not in summary_text
