"""F01: pin the multi-period same-entity material routing defect and the new contract.

Real-machine evidence: run b08511b9 (2026-09-19) failed with
'存在多份同类资料，需要先确认本轮使用的主体和期间' when one entity supplied
2024/2025/2026-06 balance sheets plus bank materials.
"""

import pytest
from openpyxl import Workbook


def make_statement(path, entity, year, month, day):
    wb = Workbook()
    ws = wb.active
    ws.title = '资产负债表'
    ws['A1'] = '资产负债表'
    ws['A2'] = f'编制单位：{entity}'
    ws['C2'] = f'{year}年{month}月{day}日'
    ws['A3'] = '资产'
    ws['B3'] = '期末余额'
    ws['A4'] = '货币资金'
    ws['B4'] = 47499.02
    wb.save(path)


def files_arg(*paths):
    return [{'id': f'f{i}', 'name': p.name, 'path': str(p)} for i, p in enumerate(paths)]


def test_fixed_same_entity_multiperiod_no_longer_raises_duplicate(tmp_path):
    """Post-fix: same-entity multi-period resolves to the latest period; the old
    duplicate-role error survives only for genuinely ambiguous cases."""
    from asset_based_agent.technical_platform.material_analysis import resolve_roles
    paths = []
    for i, (y, m, d) in enumerate([(2024, 12, 31), (2025, 12, 31), (2026, 6, 30)]):
        p = tmp_path / f'statement_{i}.xlsx'
        make_statement(p, '北京甲示例科技有限公司', y, m, d)
        paths.append(p)
    plan = {'assignments': [{'file_id': f'f{i}', 'role': 'balance_sheet', 'reason': '表头'}
                            for i in range(3)]}
    assert resolve_roles(plan, files_arg(*paths)) == {'balance_sheet': 'f2'}
    # Files whose periods/entities cannot be told apart still refuse to guess.
    with pytest.raises(ValueError, match='存在多份同类资料，需要先确认本轮使用的主体和期间'):
        resolve_roles({'assignments': [{'file_id': 'one', 'role': 'balance_sheet'},
                                       {'file_id': 'two', 'role': 'balance_sheet'}]},
                      [{'id': 'one', 'name': 'a.xlsx', 'path': str(tmp_path / 'missing_a.xlsx')},
                       {'id': 'two', 'name': 'b.xlsx', 'path': str(tmp_path / 'missing_b.xlsx')}])


def test_resolve_materials_same_entity_three_periods_selects_latest(tmp_path):
    from asset_based_agent.technical_platform.material_analysis import resolve_materials
    paths = []
    for i, (y, m, d) in enumerate([(2024, 12, 31), (2025, 12, 31), (2026, 6, 30)]):
        p = tmp_path / f'statement_{i}.xlsx'
        make_statement(p, '北京甲示例科技有限公司', y, m, d)
        paths.append(p)
    bank = tmp_path / 'bank.xlsx'
    make_statement(bank, '北京甲示例科技有限公司', 2026, 6, 30)
    plan = {'assignments': [{'file_id': f'f{i}', 'role': 'balance_sheet', 'reason': '表头'}
                            for i in range(3)]
                      + [{'file_id': 'f3', 'role': 'bank_statement', 'reason': '流水'}]}
    resolution = resolve_materials(plan, files_arg(*paths, bank))
    assert resolution.status == 'resolved'
    assert resolution.selected['balance_sheet'] == 'f2'  # 2026-06-30 is latest
    assert resolution.selected['bank_statement'] == 'f3'
    assert set(resolution.comparison_artifact_ids) == {'f0', 'f1'}
    candidates = resolution.candidates['balance_sheet']
    assert len(candidates) == 3
    assert all(c.entity_name == '北京甲示例科技有限公司' for c in candidates)
    assert [c.period_end.isoformat() for c in candidates] == ['2024-12-31', '2025-12-31', '2026-06-30']
    assert resolution.reasons  # auditable selection rationale


def test_resolve_materials_different_entities_waits_user(tmp_path):
    from asset_based_agent.technical_platform.material_analysis import resolve_materials
    a = tmp_path / 'a.xlsx'
    b = tmp_path / 'b.xlsx'
    make_statement(a, '北京甲示例科技有限公司', 2025, 12, 31)
    make_statement(b, '上海示例科技有限公司', 2026, 6, 30)
    plan = {'assignments': [{'file_id': 'f0', 'role': 'balance_sheet', 'reason': 'x'},
                            {'file_id': 'f1', 'role': 'balance_sheet', 'reason': 'y'}]}
    resolution = resolve_materials(plan, files_arg(a, b))
    assert resolution.status == 'waiting_user'
    assert not resolution.selected
    assert any('北京甲示例科技有限公司' in q and '上海示例科技有限公司' in q for q in resolution.questions)


def test_resolve_materials_same_entity_same_period_conflict_waits_user(tmp_path):
    from asset_based_agent.technical_platform.material_analysis import resolve_materials
    a = tmp_path / 'v1.xlsx'
    b = tmp_path / 'v2.xlsx'
    make_statement(a, '北京甲示例科技有限公司', 2025, 12, 31)
    make_statement(b, '北京甲示例科技有限公司', 2025, 12, 31)
    b2 = Workbook()
    ws = b2.active
    ws.title = '资产负债表'
    ws['A1'] = '资产负债表'
    ws['A2'] = '编制单位：北京甲示例科技有限公司'
    ws['C2'] = '2025年12月31日'
    ws['A4'] = '货币资金'
    ws['B4'] = 99999.99  # conflicting content, same period
    b2.save(b)
    plan = {'assignments': [{'file_id': 'f0', 'role': 'balance_sheet', 'reason': 'x'},
                            {'file_id': 'f1', 'role': 'balance_sheet', 'reason': 'y'}]}
    resolution = resolve_materials(plan, files_arg(a, b))
    assert resolution.status == 'waiting_user'
    assert any('同一期间' in q or '版本' in q for q in resolution.questions)


def test_resolve_materials_period_unidentifiable_waits_user(tmp_path):
    from asset_based_agent.technical_platform.material_analysis import resolve_materials
    a = tmp_path / 'a.xlsx'
    b = tmp_path / 'b.xlsx'
    make_statement(a, '北京甲示例科技有限公司', 2025, 12, 31)
    wb = Workbook()
    ws = wb.active
    ws.title = '资产负债表'
    ws['A1'] = '资产负债表'
    ws['A2'] = '编制单位：北京甲示例科技有限公司'  # no date anywhere
    wb.save(b)
    plan = {'assignments': [{'file_id': 'f0', 'role': 'balance_sheet', 'reason': 'x'},
                            {'file_id': 'f1', 'role': 'balance_sheet', 'reason': 'y'}]}
    resolution = resolve_materials(plan, files_arg(a, b))
    assert resolution.status == 'waiting_user'
    assert any('期间' in q for q in resolution.questions)


def test_resolve_materials_single_balance_sheet_stays_compatible(tmp_path):
    from asset_based_agent.technical_platform.material_analysis import resolve_materials
    a = tmp_path / 'a.xlsx'
    make_statement(a, '北京甲示例科技有限公司', 2026, 6, 30)
    plan = {'assignments': [{'file_id': 'f0', 'role': 'balance_sheet', 'reason': '表头'}]}
    resolution = resolve_materials(plan, files_arg(a))
    assert resolution.status == 'resolved'
    assert resolution.selected == {'balance_sheet': 'f0'}
    assert resolution.comparison_artifact_ids == ()


def test_resolve_materials_unknown_file_and_role_still_rejected(tmp_path):
    from asset_based_agent.technical_platform.material_analysis import resolve_materials
    with pytest.raises(ValueError):
        resolve_materials({'assignments': [{'file_id': 'outside', 'role': 'balance_sheet'}]},
                          [{'id': 'one', 'name': 'a.xlsx', 'path': 'a.xlsx'}])
    with pytest.raises(ValueError):
        resolve_materials({'assignments': [{'file_id': 'one', 'role': 'surprise'}]},
                          [{'id': 'one', 'name': 'a.xlsx', 'path': 'a.xlsx'}])


def test_generation_persists_plan_and_waits_user_on_entity_conflict(tmp_path):
    """Model succeeded but disambiguation is impossible: raw plan must land on disk,
    the run must enter waiting_user with concrete questions, and no worker runs."""
    from threading import Event

    from asset_based_agent.technical_platform.execution import execute_task
    from asset_based_agent.technical_platform.material_analysis import (
        MaterialAnalysisProvider,
    )
    from asset_based_agent.technical_platform.permissions import PermissionService
    from asset_based_agent.technical_platform.skills import DETAIL, digest
    from asset_based_agent.technical_platform.store import PlatformStore
    from asset_based_agent.technical_platform.task_spec import build_task_spec

    a = tmp_path / 'entity_a.xlsx'
    b = tmp_path / 'entity_b.xlsx'
    make_statement(a, '北京甲示例科技有限公司', 2025, 12, 31)
    make_statement(b, '上海示例科技有限公司', 2026, 6, 30)
    store = PlatformStore(tmp_path / 'platform.sqlite', 'user')
    project = store.create_project('test')
    session = store.create_session(project)
    store.add_file(project, a, digest(a))
    store.add_file(project, b, digest(b))
    spec = build_task_spec(store, session, '生成明细表', DETAIL, store.files(project),
                           model='model', generation_confirmed=True).to_snapshot()
    run = store.start_run(session, spec)
    PermissionService(store).authorize(run, spec, confirmed=True)

    class Client:
        calls = 0

        def analyze_materials(self, payload):
            Client.calls += 1
            ids = [f['file_id'] for f in payload['files']]
            return {'assignments': [{'file_id': i, 'role': 'balance_sheet', 'reason': '表头'}
                                    for i in ids]}

    provider = MaterialAnalysisProvider(Client(), 'model', spec['skill_instructions'])
    result = execute_task(store, run, Event(), lambda _: None, provider=provider)
    assert store.run(run)['state'] == 'waiting_user'
    assert result.get('status') == 'waiting_user'
    assert any('北京甲示例科技有限公司' in q for q in result.get('questions', ()))
    assert any('上海示例科技有限公司' in q for q in result.get('questions', ()))
    plans = list(tmp_path.rglob('material_analysis.json'))
    assert plans, 'model classification must be persisted before disambiguation'
    assert Client.calls == 1, 'no repeated model billing while waiting for the user'


def _waiting_two_entity_setup(tmp_path):
    from asset_based_agent.technical_platform.material_analysis import (
        resolution_snapshot,
        resolve_materials,
    )
    a = tmp_path / 'entity_a.xlsx'
    b = tmp_path / 'entity_b.xlsx'
    make_statement(a, '北京甲示例科技有限公司', 2025, 12, 31)
    make_statement(b, '上海示例科技有限公司', 2026, 6, 30)
    files = [{'id': 'fa', 'name': a.name, 'path': str(a)},
             {'id': 'fb', 'name': b.name, 'path': str(b)}]
    plan = {'assignments': [{'file_id': 'fa', 'role': 'balance_sheet', 'reason': '表头'},
                            {'file_id': 'fb', 'role': 'balance_sheet', 'reason': '表头'}]}
    resolution = resolve_materials(plan, files)
    assert resolution.status == 'waiting_user'
    return resolution_snapshot(resolution), files


def test_match_clarification_selects_entity_by_name(tmp_path):
    from asset_based_agent.technical_platform.material_resume import match_clarification
    pending, files = _waiting_two_entity_setup(tmp_path)
    assert match_clarification(pending, files, '本轮使用北京甲示例科技有限公司的报表') == {
        'balance_sheet': 'fa'}
    assert match_clarification(pending, files, '用上海示例科技有限公司那份') == {
        'balance_sheet': 'fb'}


def test_match_clarification_selects_by_period_wording(tmp_path):
    from asset_based_agent.technical_platform.material_resume import match_clarification
    pending, files = _waiting_two_entity_setup(tmp_path)
    assert match_clarification(pending, files, '以 2026年6月 那期为准') == {'balance_sheet': 'fb'}


def test_match_clarification_selects_by_file_name(tmp_path):
    from asset_based_agent.technical_platform.material_analysis import (
        resolution_snapshot,
        resolve_materials,
    )
    from asset_based_agent.technical_platform.material_resume import match_clarification
    a = tmp_path / '六月版.xlsx'
    b = tmp_path / '六月修订版.xlsx'
    make_statement(a, '北京甲示例科技有限公司', 2026, 6, 30)
    make_statement(b, '北京甲示例科技有限公司', 2026, 6, 30)
    files = [{'id': 'fa', 'name': a.name, 'path': str(a)},
             {'id': 'fb', 'name': b.name, 'path': str(b)}]
    plan = {'assignments': [{'file_id': 'fa', 'role': 'balance_sheet', 'reason': '表头'},
                            {'file_id': 'fb', 'role': 'balance_sheet', 'reason': '表头'}]}
    pending = resolution_snapshot(resolve_materials(plan, files))
    assert pending['status'] == 'waiting_user'
    assert match_clarification(pending, files, '以六月修订版.xlsx 为准') == {'balance_sheet': 'fb'}


def test_match_clarification_ambiguous_or_unrelated_returns_none(tmp_path):
    from asset_based_agent.technical_platform.material_resume import match_clarification
    pending, files = _waiting_two_entity_setup(tmp_path)
    assert match_clarification(pending, files, '北京甲示例科技有限公司和上海示例科技有限公司都要') is None
    assert match_clarification(pending, files, '好的，继续') is None
    assert match_clarification(pending, files, '') is None


def test_resume_after_clarification_skips_model_and_completes(tmp_path, monkeypatch):
    """Full loop: waiting_user -> user clarifies -> resume without re-billing."""
    import json
    from pathlib import Path
    from threading import Event

    from asset_based_agent.technical_platform import generation
    from asset_based_agent.technical_platform.execution import execute_task
    from asset_based_agent.technical_platform.material_analysis import (
        MaterialAnalysisProvider,
    )
    from asset_based_agent.technical_platform.material_resume import match_clarification
    from asset_based_agent.technical_platform.permissions import PermissionService
    from asset_based_agent.technical_platform.skills import DETAIL, digest
    from asset_based_agent.technical_platform.store import PlatformStore
    from asset_based_agent.technical_platform.task_spec import build_task_spec

    a = tmp_path / 'entity_a.xlsx'
    b = tmp_path / 'entity_b.xlsx'
    make_statement(a, '北京甲示例科技有限公司', 2025, 12, 31)
    make_statement(b, '上海示例科技有限公司', 2026, 6, 30)
    store = PlatformStore(tmp_path / 'platform.sqlite', 'user')
    project = store.create_project('test')
    session = store.create_session(project)
    id_a = store.add_file(project, a, digest(a))
    store.add_file(project, b, digest(b))
    spec = build_task_spec(store, session, '生成明细表', DETAIL, store.files(project),
                           model='model', generation_confirmed=True).to_snapshot()
    run = store.start_run(session, spec)
    PermissionService(store).authorize(run, spec, confirmed=True)

    class Client:
        calls = 0

        def analyze_materials(self, payload):
            Client.calls += 1
            ids = [f['file_id'] for f in payload['files']]
            return {'assignments': [{'file_id': i, 'role': 'balance_sheet', 'reason': '表头'}
                                    for i in ids]}

    provider = MaterialAnalysisProvider(Client(), 'model', spec['skill_instructions'])
    result = execute_task(store, run, Event(), lambda _: None, provider=provider)
    assert store.run(run)['state'] == 'waiting_user'
    assert Client.calls == 1

    override = match_clarification(result['pending'], store.files(project),
                                   '本轮使用北京甲示例科技有限公司')
    assert override == {'balance_sheet': id_a}
    store.set_material_resolution_override(run, override)

    jobs = []

    class FakePopen:
        def __init__(self, command, cwd=None, env=None, stdout=None, stderr=None,
                     creationflags=0, **kwargs):
            self.returncode = 0
            jobs.append(json.loads(Path(command[-1]).read_text(encoding='utf-8')))
            Path(cwd, 'status.json').write_text(json.dumps({'ok': True}), encoding='utf-8')
            output = Path(cwd, 'output')
            output.mkdir(exist_ok=True)
            output.joinpath('detail_workbook.xlsx').write_bytes(b'fake-workbook')

        def poll(self):
            return self.returncode

        def terminate(self):
            pass

        def wait(self, timeout=None):
            return 0

        def kill(self):
            pass

    monkeypatch.setattr(generation.subprocess, 'Popen', FakePopen)
    result = execute_task(store, run, Event(), lambda _: None, provider=provider)
    assert store.run(run)['state'] == 'succeeded'
    assert result['ok'] is True
    assert Client.calls == 1, 'resume must not re-bill the completed model analysis'
    assert len(jobs) == 1
    assert jobs[0]['inputs']['balance_sheet'].endswith('balance_sheet.xlsx')
    assert 'financial_statements' not in jobs[0]['inputs']


def test_override_cannot_widen_beyond_candidates(tmp_path):
    import pytest as _pytest

    from asset_based_agent.technical_platform.material_analysis import (
        apply_resolution_override,
    )
    pending, files = _waiting_two_entity_setup(tmp_path)
    with _pytest.raises(ValueError):
        apply_resolution_override({'balance_sheet': 'outside'}, files, pending)
    with _pytest.raises(ValueError):
        apply_resolution_override({}, files, pending)
    resolved = apply_resolution_override({'balance_sheet': 'fa'}, files, pending)
    assert resolved.status == 'resolved'
    assert resolved.selected == {'balance_sheet': 'fa'}
    # The rejected candidate is a different entity, not comparison material.
    assert resolved.comparison_artifact_ids == ()


def _waiting_run_window(tmp_path):
    """Real store with a waiting_user run plus a minimal window stub."""
    from threading import Event
    from types import SimpleNamespace

    from asset_based_agent.technical_platform.execution import execute_task
    from asset_based_agent.technical_platform.material_analysis import (
        MaterialAnalysisProvider,
    )
    from asset_based_agent.technical_platform.permissions import PermissionService
    from asset_based_agent.technical_platform.skills import DETAIL, digest
    from asset_based_agent.technical_platform.store import PlatformStore
    from asset_based_agent.technical_platform.task_spec import build_task_spec

    a = tmp_path / 'entity_a.xlsx'
    b = tmp_path / 'entity_b.xlsx'
    make_statement(a, '北京甲示例科技有限公司', 2025, 12, 31)
    make_statement(b, '上海示例科技有限公司', 2026, 6, 30)
    store = PlatformStore(tmp_path / 'platform.sqlite', 'user')
    project = store.create_project('test')
    session = store.create_session(project)
    store.add_file(project, a, digest(a))
    store.add_file(project, b, digest(b))
    spec = build_task_spec(store, session, '生成明细表', DETAIL, store.files(project),
                           model='model', generation_confirmed=True).to_snapshot()
    run = store.start_run(session, spec)
    PermissionService(store).authorize(run, spec, confirmed=True)

    class Client:
        def analyze_materials(self, payload):
            ids = [f['file_id'] for f in payload['files']]
            return {'assignments': [{'file_id': i, 'role': 'balance_sheet', 'reason': '表头'}
                                    for i in ids]}

    execute_task(store, run, Event(), lambda _: None,
                 provider=MaterialAnalysisProvider(Client(), 'model', spec['skill_instructions']))
    assert store.run(run)['state'] == 'waiting_user'
    texts = []
    window = SimpleNamespace(
        store=store, session_id=session, client=None,
        status=SimpleNamespace(setText=texts.append),
        render_messages=lambda: None,
    )
    return store, run, session, window, texts


def test_ui_hook_ignores_sessions_without_waiting_run(tmp_path):
    from asset_based_agent.technical_platform.app import PlatformWindow
    from asset_based_agent.technical_platform.store import PlatformStore
    store = PlatformStore(tmp_path / 'platform.sqlite', 'user')
    project = store.create_project('test')
    session = store.create_session(project)
    from types import SimpleNamespace
    fake = SimpleNamespace(store=store, session_id=session)
    assert PlatformWindow.try_resume_waiting_run(fake, '随便说点什么') is False


def test_ui_hook_reasks_when_clarification_stays_ambiguous(tmp_path):
    from asset_based_agent.technical_platform.app import PlatformWindow
    store, run, session, window, texts = _waiting_run_window(tmp_path)
    assert PlatformWindow.try_resume_waiting_run(window, '好的，继续') is True
    assert store.run(run)['state'] == 'waiting_user'
    last = store.messages(session)[-1]
    assert last['role'] == 'assistant' and '无法唯一确定' in last['text']
    assert any('等待补充信息' in t for t in texts)


def test_ui_hook_cancels_waiting_run_on_user_request(tmp_path):
    from asset_based_agent.technical_platform.app import PlatformWindow
    store, run, session, window, texts = _waiting_run_window(tmp_path)
    assert PlatformWindow.try_resume_waiting_run(window, '算了，取消这个任务') is True
    assert store.run(run)['state'] == 'cancelled'
    assert any('已取消' in m['text'] for m in store.messages(session) if m['role'] == 'assistant')
