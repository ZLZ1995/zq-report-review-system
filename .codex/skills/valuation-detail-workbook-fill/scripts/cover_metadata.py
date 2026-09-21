"""Read cover identity and reporting date from explicitly supplied statements."""

import calendar
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def statement_date(value):
    if isinstance(value, (datetime, date)):
        return datetime(value.year, value.month, value.day)
    if not isinstance(value, str):
        return None
    matches = re.findall(r'(?<!\d)(\d{4})\s*[-/年]\s*(\d{1,2})\s*[-/月]\s*(\d{1,2})(?:日)?(?!\d)', value)
    if not matches:
        return None
    return max(datetime(*map(int, match)) for match in matches)


def read_statement_metadata(path):
    path = Path(path)
    records = []
    if path.suffix.lower() == '.xls':
        import xlrd
        book = xlrd.open_workbook(str(path))
        try:
            sheet = next((s for s in book.sheets() if '资产负债表' in s.name), book.sheet_by_index(0))
            sheet_name = sheet.name
            for r in range(min(sheet.nrows, 12)):
                for c in range(min(sheet.ncols, 15)):
                    cell = sheet.cell(r, c)
                    value = xlrd.xldate_as_datetime(cell.value, book.datemode) if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                    records.append((r + 1, c + 1, value))
        finally:
            book.release_resources()
    else:
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = next((s for s in book if '资产负债表' in s.title), book.worksheets[0])
            sheet_name = sheet.title
            records = [(c.row, c.column, c.value) for row in sheet.iter_rows(max_row=12, max_col=15)
                       for c in row if c.value is not None]
        finally:
            book.close()
    company = None
    dates = []
    company_cell = None
    by_position = {(r, c): v for r, c, v in records}
    for r, c, value in records:
        text = str(value or '').strip()
        if text in {'资产', '流动资产：', '流动资产:', '流动资产'}:
            break
        match = re.match(r'^(?:编制单位|单位名称)\s*[:：]\s*(.*)$', text)
        if not match:
            coded = re.match(r'^公司\s*[=:：]\s*[A-Za-z0-9]+\s*[(（](.+)[)）]\s*$', text)
            dashed = re.match(r'^公司\s*[=:：]\s*[A-Za-z0-9]+\s*-\s*(\S[^期]*?)\s*(?:期间.*)?$', text)
            if coded or dashed:
                name = (coded or dashed).group(1).strip()
                if name:
                    company = name
                    company_cell = f'{get_column_letter(c)}{r}'
        if match:
            name = match.group(1).strip()
            source_col = c
            if not name:
                for source_col in range(c + 1, c + 6):
                    name = str(by_position.get((r, source_col), '') or '').strip()
                    if name:
                        break
            if name:
                company = name
                company_cell = f'{get_column_letter(source_col)}{r}'
        parsed = statement_date(value)
        if parsed is not None and not any(word in text for word in ('打印', '导出', '编制日期', '填表日期')):
            dates.append((parsed, f'{get_column_letter(c)}{r}'))
    if not dates:
        for r, c, value in records:
            text = str(value or '').strip()
            match = re.search(r'(?:本期|期间)\s*[:：]?\s*(\d{4})\s*[-/年]\s*(\d{1,2})', text)
            if match:
                year, month = int(match.group(1)), int(match.group(2))
                if 1 <= month <= 12:
                    dates.append((datetime(year, month, calendar.monthrange(year, month)[1]),
                                  f'{get_column_letter(c)}{r}'))
                    break
    if not company or not dates:
        raise ValueError(f'statement_cover_metadata_missing: {path}')
    report_date, date_cell = max(dates)
    return {'source': str(path), 'sheet': sheet_name, 'company': company,
            'report_date': report_date, 'company_cell': company_cell, 'date_cell': date_cell}


def select_latest_statement(paths):
    candidates = [read_statement_metadata(p) for p in dict.fromkeys(map(str, paths))]
    if not candidates:
        raise ValueError('statement_inputs_empty')
    if len({item['company'] for item in candidates}) != 1:
        raise ValueError('statement_companies_conflict')
    return max(candidates, key=lambda item: item['report_date'])


def cover_assignments(wb, metadata):
    names = [name for name in ('封面', '封面页') if name in wb.sheetnames]
    if len(names) != 1:
        raise ValueError('cover_sheet_missing_or_ambiguous')
    ws = wb[names[0]]
    if '被评估单位' not in str(ws['D7'].value) or '评估基准日' not in str(ws['D9'].value):
        raise ValueError('cover_layout_not_confirmed')
    dt = metadata.get('report_date')
    if not metadata.get('company') or not isinstance(dt, (datetime, date)):
        raise ValueError('cover_metadata_missing')
    assignments = {'F7': metadata['company'], 'F9': dt.year, 'H9': dt.month, 'J9': dt.day}
    for address in assignments:
        if ws[address].data_type == 'f':
            raise ValueError(f'cover_target_contains_formula: {address}')
        for region in ws.merged_cells.ranges:
            if address in region and (address != region.start_cell.coordinate or address != 'F7'):
                raise ValueError(f'cover_target_merge_conflict: {address}')
    return ws, assignments


def validate_cover(wb, metadata):
    ws, assignments = cover_assignments(wb, metadata)
    if any(ws[cell].value != value for cell, value in assignments.items()):
        raise ValueError('cover_value_verification_failed')
    return {'status': 'pass', 'source': metadata.get('source'),
            'source_sheet': metadata.get('sheet'), 'company_cell': metadata.get('company_cell'),
            'date_cell': metadata.get('date_cell'), 'sheet': ws.title, 'writes': assignments}
