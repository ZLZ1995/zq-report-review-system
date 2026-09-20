"""K08: cover metadata must read A8T-format statement headers.

Real A8T statements carry `公司=CODE (全称)` and `本期：YYYY-MM` instead of
`编制单位：X` + a full date. The cover writer must derive the month-end
valuation date deterministically from the in-table period, never the name.
"""
import importlib.util
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src'))

from openpyxl import Workbook

SCRIPTS = ROOT / '.codex' / 'skills' / 'valuation-detail-workbook-fill' / 'scripts'


def load_cover_metadata():
    spec = importlib.util.spec_from_file_location('cover_metadata', SCRIPTS / 'cover_metadata.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def make_a8t_statement(path):
    wb = Workbook()
    ws = wb.active
    ws.title = '销项 1 (A8T)'
    ws['A1'] = 'PRC-资产负债表'
    ws['A3'] = '本期：2026-07'
    ws['A7'] = '公司=A8T (阿里云飞天（北京）云计算有限公司)'
    ws['A9'] = '资产'
    ws['E9'] = '期末余额'
    wb.save(path)


class CoverMetadataA8TTest(unittest.TestCase):
    def setUp(self):
        import tempfile
        self.tmp = tempfile.TemporaryDirectory(dir=r'D:\ZQ-Acceptance\tmp-pytest-basetemp')
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_company_code_parenthesized_full_name(self):
        module = load_cover_metadata()
        path = self.dir / 'A8T-BS202607.xlsx'
        make_a8t_statement(path)
        metadata = module.read_statement_metadata(path)
        self.assertEqual('阿里云飞天（北京）云计算有限公司', metadata['company'])

    def test_year_month_period_derives_month_end(self):
        module = load_cover_metadata()
        path = self.dir / 'A8T-BS202607.xlsx'
        make_a8t_statement(path)
        metadata = module.read_statement_metadata(path)
        self.assertEqual((2026, 7, 31),
                         (metadata['report_date'].year, metadata['report_date'].month,
                          metadata['report_date'].day))

    def test_legacy_header_format_still_works(self):
        module = load_cover_metadata()
        path = self.dir / 'legacy.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = '资产负债表'
        ws['B2'] = '资产负债表'
        ws['B3'] = '2026年6月30日'
        ws['B5'] = '编制单位：北京甲示例科技有限公司'
        ws['B6'] = '资产'
        wb.save(path)
        metadata = module.read_statement_metadata(path)
        self.assertEqual('北京甲示例科技有限公司', metadata['company'])
        self.assertEqual((2026, 6, 30), (metadata['report_date'].year,
                                         metadata['report_date'].month,
                                         metadata['report_date'].day))


if __name__ == '__main__':
    unittest.main()
