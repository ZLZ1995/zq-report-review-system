from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


PREFERRED_SHEETS = [
    "银行存款",
    "应收账款",
    "预付账款",
    "其他应收款",
    "应付账款",
    "预收账款",
    "其他应付款",
    "应交税费",
    "固定资产汇总",
    "无形资产汇总",
    "流动汇总",
    "非流动资产汇总",
    "流动负债汇总",
    "分类汇总",
    "资产负债表",
]


def clean(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def find_header_row(ws, max_rows: int = 12) -> int:
    for r in range(1, min(ws.max_row, max_rows) + 1):
        row_text = "|".join(clean(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 12) + 1))
        if "序号" in row_text and ("账面值" in row_text or "账面价值" in row_text or "欠款单位名称" in row_text):
            return r
    return 5


def find_total_row(ws, start_row: int, max_scan: int = 160) -> int:
    for r in range(start_row, min(ws.max_row, max_scan) + 1):
        row_text = "|".join(clean(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 12) + 1))
        if "合计" in row_text or "净值" in row_text or "净额" in row_text:
            return r
    return min(ws.max_row, max_scan)


def find_footer_row(ws, total_row: int) -> int:
    keywords = ("被评估单位填表人", "评估人员", "填表日期", "负责人")
    for r in range(total_row + 1, ws.max_row + 1):
        row_text = "|".join(clean(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 12) + 1))
        if any(key in row_text for key in keywords):
            return r
    return ws.max_row


def collect_formula_cells(ws) -> list[str]:
    formulas = []
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formulas.append(cell.coordinate)
    return formulas


def scan_sheet(ws) -> dict[str, Any]:
    header_row = find_header_row(ws)
    data_start_row = header_row + 1
    total_row = find_total_row(ws, data_start_row)
    footer_start_row = find_footer_row(ws, total_row)
    formula_cells = collect_formula_cells(ws)
    protected_text_ranges = []
    if footer_start_row <= ws.max_row:
        protected_text_ranges.append({"type": "footer_owner_date", "start_row": footer_start_row, "end_row": ws.max_row})
    protected_text_ranges.append({"type": "title_header", "start_row": 1, "end_row": header_row})
    writable_body = {
        "start_row": data_start_row,
        "end_row": max(data_start_row, total_row - 1),
        "start_col": 1,
        "end_col": ws.max_column,
    }
    writable_input_candidates = []
    for r in range(writable_body["start_row"], writable_body["end_row"] + 1):
        for c in range(writable_body["start_col"], writable_body["end_col"] + 1):
            cell = ws.cell(r, c)
            if isinstance(cell, MergedCell):
                continue
            if cell.coordinate in formula_cells:
                continue
            writable_input_candidates.append(cell.coordinate)
    summary_zone = {
        "start_row": total_row,
        "end_row": max(total_row, footer_start_row - 1),
        "start_col": 1,
        "end_col": ws.max_column,
    }
    return {
        "sheet": ws.title,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "header_row": header_row,
        "data_start_row": data_start_row,
        "total_row": total_row,
        "footer_start_row": footer_start_row,
        "merged_ranges": [str(item) for item in ws.merged_cells.ranges],
        "formula_cells": formula_cells,
        "writable_body": writable_body,
        "writable_input_candidates": writable_input_candidates,
        "summary_zone": summary_zone,
        "protected_text_ranges": protected_text_ranges,
    }


def build_formula_protection(layout_map: dict[str, Any]) -> dict[str, Any]:
    report = {"sheets": {}, "formula_cell_count": 0}
    for sheet_name, meta in layout_map["sheets"].items():
        report["sheets"][sheet_name] = {
            "formula_cells": meta["formula_cells"],
            "protected_text_ranges": meta["protected_text_ranges"],
            "readonly_summary_zone": meta["summary_zone"],
            "writable_body": meta["writable_body"],
        }
        report["formula_cell_count"] += len(meta["formula_cells"])
    return report


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--layout-output", required=True)
    parser.add_argument("--protection-output", required=True)
    parser.add_argument('--execution-scope')
    args = parser.parse_args()

    wb = load_workbook(Path(args.template))
    selected_sheets = [name for name in wb.sheetnames if name in PREFERRED_SHEETS]
    if args.execution_scope:
        scope = json.loads(Path(args.execution_scope).read_text('utf-8'))
        if scope['selected_mode'] != 'full_template':
            selected = set(scope['active_detail_sheets']) | set(scope['required_dependency_sheets'])
            selected_sheets = [name for name in wb.sheetnames if name in selected]
    layout_map = {
        "template": str(Path(args.template)),
        "sheet_order": selected_sheets,
        "sheets": {},
    }
    for name in selected_sheets:
        layout_map["sheets"][name] = scan_sheet(wb[name])
    wb.close()

    layout_output = Path(args.layout_output)
    protection_output = Path(args.protection_output)
    layout_output.parent.mkdir(parents=True, exist_ok=True)
    protection_output.parent.mkdir(parents=True, exist_ok=True)
    layout_output.write_text(json.dumps(layout_map, ensure_ascii=False, indent=2), encoding="utf-8")
    protection_output.write_text(json.dumps(build_formula_protection(layout_map), ensure_ascii=False, indent=2), encoding="utf-8")
    print(layout_output.as_posix())
    print(protection_output.as_posix())


if __name__ == "__main__":
    main()
