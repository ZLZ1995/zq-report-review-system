"""K03: local read-only material pre-identification summaries.

Fixtures mirror the real A8T source files (titles, entity/period markers,
hidden XDO_METADATA sheet). Summaries must be semantic-parity across
.xls/.xlsx, never include hidden sheets or local paths, and flag any
derived period instead of silently guessing.
"""
import json
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src'))

import openpyxl
import xlwt

from asset_based_agent.technical_platform import material_summary


def make_bs_xls(path):
    book = xlwt.Workbook()
    sheet = book.add_sheet('销项 1 (A8T)')
    sheet.write(0, 0, 'PRC-资产负债表')
    sheet.write(2, 0, '本期：2026-07')
    sheet.write(6, 0, '公司=A8T (阿里云飞天（北京）云计算有限公司)')
    sheet.write(8, 0, '资产')
    sheet.write(8, 4, '期末余额')
    sheet.write(8, 6, '期初余额')
    hidden = book.add_sheet('XDO_METADATA')
    hidden.write(0, 0, 'SECRET-HIDDEN-TOKEN')
    hidden.visibility = 1
    book.save(str(path))


def make_pl_xls(path):
    book = xlwt.Workbook()
    sheet = book.add_sheet('销项 1 (A8T)')
    sheet.write(0, 0, '利润表-PRC-PL')
    sheet.write(2, 0, '本期：2026-07')
    sheet.write(6, 0, '公司=A8T (阿里云飞天（北京）云计算有限公司)')
    sheet.write(8, 0, '项目')
    sheet.write(8, 3, '本期金额')
    sheet.write(8, 5, '本年累计金额')
    book.save(str(path))


def make_tb_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet0'
    ws['A1'] = '科目汇总试算表'
    ws['A2'] = '公司: A8T - 阿里云飞天（北京）云计算有限公司       期间: 2026-07    PRCGAAP'
    ws['A4'] = '科目编码'
    ws['B4'] = '科目名称'
    ws['C4'] = '期末借方'
    wb.save(path)


def make_tbd_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet0'
    ws['A1'] = '总帐明细帐追溯报表'
    ws['A2'] = '公司：A8T-A8T帐套：PRCGAAP(CNY)'
    ws['A3'] = '期间'
    ws['B3'] = '凭证号'
    ws['C3'] = '摘要'
    ws['D3'] = '公司描述'
    for row, period in ((4, '2026-01'), (5, '2026-03'), (6, '2026-07')):
        ws.cell(row=row, column=1, value=period)
        ws.cell(row=row, column=4, value='阿里云飞天（北京）云计算有限公司')
    wb.save(path)


def make_cf_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '现金流量表'
    ws['A1'] = '现金流量表（未经审计）'
    for row, label in ((2, '识别号：'), (3, '名称：'), (4, '报表时期：'), (5, '申报状态：'), (6, '金额单位：')):
        ws.cell(row=row, column=1, value=label)
    ws['A8'] = '项目'
    ws['B8'] = '本期金额'
    wb.save(path)


class MaterialSummaryTest(unittest.TestCase):
    def setUp(self):
        import tempfile
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def summarize(self, path):
        return material_summary.summarize_file(path, artifact_id='file-1', name=path.name)

    def test_required_keys(self):
        path = self.dir / 'A8T-BS202607.xls'
        make_bs_xls(path)
        summary = self.summarize(path)
        self.assertEqual(
            {'artifact_id', 'name', 'format', 'readable', 'document_type',
             'entity_name', 'period_start', 'period_end', 'sheet_names',
             'header_evidence', 'confidence', 'warnings'},
            set(summary))
        json.dumps(summary, ensure_ascii=False)

    def test_xls_balance_sheet(self):
        path = self.dir / 'A8T-BS202607.xls'
        make_bs_xls(path)
        summary = self.summarize(path)
        self.assertEqual('file-1', summary['artifact_id'])
        self.assertEqual('A8T-BS202607.xls', summary['name'])
        self.assertEqual('xls', summary['format'])
        self.assertTrue(summary['readable'])
        self.assertEqual('balance_sheet', summary['document_type'])
        self.assertEqual('A8T', summary['entity_name'])
        self.assertEqual('2026-07-31', summary['period_end'])
        self.assertEqual(['销项 1 (A8T)'], summary['sheet_names'])
        self.assertIn('资产负债表', summary['header_evidence'])
        self.assertIn('期末余额', summary['header_evidence'])
        self.assertTrue(any('推导' in w or '月末' in w for w in summary['warnings']),
                        summary['warnings'])
        self.assertGreaterEqual(summary['confidence'], 0.8)

    def test_hidden_sheet_never_enters_summary(self):
        path = self.dir / 'A8T-BS202607.xls'
        make_bs_xls(path)
        summary = self.summarize(path)
        self.assertNotIn('XDO_METADATA', summary['sheet_names'])
        blob = json.dumps(summary, ensure_ascii=False)
        self.assertNotIn('SECRET-HIDDEN-TOKEN', blob)
        self.assertNotIn('XDO_METADATA', blob)

    def test_no_absolute_path_in_summary(self):
        path = self.dir / 'A8T-BS202607.xls'
        make_bs_xls(path)
        blob = json.dumps(self.summarize(path), ensure_ascii=False)
        self.assertNotIn(str(self.dir), blob)
        self.assertNotIn(str(path), blob)

    def test_xls_income_statement(self):
        path = self.dir / 'A8T-PL202607.xls'
        make_pl_xls(path)
        summary = self.summarize(path)
        self.assertEqual('income_statement', summary['document_type'])
        self.assertEqual('A8T', summary['entity_name'])
        self.assertEqual('2026-07-31', summary['period_end'])

    def test_trial_balance_xlsx(self):
        path = self.dir / '2026.07_TB.xlsx'
        make_tb_xlsx(path)
        summary = self.summarize(path)
        self.assertEqual('xlsx', summary['format'])
        self.assertEqual('trial_balance', summary['document_type'])
        self.assertEqual('A8T', summary['entity_name'])
        self.assertEqual('2026-07-31', summary['period_end'])
        self.assertIn('科目汇总试算表', summary['header_evidence'])

    def test_journal_xlsx_period_from_data(self):
        path = self.dir / '2026.1-7_TBD.xlsx'
        make_tbd_xlsx(path)
        summary = self.summarize(path)
        self.assertEqual('journal', summary['document_type'])
        self.assertEqual('A8T', summary['entity_name'])
        self.assertEqual('2026-07-31', summary['period_end'])

    def test_cash_flow_with_empty_headers_warns(self):
        path = self.dir / 'A8T-CF202607.xlsx'
        make_cf_xlsx(path)
        summary = self.summarize(path)
        self.assertEqual('cash_flow_statement', summary['document_type'])
        self.assertIsNone(summary['entity_name'])
        self.assertIsNone(summary['period_end'])
        text = ' '.join(summary['warnings'])
        self.assertIn('主体', text)
        self.assertIn('期间', text)
        self.assertLess(summary['confidence'], 0.8)

    def test_xls_xlsx_semantic_parity(self):
        xls_path = self.dir / 'bs.xls'
        make_bs_xls(xls_path)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '销项 1 (A8T)'
        ws['A1'] = 'PRC-资产负债表'
        ws['A3'] = '本期：2026-07'
        ws['A7'] = '公司=A8T (阿里云飞天（北京）云计算有限公司)'
        ws['A9'] = '资产'
        ws['E9'] = '期末余额'
        ws['G9'] = '期初余额'
        xlsx_path = self.dir / 'bs.xlsx'
        wb.save(xlsx_path)
        xls_summary = self.summarize(xls_path)
        xlsx_summary = self.summarize(xlsx_path)
        for key in ('document_type', 'entity_name', 'period_start', 'period_end', 'sheet_names'):
            self.assertEqual(xls_summary[key], xlsx_summary[key], key)

    def test_unreadable_file(self):
        path = self.dir / 'broken.xls'
        path.write_bytes(b'\x00\x01not an ole2 workbook')
        summary = self.summarize(path)
        self.assertFalse(summary['readable'])
        self.assertEqual('other', summary['document_type'])
        self.assertTrue(summary['warnings'])
        self.assertLess(summary['confidence'], 0.5)

    def test_unsupported_extension(self):
        path = self.dir / 'notes.csv'
        path.write_text('a,b,c', encoding='utf-8')
        summary = self.summarize(path)
        self.assertFalse(summary['readable'])
        self.assertTrue(any('不支持' in w or '扩展名' in w for w in summary['warnings']),
                        summary['warnings'])


if __name__ == '__main__':
    unittest.main()
