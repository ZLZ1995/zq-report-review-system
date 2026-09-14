from pathlib import Path
from shutil import copy2
import argparse
import json

from docx import Document


def cell_text(cell):
    return " ".join(cell.text.replace("\n", "|").split())


def extract(path):
    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    event_paras = [p for p in paragraphs if p.startswith("（")]
    tables = []
    for idx, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append([cell_text(cell) for cell in row.cells])
        tables.append(
            {
                "idx": idx,
                "rows": len(table.rows),
                "cols": len(table.columns),
                "first_rows": rows[:4],
                "last_rows": rows[-3:],
                "all_rows": rows,
            }
        )
    return {
        "path": str(path),
        "paragraph_count": len(paragraphs),
        "event_paragraphs": event_paras,
        "table_count": len(tables),
        "tables": tables,
    }


def main():
    parser = argparse.ArgumentParser(description="Compare generated and reference gongshang DOCX structure.")
    parser.add_argument("--generated", required=True, help="Generated DOCX path.")
    parser.add_argument("--reference", required=True, help="Human-corrected/reference DOCX path.")
    parser.add_argument("--output", required=True, help="Output JSON path.")
    args = parser.parse_args()
    output = Path(args.output)
    work = output.parent / "gongshang_compare_work"
    work.mkdir(parents=True, exist_ok=True)
    base_ascii = work / "generated.docx"
    human_ascii = work / "reference.docx"
    copy2(Path(args.generated), base_ascii)
    copy2(Path(args.reference), human_ascii)
    result = {
        "generated": extract(base_ascii),
        "reference": extract(human_ascii),
        "ascii_base": str(base_ascii),
        "ascii_human": str(human_ascii),
    }
    generated_shapes = [(t["cols"], t["rows"]) for t in result["generated"]["tables"]]
    reference_shapes = [(t["cols"], t["rows"]) for t in result["reference"]["tables"]]
    result["diff_count"] = 0
    if result["generated"]["event_paragraphs"] != result["reference"]["event_paragraphs"]:
        result["diff_count"] += 1
    if generated_shapes != reference_shapes:
        result["diff_count"] += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(output)
    print("generated events", len(result["generated"]["event_paragraphs"]), "tables", result["generated"]["table_count"])
    print("reference events", len(result["reference"]["event_paragraphs"]), "tables", result["reference"]["table_count"])


if __name__ == "__main__":
    main()
