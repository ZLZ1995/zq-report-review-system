from pathlib import Path

import pytest

from asset_based_agent.technical_platform import generation_worker
from asset_based_agent.technical_platform.skills import HISTORY


def test_history_release_binds_the_approved_template():
    from asset_based_agent.technical_platform.generation import locked_template
    approved = Path(__file__).resolve().parents[2] / 'assets/builtin_templates' / HISTORY.id / 'template.docx'
    assert locked_template(HISTORY.id).read_bytes() == approved.read_bytes()


def test_history_worker_has_a_real_adapter(tmp_path, monkeypatch):
    import json

    from docx import Document
    from openpyxl import Workbook

    source = tmp_path / 'source.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = '变更信息'
    ws.append(['变更日期', '变更事项', '变更前', '变更后'])
    ws.append(['2026-01-01', '法定代表人变更', '张甲', '李乙'])
    wb.save(source)
    root = Path(__file__).resolve().parents[2]
    template = root / 'assets/builtin_templates' / HISTORY.id / 'template.docx'
    job = tmp_path / 'job.json'
    job.write_text(json.dumps({'skill_id': HISTORY.id, 'inputs': {
        'source_excel': str(source), 'template': str(template)}}, ensure_ascii=False), encoding='utf-8')
    assert generation_worker.main(job) == 0
    output = tmp_path / 'output'
    doc = Document(output / 'history_fragment.docx')
    assert len(doc.tables) == 1
    assert '张甲' in doc.tables[0].cell(1, 1).text
    assert '李乙' in doc.tables[0].cell(1, 2).text
    assert not any('【变更' in p.text for p in doc.paragraphs)
    assert json.loads((output / 'history_validation.json').read_text('utf-8'))['ok']


@pytest.mark.parametrize('before,after,expected', [
    ('甲公司 货币100万人民币', '乙公司 货币100万人民币', 0),
    ('甲公司;乙公司', '丙公司;丁公司', 2),
])
def test_equity_generation_requires_amount_evidence(tmp_path, before, after, expected):
    import json

    from openpyxl import Workbook
    source = tmp_path / 'source.xlsx'
    wb = Workbook()
    wb.active.title = '变更信息'
    wb.active.append(['变更日期', '变更事项', '变更前', '变更后'])
    wb.active.append(['2026-01-01', '股东出资变更', before, after])
    wb.save(source)
    template = Path(__file__).resolve().parents[2] / 'assets/builtin_templates' / HISTORY.id / 'template.docx'
    original = template.read_bytes(), source.read_bytes()
    job = tmp_path / 'job.json'
    job.write_text(json.dumps({'skill_id': HISTORY.id, 'inputs': {'source_excel': str(source), 'template': str(template)}}), encoding='utf-8')
    assert generation_worker.main(job) == expected
    assert (tmp_path / 'output/history_fragment.docx').exists() == (expected == 0)
    assert original == (template.read_bytes(), source.read_bytes())


def test_more_than_five_events_preserve_template_package_and_format(tmp_path):
    import json
    from zipfile import ZipFile

    from docx import Document
    from docx.oxml.ns import qn
    from openpyxl import Workbook
    source = tmp_path / 'source.xlsx'
    wb = Workbook()
    wb.active.title = '变更信息'
    wb.active.append(['变更日期', '变更事项', '变更前', '变更后'])
    for day in range(1, 9):
        wb.active.append([f'2026-01-{day:02}', '公司住所变更', f'示例市原办公地址第{day}号', f'示例市新办公地址第{day}号'])
    wb.save(source)
    template = Path(__file__).resolve().parents[2] / 'assets/builtin_templates' / HISTORY.id / 'template.docx'
    job = tmp_path / 'job.json'
    job.write_text(json.dumps({'skill_id': HISTORY.id, 'inputs': {'source_excel': str(source), 'template': str(template)}}), encoding='utf-8')
    assert generation_worker.main(job) == 0
    final = tmp_path / 'output/history_fragment.docx'
    doc = Document(final)
    assert len(doc.tables) == 8
    assert all(p.paragraph_format.keep_with_next for p in doc.paragraphs if '详情如下' in p.text)
    with ZipFile(template) as src, ZipFile(final) as dst:
        assert src.namelist() == dst.namelist()
        assert all(src.read(n) == dst.read(n) for n in src.namelist() if n != 'word/document.xml')
    for p in doc._element.xpath('.//w:p'):
        in_table = any(a.tag == qn('w:tc') for a in p.iterancestors())
        assert p.xpath('./w:pPr/w:ind/@w:firstLineChars') == ['0' if in_table else '200']
        assert all(x == ('18' if in_table else '24') for x in p.xpath('./w:r/w:rPr/w:sz/@w:val'))


def test_history_task_runs_subprocess_and_exposes_conversation_artifact(tmp_path):
    from threading import Event

    from openpyxl import Workbook

    from asset_based_agent.technical_platform.execution import execute_task
    from asset_based_agent.technical_platform.generation import artifact_path
    from asset_based_agent.technical_platform.skills import digest
    from asset_based_agent.technical_platform.store import PlatformStore
    from asset_based_agent.technical_platform.task_spec import build_task_spec
    source = tmp_path / 'source.xlsx'
    wb = Workbook()
    wb.active.title = '变更信息'
    wb.active.append(['变更日期', '变更事项', '变更前', '变更后'])
    wb.active.append(['2026-01-01', '法定代表人变更', '张甲', '李乙'])
    wb.save(source)
    store = PlatformStore(tmp_path / 'platform.sqlite', 'fixture-user')
    project = store.create_project('生成验收样例')
    session = store.create_session(project)
    store.add_file(project, source, digest(source))
    files = store.files(project)
    spec = build_task_spec(store, session, '生成工商历史沿革', HISTORY, files,
                           input_roles={'source_excel': files[0]['id']}, generation_confirmed=True)
    run = store.start_run(session, spec.to_snapshot())
    from asset_based_agent.technical_platform.permissions import PermissionService
    PermissionService(store).authorize(run, spec.to_snapshot(), confirmed=True)
    result = execute_task(store, run, Event(), lambda _: None)
    assert result['ok'] is True
    assert store.run(run)['state'] == 'succeeded'
    assert result['model_called'] is False
    with store.connect() as db:
        assert db.execute('SELECT COUNT(*) FROM execution_results WHERE run=?', (run,)).fetchone()[0] == 1
    assert artifact_path(store, session, run, 0).name == 'history_fragment.docx'
