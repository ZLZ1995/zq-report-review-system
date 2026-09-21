import json
import re
import shutil
import subprocess
import sys
import time
import hashlib
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from docx import Document

SKILL = Path(__file__).resolve().parents[1]

def clean(value):
    return re.sub(r'[\s：:（）()、]', '', str(value or ''))

def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def extract(ws, item, header):
    hits = []
    for row in ws:
        for cell in row:
            label = clean(cell.value)
            if label == item or (item == 'equity' and label in
                    ('所有者权益或股东权益', '所有者权益或股东权益合计')) or (
                    item in ('营业收入', '利润总额', '净利润') and
                    re.match(r'^[一二三四五六七八九十]+[.．]?'+item, label)):
                # Section-only labels are not data rows.
                cols = [c.column for c in ws[4] if clean(c.value) == header and c.column > cell.column]
                if not cols:
                    continue
                value_cell = ws.cell(cell.row, min(cols))
                if value_cell.value is None and item == 'equity':
                    continue
                hits.append(value_cell)
    if len(hits) != 1:
        raise ValueError(f'{ws.title} {item}: expected one match, found {len(hits)}')
    c = hits[0]
    if not isinstance(c.value, (int, float)) or isinstance(c.value, bool):
        raise ValueError(f'{ws.title}!{c.coordinate}: missing numeric value or formula cache')
    return {'sheet': ws.title, 'cell': c.coordinate, 'yuan': str(Decimal(str(c.value)))}

def replace(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(text)

def main():
    start = time.perf_counter()
    utc_start = datetime.now(timezone.utc).isoformat()
    cfg = json.loads(Path(sys.argv[1]).read_text(encoding='utf-8'))
    out = Path(cfg['output_dir'])
    out.mkdir(parents=True, exist_ok=False)
    timings = {'started_utc': utc_start, 'stages_seconds': {}, 'visual_review': 'pending'}
    checkpoint = start
    def stage(name):
        nonlocal checkpoint
        now = time.perf_counter()
        timings['stages_seconds'][name] = round(now - checkpoint, 4)
        checkpoint = now
    try:
        sources = cfg['sources']
        if len(sources) != 3:
            raise ValueError('Three source workbooks required')
        original_hashes = []
        copies = []
        for i, source in enumerate(sources):
            original_hashes.append(digest(source['path']))
            copy = out / f'input_{i}.xlsx'
            shutil.copy2(source['path'], copy)
            copies.append(copy)
        stage('source_copy_and_hash')
        records = []
        for source, path in zip(sources, copies):
            wb = load_workbook(path, data_only=True)
            try:
                bs, pl = wb['资产负债表'], wb['利润表']
                for ws in (bs, pl):
                    meta = ''.join(str(c.value or '') for row in ws.iter_rows(max_row=3) for c in row)
                    if cfg['company'] not in meta:
                        raise ValueError('Company mismatch')
                    date_match = re.search(r'(\d{4})年(\d{1,2})月(?:(\d{1,2})日)?', meta)
                    dt = datetime.fromisoformat(source['date'])
                    if not date_match or (int(date_match[1]), int(date_match[2])) != (dt.year, dt.month):
                        raise ValueError('Reporting period mismatch')
                    if ws == bs and (not date_match[3] or int(date_match[3]) != dt.day):
                        raise ValueError('Balance sheet date mismatch')
                fields = {key: extract(bs, item, '期末余额') for key, item in
                          [('assets', '资产总计'), ('liabilities', '负债合计'), ('equity', 'equity')]}
                fields.update({key: extract(pl, item, '本年金额') for key, item in
                               [('revenue', '营业收入'), ('pretax_profit', '利润总额'), ('net_profit', '净利润')]})
                difference = Decimal(fields['assets']['yuan']) - Decimal(fields['liabilities']['yuan']) - Decimal(fields['equity']['yuan'])
                if abs(difference) > Decimal('0.01'):
                    raise ValueError('Balance sheet does not reconcile')
                records.append({'source': source, 'fields': fields, 'balance_difference_yuan': str(difference)})
            finally:
                wb.close()
        stage('extract_and_reconcile')
        doc = Document(cfg.get('template', str(SKILL / 'assets' / 'financial_table.docx')))
        if len(doc.tables) != 1 or len(doc.tables[0].rows) != 8 or len(doc.tables[0].columns) != 4:
            raise ValueError('Template layout mismatch')
        table = doc.tables[0]
        zeros = []
        for i, rec in enumerate(records, 1):
            replace(table.cell(0, i).paragraphs[0], rec['source']['balance_label'])
            replace(table.cell(4, i).paragraphs[0], rec['source']['income_label'])
            for row, key in [(1,'assets'),(2,'liabilities'),(3,'equity'),(5,'revenue'),(6,'pretax_profit'),(7,'net_profit')]:
                yuan = Decimal(rec['fields'][key]['yuan'])
                wan = (yuan / Decimal('10000')).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
                formatted = f'{wan:,.2f}' if wan else '0.00'
                replace(table.cell(row,i).paragraphs[0], formatted)
                rec['fields'][key]['display_wan'] = formatted
                if yuan and not wan:
                    zeros.append(f"{rec['source']['balance_label']}的{table.cell(row,0).text.replace('（万元）','')}为{yuan:.2f}元")
        for sec in doc.sections:
            replace(sec.header.paragraphs[0], cfg.get('report_header', cfg['company']+'财务状况简表'))
            for t in sec.footer.tables:
                replace(t.cell(0,0).paragraphs[0], cfg.get('footer_org', ''))
        period_text = '、'.join(s['income_label'] for s in sources)
        note = f'注：上述{period_text}财务数据来源于企业提供的财务报表。表内金额以万元列示并保留两位小数。'
        if zeros:
            note += '；'.join(zeros)+'，四舍五入后显示为0.00万元。'
        for p in doc.paragraphs:
            if p.text.startswith('注：'):
                replace(p, note)
        doc.core_properties.title = cfg['company']+'财务状况简表'
        output = out / 'financial_brief.docx'
        doc.save(output)
        check = Document(output)
        for i, rec in enumerate(records,1):
            for row,key in [(1,'assets'),(2,'liabilities'),(3,'equity'),(5,'revenue'),(6,'pretax_profit'),(7,'net_profit')]:
                assert check.tables[0].cell(row,i).text == rec['fields'][key]['display_wan']
        for source, old_hash in zip(sources, original_hashes):
            if digest(source['path']) != old_hash:
                raise ValueError('Source file changed during run')
        (out/'extraction.json').write_text(json.dumps({'records':records, 'source_sha256':original_hashes, 'source_unchanged':True}, ensure_ascii=False, indent=2), encoding='utf-8')
        stage('build_docx_and_validate')
        pdf = out/'financial_brief.pdf'
        result = subprocess.run(['cscript.exe','//nologo',str(SKILL/'scripts'/'export_word_pdf.vbs'),str(output),str(pdf)], capture_output=True, timeout=60)
        (out/'word_export.log').write_bytes(result.stdout + result.stderr)
        if result.returncode or not pdf.exists():
            raise RuntimeError('Word export failed; see word_export.log')
        stage('word_export_pdf')
        poppler = cfg.get('poppler')
        if poppler:
            subprocess.run([poppler,'-png','-r','150',str(pdf),str(out/'page')],check=True,capture_output=True,timeout=60)
        else:
            from pypdfium2 import PdfDocument

            document = PdfDocument(str(pdf))
            try:
                for index in range(len(document)):
                    page = document[index]
                    try:
                        image = page.render(scale=150 / 72).to_pil()
                        image.save(out / f'page-{index + 1}.png')
                    finally:
                        page.close()
            finally:
                document.close()
        pages = list(out.glob('page-*.png'))
        if len(pages) != 1:
            raise ValueError(f'Expected single page, found {len(pages)}')
        stage('render_png')
        timings['page_count'] = len(pages)
        timings['status'] = 'machine_checks_passed_visual_pending'
    except Exception as error:
        timings['status'] = 'failed'
        timings['error'] = str(error)
        raise
    finally:
        timings['machine_seconds'] = round(time.perf_counter()-start,4)
        timings['machine_finished_utc'] = datetime.now(timezone.utc).isoformat()
        (out/'timing.json').write_text(json.dumps(timings, ensure_ascii=False, indent=2),encoding='utf-8')
        print(json.dumps(timings,ensure_ascii=True))

if __name__ == '__main__':
    main()
