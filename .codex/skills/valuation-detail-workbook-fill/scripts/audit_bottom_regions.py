from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import get_column_letter


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def build_bottom_region_cells(ws, start_row: int, end_row: int) -> list[str]:
    cells: list[str] = []
    for r in range(start_row, end_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            if isinstance(cell, MergedCell):
                continue
            cells.append(f"{get_column_letter(c)}{r}")
    return cells


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--layout-map", required=True)
    parser.add_argument("--summary-registry", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    wb = load_workbook(args.template, data_only=False)
    layout_map = load_json(Path(args.layout_map))
    summary_registry = load_json(Path(args.summary_registry))

    audit = {"template": args.template, "sheet_count": len(wb.sheetnames), "sheets": {}}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        layout_meta = (layout_map.get("sheets") or {}).get(sheet_name, {})
        summary_meta = (summary_registry.get("sheets") or {}).get(sheet_name, {})
        summary_zone = summary_meta.get("summary_zone") or layout_meta.get("summary_zone") or {}
        if not summary_zone:
            audit["sheets"][sheet_name] = {
                "chain_role": summary_meta.get("chain_role", "untracked"),
                "bottom_region_present": False,
            }
            continue
        start_row = int(summary_zone["start_row"])
        end_row = int(summary_zone["end_row"])
        bottom_cells = build_bottom_region_cells(ws, start_row, end_row)
        formula_cells = set(summary_meta.get("formula_output_cells", []))
        protected_text_ranges = summary_meta.get("protected_text_ranges", [])
        legal_input_cells = set(summary_meta.get("legal_input_cells", []))
        forbidden_non_formula = set(summary_meta.get("forbidden_non_formula_cells", []))

        formula_output_cells = [cell for cell in bottom_cells if cell in formula_cells]
        legal_inputs = [cell for cell in bottom_cells if cell in legal_input_cells]
        forbidden_non_formula_cells = [cell for cell in bottom_cells if cell in forbidden_non_formula]

        review_needed_cells = [
            cell
            for cell in bottom_cells
            if cell not in formula_cells and cell not in legal_input_cells and cell not in forbidden_non_formula
        ]

        audit["sheets"][sheet_name] = {
            "chain_role": summary_meta.get("chain_role", "untracked"),
            "bottom_region_present": True,
            "summary_zone": summary_zone,
            "formula_output_cells": formula_output_cells,
            "legal_input_cells": legal_inputs,
            "forbidden_non_formula_cells": forbidden_non_formula_cells,
            "review_needed_cells": review_needed_cells,
            "protected_text_ranges": protected_text_ranges,
        }

    wb.close()
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(out.as_posix())


if __name__ == "__main__":
    main()
