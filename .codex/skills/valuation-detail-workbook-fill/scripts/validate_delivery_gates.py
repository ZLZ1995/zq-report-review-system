from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from workflow_contract import validate_delivery_gate_report


DATE_HEADER_NAMES = {"发生日期"}
COUNTERPARTY_HEADER_NAMES = {"结算对象", "欠款单位名称", "往来单位", "预付款单位名称", "债权人名称"}
SIX_COUNTERPARTY_SHEETS = {"应收账款", "预付账款", "其他应收款", "应付账款", "预收账款", "其他应付款"}
FORBIDDEN_PLACEHOLDER_TERMS = {
    "报表差额占位",
    "报表差额",
    "其余明细合计",
    "详见缺资料清单",
    "税费调整项",
}
COUNTERPARTY_FORBIDDEN_TERMS = {
    "技术服务收入",
    "押金",
    "待抵扣进项税",
    "内部资金拆借",
    "应交所得税",
    "综合本位币",
    "人民币",
}
FIXED_FOOTER_TERMS = {"合计", "坏账准备", "评估风险损失", "净额"}
DATE_VALUE_RE = re.compile(r"^\d{4}/\d{2}/\d{2}$")
DATE_LIKE_RE = re.compile(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}$")


def clean(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def compact(value: Any) -> str:
    return clean(value).replace(" ", "")


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def xlsx_zip_valid(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as archive:
            return archive.testzip() is None
    except Exception:
        return False


def header_positions(ws) -> dict[str, int]:
    positions: dict[str, int] = {}
    for row_idx in range(1, min(ws.max_row, 10) + 1):
        for col_idx in range(1, ws.max_column + 1):
            text = clean(ws.cell(row_idx, col_idx).value)
            if text and text not in positions:
                positions[text] = col_idx
    return positions


def find_total_row(ws) -> int | None:
    for row_idx in range(1, ws.max_row + 1):
        row_text = "".join(clean(ws.cell(row_idx, col).value) for col in range(1, min(ws.max_column, 12) + 1))
        if "合计" in row_text.replace(" ", ""):
            return row_idx
    return None


def scan_workbook(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, data_only=False)
    bad_date_cells = []
    forbidden_placeholder_cells = []
    semantic_forbidden_cells = []
    bank_account_non_digit_cells = []
    broken_footer_sheets = []
    for ws in wb.worksheets:
        headers = header_positions(ws)
        date_cols = [col for name, col in headers.items() if name in DATE_HEADER_NAMES]
        counterparty_cols = {
            col for name, col in headers.items() if name in COUNTERPARTY_HEADER_NAMES
        } if ws.title in SIX_COUNTERPARTY_SHEETS else set()
        for row in ws.iter_rows():
            for cell in row:
                text = clean(cell.value)
                if not text:
                    continue
                if any(term in text for term in FORBIDDEN_PLACEHOLDER_TERMS):
                    forbidden_placeholder_cells.append({"sheet": ws.title, "cell": cell.coordinate, "value": text})
                if ws.title == "银行存款" and text == "人民币":
                    continue
                if cell.column in counterparty_cols and any(term == text for term in COUNTERPARTY_FORBIDDEN_TERMS):
                    semantic_forbidden_cells.append({"sheet": ws.title, "cell": cell.coordinate, "value": text})
        for date_col in date_cols:
            for row_idx in range(1, ws.max_row + 1):
                text = clean(ws.cell(row_idx, date_col).value)
                if DATE_LIKE_RE.match(text) and not DATE_VALUE_RE.match(text):
                    bad_date_cells.append({"sheet": ws.title, "cell": ws.cell(row_idx, date_col).coordinate, "value": text})
        if ws.title == "银行存款":
            for row in ws.iter_rows():
                for cell in row:
                    if cell.row <= 5:
                        continue
                    text = clean(cell.value)
                    if "账号" in clean(ws.cell(5, cell.column).value) and text and not text.isdigit():
                        bank_account_non_digit_cells.append({"sheet": ws.title, "cell": cell.coordinate, "value": text})
        if ws.title in {"应收账款", "其他应收款"}:
            total_row = find_total_row(ws)
            if total_row is not None:
                tail_text = "".join(
                    compact(ws.cell(row_idx, col).value)
                    for row_idx in range(total_row, min(ws.max_row, total_row + 8) + 1)
                    for col in range(1, min(ws.max_column, 12) + 1)
                )
                missing = sorted(term for term in FIXED_FOOTER_TERMS if compact(term) not in tail_text)
                if missing:
                    broken_footer_sheets.append({"sheet": ws.title, "total_row": total_row, "missing_terms": missing})
    return {
        "sheet_count": len(wb.worksheets),
        "bad_date_cells": bad_date_cells,
        "forbidden_placeholder_cells": forbidden_placeholder_cells,
        "semantic_forbidden_cells": semantic_forbidden_cells,
        "bank_account_non_digit_cells": bank_account_non_digit_cells,
        "broken_footer_sheets": broken_footer_sheets,
    }


def load_optional_json(path: Path | None) -> dict[str, Any]:
    if not path:
        return {}
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def build_report(args: argparse.Namespace) -> dict[str, Any]:
    workbook = Path(args.workbook)
    scan = scan_workbook(workbook)
    validation = load_optional_json(Path(args.validation_report) if args.validation_report else None)
    semantic = load_optional_json(Path(args.semantic_validation_report) if args.semantic_validation_report else None)
    report = {
        "workbook": str(workbook),
        "xlsx_zip_valid": xlsx_zip_valid(workbook),
        "balance_sheet_balanced": bool(validation.get("assets_equal_liabilities_equity", args.assume_balanced)),
        "classification_j4": validation.get("classification_j4", args.classification_j4 or ""),
        "classification_difference_count": len(validation.get("difference_hits", [])),
        "semantic_failure_count": int(semantic.get("failure_count", 0) or 0) + len(scan["semantic_forbidden_cells"]),
        "six_counterparty_semantic_anomaly_count": int(semantic.get("six_counterparty_semantic_anomaly_count", 0) or 0),
        "six_counterparty_semantic_clear": bool(semantic.get("six_counterparty_semantic_clear", False)),
        "bad_date_format_count": len(scan["bad_date_cells"]),
        "forbidden_placeholder_count": len(scan["forbidden_placeholder_cells"]),
        "bank_account_non_digit_count": len(scan["bank_account_non_digit_cells"]),
        "broken_footer_count": len(scan["broken_footer_sheets"]),
        "scan": scan,
    }
    report["gate_issues"] = validate_delivery_gate_report(report)
    report["status"] = "pass" if not report["gate_issues"] else "fail"
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate valuation detail workbook P0 delivery gates.")
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--validation-report", default="")
    parser.add_argument("--semantic-validation-report", default="")
    parser.add_argument("--classification-j4", default="")
    parser.add_argument("--assume-balanced", action="store_true")
    args = parser.parse_args()
    report = build_report(args)
    write_json(Path(args.output), report)
    print(json.dumps({"status": report["status"], "gate_issue_count": len(report["gate_issues"])}, ensure_ascii=False))
    return 0 if report["status"] == "pass" else 2


if __name__ == "__main__":
    raise SystemExit(main())
