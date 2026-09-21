"""K08: writer/reviewer consistency fixes driven by the A8T real-file retest.

Covers the remaining validation-failure classes found after the journal
parsing and candidate-unification fixes:

1. group_rows must not clobber a VALID TB counterparty with the dominant
   journal entity of the same account (淘宝/阿里上海 were being relabeled
   天猫), while a forbidden placeholder (默认值) may be resolved from
   journal evidence through the shared resolve_tb_counterparty rule.
2. The debit-balance suspense row 2241990000 (待查资金入账) must keep its
   net negative amount so the 其他应付款 detail reconciles to the BS.
3. stage2 must write the BS-sourced 其他流动负债 (预提费用) detail row so
   the 分类汇总 chain reconciles.
4. Locked-template writes must record the cells actually written so the
   reviewer never demands values from cells the lock made unwritable.
5. The semantic placeholder gate keeps the existing small-balance
   tolerance for every mandatory detail sheet, not just 其他应付款.
"""
import importlib.util
import sys
import unittest
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src'))

from openpyxl import Workbook

SCRIPTS = ROOT / '.codex' / 'skills' / 'valuation-detail-workbook-fill' / 'scripts'


def load_pipeline():
    sys.path.insert(0, str(SCRIPTS))
    spec = importlib.util.spec_from_file_location(
        'run_detail_workbook_pipeline', SCRIPTS / 'run_detail_workbook_pipeline.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def tb_row(code, name, aux, debit_end=0.0, credit_end=0.0):
    return {'tb_code': code, 'tb_account_name': name, 'aux_name': aux,
            'debit_end': debit_end, 'credit_end': credit_end, 'source_row': 1}


def mapping_entry(code, name, aux, sheet):
    return {'tb_code': code, 'tb_account_name': name, 'aux_name': aux,
            'target_sheet': sheet, 'detail_policy': 'detail_fillable'}


def journal_match(entity, debit=0.0, credit=0.0):
    return {'entity': entity, 'summary': 'S', 'business_desc': 'B',
            'line_desc': 'L', 'gl_date': datetime(2026, 1, 1),
            'debit': debit, 'credit': credit}


class ResolveTbCounterpartyTest(unittest.TestCase):
    def test_valid_tb_name_wins_over_dominant_journal_entity(self):
        p = load_pipeline()
        index = {'1124050000': [journal_match('浙江天猫技术有限公司', credit=99999)]}
        got = p.resolve_tb_counterparty('1124050000', '淘宝（中国）软件有限公司', index)
        self.assertEqual('淘宝（中国）软件有限公司', got)

    def test_forbidden_name_resolved_from_journal(self):
        p = load_pipeline()
        index = {'2202010000': [journal_match('安永(中国)企业咨询有限公司', credit=10600)]}
        got = p.resolve_tb_counterparty('2202010000', '默认值', index)
        self.assertEqual('安永(中国)企业咨询有限公司', got)

    def test_suspense_counterparty_allowed_by_exception_list(self):
        p = load_pipeline()
        got = p.resolve_tb_counterparty('2241990000', '其他应付-其他-待查资金入账-总账', {})
        self.assertEqual('待查资金入账', got)

    def test_forbidden_name_without_journal_evidence_stays_unresolved(self):
        p = load_pipeline()
        self.assertEqual('', p.resolve_tb_counterparty('2203040000', '非银行利息收入-集团内借贷', {}))


class GroupRowsGuardTest(unittest.TestCase):
    def test_valid_aux_name_not_relabeled(self):
        p = load_pipeline()
        mapping = {'account_mappings': [
            mapping_entry('1124050000', '内部往来-集团内公司借款（本金）', '淘宝（中国）软件有限公司', '其他应付款')]}
        tb = [tb_row('1124050000', '内部往来-集团内公司借款（本金）', '淘宝（中国）软件有限公司', credit_end=34389.85)]
        index = {'1124050000': [journal_match('浙江天猫技术有限公司', credit=13550942.26)]}
        plan = {sheet: rows for sheet, rows, *_ in p.group_rows_for_y71(mapping, tb, {'其他应付款': 34389.85}, index, {})}
        rows = plan['其他应付款']
        names = [r['counterparty'] for r in rows]
        self.assertIn('淘宝（中国）软件有限公司', names)
        self.assertNotIn('浙江天猫技术有限公司', names)

    def test_suspense_row_keeps_net_negative_amount(self):
        p = load_pipeline()
        mapping = {'account_mappings': [
            mapping_entry('2241990000', '其他应付款-待查资金入账', '其他应付-其他-待查资金入账-总账', '其他应付款')]}
        tb = [tb_row('2241990000', '其他应付款-待查资金入账', '其他应付-其他-待查资金入账-总账', debit_end=32537.99)]
        plan = {sheet: rows for sheet, rows, *_ in p.group_rows_for_y71(mapping, tb, {'其他应付款': -32537.99}, {}, {})}
        rows = plan['其他应付款']
        self.assertEqual(1, len(rows), rows)
        self.assertAlmostEqual(-32537.99, rows[0]['book_value'])


class Stage2OtherCurrentLiabilitiesTest(unittest.TestCase):
    def test_accrued_expense_row_written_from_bs(self):
        p = load_pipeline()
        wb = Workbook()
        ws = wb.active
        ws.title = '其他流动负债'
        ws['H6'] = '=G6'
        bs = {'values': {'其他流动负债': 109884.54}, 'values_prior': {}}
        writes = p.stage2_postfix_key_sheets(wb, bs)
        self.assertEqual('预提费用', ws['B6'].value)
        self.assertEqual(109884.54, ws['G6'].value)
        self.assertEqual('=G6', ws['H6'].value, '公式单元格不得被静态值覆盖')
        cells = {w['cell'] for w in writes}
        self.assertIn('G6', cells)
        self.assertNotIn('H6', cells)


class LockedWriteCellTrackingTest(unittest.TestCase):
    def test_written_record_lists_only_actually_written_cells(self):
        p = load_pipeline()
        wb = Workbook()
        ws = wb.active
        ws.title = '预收账款'
        for col in 'ABCDEFGHI':
            ws[f'{col}5'] = '表头'
        ws['A27'] = '合计'
        ws._locked_template_layout = True
        confirmed = [f'{col}{row}' for col in 'ABCDEFGI' for row in range(6, 27)]
        registry = {'selected_sheets': {'预收账款': {
            'chain_role': 'detail',
            'confirmed_input_cells': confirmed,
            'detail_body_cells': list(confirmed),
            'forbidden_non_formula_cells': [],
            'role_rules': {'allow_title_inputs': False, 'allow_detail_body_inputs': True,
                           'allow_balance_sheet_inputs': False, 'allow_summary_inputs': False},
            'formula_cells': [],
        }}}
        protection = {'sheets': {'预收账款': {'formula_cells': []}}}
        rows = [{'counterparty': '', 'book_value': -0.32, 'sub_name': '短期预收账款/递延收入-销项税暂估',
                 'remark': '', 'fill_mode': 'placeholder_only'}]
        written = p.write_simple_detail_sheet(ws, rows, 'detail', protection, registry, [])
        self.assertEqual(1, len(written))
        cells = set(written[0].get('_written_cells', set()))
        self.assertIn('G6', cells)
        self.assertNotIn('H6', cells, '锁定模板未确认列不得计入已写记录')


class SemanticToleranceTest(unittest.TestCase):
    def test_small_balance_placeholder_tolerated_on_any_mandatory_sheet(self):
        import tempfile
        p = load_pipeline()
        wb = Workbook()
        ws = wb.active
        ws.title = '预收账款'
        ws['B6'] = ''
        ws['G6'] = -0.32
        tmp = Path(tempfile.mkdtemp(prefix='k08_sem_')) / 'wb.xlsx'
        wb.save(tmp)
        report = p.build_semantic_validation_report(
            tmp, [], [], {'anomalies': []},
            [{'_sheet': '预收账款', 'fill_mode': 'placeholder_only', 'book_value': -0.32}],
            {'预收账款': -0.32})
        reasons = [f['reason'] for f in report['failures']]
        self.assertNotIn('placeholder_only_not_acceptable_for_detail_sheet', reasons)


if __name__ == '__main__':
    unittest.main()
