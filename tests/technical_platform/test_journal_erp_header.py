"""K08: journal loader must recognise Oracle-ERP TBD exports.

The A8T journal (序时账) is an Oracle GL export with title rows, the
header on row 3 and 52 columns keyed by names such as 会计科目代码 /
本币借项发生额 / 供应商名称 / 往来描述. The xlsx loader only probed for
the Kingdee-style header (科目编码/科目名称/方向/金额) and silently fell
back to fixed letter defaults, producing garbage rows whose tb_code was
"日记帐摘要" and whose counterparty fields were empty — every downstream
journal match then failed against entries that actually exist.

The ERP carries two legitimate counterparty columns per line (供应商名称
from the AP subledger, 往来描述 from the intercompany segment). The
pipeline invariant "vendor_name == counterparty_desc" was written when
both fields came from a single column, so the loader must pick ONE
evidence-backed value per row: 内部往来 accounts key on 往来描述, all
other accounts on 供应商名称, each falling back to the other column when
its own value is missing or a forbidden placeholder such as 默认值.
"""
import importlib.util
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'src'))

from openpyxl import Workbook

SCRIPTS = ROOT / '.codex' / 'skills' / 'valuation-detail-workbook-fill' / 'scripts'


def load_pipeline():
    sys.path.insert(0, str(SCRIPTS))
    spec = importlib.util.spec_from_file_location(
        'run_detail_workbook_pipeline', SCRIPTS / 'run_detail_workbook_pipeline.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def make_erp_tbd(path):
    """Minimal Oracle-ERP TBD layout: 2 title rows, header on row 3,
    key columns at their real letter positions."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet0'
    ws['A1'] = '总帐明细帐追溯报表'
    ws['A2'] = '公司：A8T-A8T帐套：PRCGAAP(CNY)'
    header = {
        'A': '公司代码', 'B': '期间', 'C': 'GL日期', 'K': '日记帐名称',
        'L': '日记帐摘要', 'N': '日记帐行说明',
        'U': '本币借项发生额', 'V': '本币贷项发生额',
        'Z': '供应商名称', 'AF': '会计科目代码', 'AR': '会计科目描述',
        'AT': '往来描述',
    }
    for col, name in header.items():
        ws[f'{col}3'] = name
    # 内部往来行：供应商名称是外部供应商，往来描述才是集团内客商
    ws['C4'] = '2026-01-08'
    ws['L4'] = 'AP_CLEAR'
    ws['N4'] = 'PO8822931449'
    ws['U4'] = '6757.25'
    ws['V4'] = '0'
    ws['Z4'] = '中国平安财产保险股份有限公司浙江分公司'
    ws['AF4'] = '1124020000'
    ws['AR4'] = '内部往来-代垫/代收款项'
    ws['AT4'] = '杭州阿里云飞天信息技术有限公司'
    # 普通负债行：往来描述是禁用占位词，供应商名称（员工）兜底
    ws['C5'] = '2026-01-20'
    ws['L5'] = 'PREPAYMENT'
    ws['U5'] = '0'
    ws['V5'] = '264'
    ws['Z5'] = '管仁良'
    ws['AF5'] = '2241120000'
    ws['AR5'] = '其他应付款-员工款'
    ws['AT5'] = '默认值'
    # 外部供应商行：供应商名称是真实客商，往来描述是集团段杂项
    ws['C6'] = '2026-01-08'
    ws['L6'] = 'AP_CLEAR'
    ws['U6'] = '0'
    ws['V6'] = '6757.25'
    ws['Z6'] = '中国平安财产保险股份有限公司浙江分公司'
    ws['AF6'] = '2202010000'
    ws['AR6'] = '应付账款-外部供应商 （AP模块）'
    ws['AT6'] = '阿里巴巴集團其他少數股權投資'
    wb.save(path)


class TestErpJournalHeader(unittest.TestCase):
    def setUp(self):
        import tempfile
        self.tmp = Path(tempfile.mkdtemp(prefix='k08_journal_')) / 'tbd.xlsx'
        make_erp_tbd(self.tmp)
        self.pipeline = load_pipeline()
        self.rows = self.pipeline.load_journal_rows_from_xlsx_xml(self.tmp)

    def test_erp_layout_columns_mapped_by_name(self):
        rows = self.rows
        self.assertEqual(len(rows), 3, rows)
        first = rows[0]
        self.assertEqual(first['tb_code'], '1124020000')
        self.assertEqual(first['account_name'], '内部往来-代垫/代收款项')
        self.assertAlmostEqual(first['debit'], 6757.25)
        self.assertAlmostEqual(first['credit'], 0.0)
        self.assertIn('2026-01-08', str(first['gl_date']))
        self.assertEqual(rows[1]['tb_code'], '2241120000')
        self.assertAlmostEqual(rows[1]['credit'], 264.0)

    def test_title_rows_not_parsed_as_entries(self):
        for row in self.rows:
            self.assertNotEqual(row['tb_code'], '日记帐摘要')
            self.assertNotEqual(row['tb_code'], '会计科目代码')

    def test_intercompany_rows_key_on_wanglai_desc(self):
        row = self.rows[0]
        self.assertEqual(row['vendor_name'], '杭州阿里云飞天信息技术有限公司')
        self.assertEqual(row['counterparty_desc'], '杭州阿里云飞天信息技术有限公司')

    def test_forbidden_wanglai_falls_back_to_vendor(self):
        row = self.rows[1]
        self.assertEqual(row['vendor_name'], '管仁良')
        self.assertEqual(row['counterparty_desc'], '管仁良')

    def test_external_ap_rows_key_on_vendor_name(self):
        row = self.rows[2]
        self.assertEqual(row['vendor_name'], '中国平安财产保险股份有限公司浙江分公司')
        self.assertEqual(row['counterparty_desc'], '中国平安财产保险股份有限公司浙江分公司')

    def test_entity_index_builds_without_conflict(self):
        index = self.pipeline.build_journal_entity_index(self.rows)
        entities = [e['entity'] for e in index['1124020000']]
        self.assertEqual(entities, ['杭州阿里云飞天信息技术有限公司'])


if __name__ == '__main__':
    unittest.main()
