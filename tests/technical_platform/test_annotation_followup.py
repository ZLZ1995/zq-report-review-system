from asset_based_agent.technical_platform.annotations import annotation_offer


def test_offer_requires_completed_review_with_issues():
    assert annotation_offer('succeeded', {'kind': 'review', 'issues': [{'description': '问题'}]})
    assert not annotation_offer('cancelled', {'kind': 'review', 'issues': [{}]})
    assert not annotation_offer('succeeded', {'kind': 'review', 'issues': []})
    assert not annotation_offer('succeeded', {'kind': 'preflight', 'issues': [{}]})


def test_docx_copy_keeps_original_and_unrelated_parts(tmp_path):
    from zipfile import ZipFile

    from docx import Document

    from asset_based_agent.technical_platform.annotations import annotate_docx
    from asset_based_agent.technical_platform.skills import digest
    source, output = tmp_path / 'original.docx', tmp_path / 'comments.docx'
    doc = Document()
    doc.add_paragraph('原文测试，金额单位需要核对。')
    doc.add_paragraph('正常内容不能修改。')
    doc.save(source)
    before = digest(source)
    applied, skipped = annotate_docx(source, output, [(1, {
        'original_text': '原文测试，金额单位需要核对。', 'description': '单位需核对',
        'requires_verification': True, 'recommendation': '核对来源单位'})])
    assert applied == [1] and skipped == []
    assert digest(source) == before
    assert [p.text for p in Document(output).paragraphs] == [p.text for p in doc.paragraphs]
    with ZipFile(output) as archive:
        assert '待核实' in archive.read('word/comments.xml').decode('utf-8')


def test_ambiguous_docx_does_not_create_output(tmp_path):
    from docx import Document

    from asset_based_agent.technical_platform.annotations import annotate_docx
    source, output = tmp_path / 'original.docx', tmp_path / 'comments.docx'
    doc = Document()
    doc.add_paragraph('重复内容')
    doc.add_paragraph('重复内容')
    doc.save(source)
    applied, skipped = annotate_docx(source, output, [(1, {'original_text': '重复内容'})])
    assert not applied and len(skipped) == 1
    assert not output.exists()


def test_excel_marks_only_visible_verified_cell_preserving_hidden_parts(tmp_path):
    from openpyxl import Workbook, load_workbook

    from asset_based_agent.technical_platform.annotations import annotate_excel
    source, output = tmp_path / 'original.xlsx', tmp_path / 'comments.xlsx'
    book = Workbook()
    book.active.title = '汇总表'
    book.active['A1'] = '100万元'
    hidden = book.create_sheet('隐藏表')
    hidden['A1'] = '秘密'
    hidden.sheet_state = 'hidden'
    book.save(source)
    applied, skipped = annotate_excel(source, output, [
        (1, {'location': {'table': '汇总表', 'cell': 'A1'}, 'original_text': '100万元', 'description': '核对单位'}),
        (2, {'location': {'table': '隐藏表', 'cell': 'A1'}, 'original_text': '秘密', 'description': '不得标注'})])
    assert applied == [1] and [x['issue'] for x in skipped] == [2]
    result = load_workbook(output)
    assert result['汇总表']['A1'].value == '100万元'
    assert result['汇总表']['A1'].comment is not None
    assert result['隐藏表']['A1'].comment is None


def test_annotation_question_once_and_decline_does_not_write(tmp_path, monkeypatch):
    import os
    os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtWidgets import QApplication, QMessageBox
    from asset_based_agent.technical_platform.app import PlatformWindow
    from asset_based_agent.technical_platform.store import PlatformStore
    app = QApplication.instance() or QApplication([])
    assert app
    store = PlatformStore(tmp_path / 'state.db', 'test')
    project = store.create_project('批注询问')
    session = store.create_session(project)
    run_id = store.start_run(session, {'files': []})
    store.claim_run(run_id)
    store.save_result(run_id, {'kind': 'review', 'issues': [{'description': '需要核对'}]})
    store.transition(run_id, 'validating', 'test')
    store.transition(run_id, 'succeeded', 'test')
    window = PlatformWindow(store)
    window.reload_projects(project)
    calls = []
    monkeypatch.setattr(QMessageBox, 'question', lambda *args: calls.append(True) or QMessageBox.StandardButton.No)
    try:
        window.offer_annotations(run_id)
        window.offer_annotations(run_id)
        assert calls == [True]
        assert '是否生成带问题标记和批注' in window.transcript.toPlainText()
        assert not list(tmp_path.glob('*.docx'))
        assert store.run(run_id)['state'] == 'succeeded'
    finally:
        window.close()
