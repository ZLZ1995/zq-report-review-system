"""Run the approved template with synthetic, balanced source data; never bypass gates."""
import hashlib
import json
import sys
from pathlib import Path
from tempfile import mkdtemp

from openpyxl import Workbook

from asset_based_agent.technical_platform.generation_worker import main


def run():
    root = Path(__file__).resolve().parents[1]
    frozen = '--exe' in sys.argv
    parent = root / 'outputs/detail_generation_acceptance'
    parent.mkdir(parents=True, exist_ok=True)
    work = Path(mkdtemp(prefix='run_', dir=parent))
    tb, bs = work / 'trial_balance.xlsx', work / 'balance_sheet.xlsx'
    journal = work / 'journal.xlsx'
    wb = Workbook()
    ws = wb.active
    ws['A1'], ws['B3'] = '科目汇总试算表', '科目代码'
    for row, code, account, party, debit, credit in [
        (5, '1122', '应收账款', '验收甲有限公司', 1000, 0),
        (6, '2202', '应付账款', '验收乙有限公司', 0, 1000),
    ]:
        for column, value in [('B', code), ('D', account), ('E', party), ('J', debit), ('K', credit)]:
            ws[f'{column}{row}'] = value
    wb.save(tb)
    wb.close()
    wb = Workbook()
    ws = wb.active
    ws.title = '资产负债表'
    ws['B2'], ws['B3'], ws['B5'] = '资产负债表', '2026年3月31日', '编制单位：合成验收有限公司'
    left = [('货币资金', 0), ('应收账款', 1000), ('预付账款', 0), ('其他应收款', 0),
            ('存货', 0), ('流动资产合计', 1000), ('固定资产', 0), ('非流动资产合计', 0), ('资产总计', 1000)]
    right = [('应付账款', 1000), ('预收账款', 0), ('其他应付款', 0), ('应交税费', 0),
             ('应付职工薪酬', 0), ('流动负债合计', 1000), ('非流动负债合计', 0), ('负债合计', 1000),
             ('实收资本', 0), ('所有者权益合计', 0), ('负债和所有者权益总计', 1000)]
    for row, (label, amount) in enumerate(left, 7):
        ws[f'B{row}'], ws[f'C{row}'], ws[f'D{row}'] = label, amount, amount
    for row, (label, amount) in enumerate(right, 7):
        ws[f'E{row}'], ws[f'F{row}'], ws[f'G{row}'] = label, amount, amount
    wb.save(bs)
    wb.close()
    wb = Workbook()
    wb.active.append(['业务日期', '摘要', '科目编码', '科目名称', '借方', '贷方', '往来单位'])
    wb.active.append(['2026/03/01', '销售服务', '1122', '应收账款', 1000, 0, '验收甲有限公司'])
    wb.active.append(['2026/03/01', '采购服务', '2202', '应付账款', 0, 1000, '验收乙有限公司'])
    wb.save(journal)
    wb.close()
    bank_inputs = {}
    if '--bank' in sys.argv:
        from openpyxl import load_workbook
        wb = load_workbook(tb)
        for address, value in {'B5': '1002', 'D5': '银行存款', 'E5': '测试银行',
                               'B6': '4001', 'D6': '实收资本', 'E6': '测试股东'}.items():
            wb.active[address] = value
        wb.save(tb)
        wb.close()
        wb = load_workbook(bs)
        for row, (label, _) in enumerate(left, 7):
            amount = 1000 if label in {'货币资金', '流动资产合计', '资产总计'} else 0
            wb.active[f'C{row}'] = wb.active[f'D{row}'] = amount
        for row, (label, _) in enumerate(right, 7):
            amount = 1000 if label in {'实收资本', '所有者权益合计', '负债和所有者权益总计'} else 0
            wb.active[f'F{row}'] = wb.active[f'G{row}'] = amount
        wb.save(bs)
        wb.close()
        bank = work / 'bank.xlsx'
        wb = Workbook()
        wb.active.append(['账号', '账户名称', '开户行', '交易日期', '账户余额'])
        wb.active.append(['1234567890123', '合成验收有限公司', '测试银行', '2026/03/31', 1000])
        wb.save(bank)
        wb.close()
        bank_inputs['bank_statement'] = str(bank)
    template = root / 'assets/builtin_templates/valuation-detail-workbook-fill/template.xlsx'
    paths = [tb, bs, journal, template]
    paths.extend(Path(path) for path in bank_inputs.values())
    before = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}
    job = work / 'job.json'
    job.write_text(json.dumps({'skill_id': 'valuation-detail-workbook-fill', 'inputs': {
        'trial_balance': str(tb), 'balance_sheet': str(bs), 'journal': str(journal),
        'template': str(template), **bank_inputs}}, ensure_ascii=False), 'utf-8')
    if '--wps' in sys.argv:
        # Force the WPS backend for acceptance only; production stays auto-select.
        sys.path.insert(0, str(root / '.codex/skills/valuation-detail-workbook-fill/scripts'))
        import recalculate_readonly
        import win32com.client
        factory = recalculate_readonly.create_calculation_application
        def dispatch(progid):
            if progid == 'Excel.Application':
                raise OSError('WPS-only acceptance')
            return win32com.client.DispatchEx(progid)
        recalculate_readonly.create_calculation_application = lambda: factory(dispatch)
    print(f'Acceptance directory: {work}', flush=True)
    if frozen:
        import subprocess
        executable = root / 'dist/technical_platform/ZQ技术平台/ZQ技术平台.exe'
        code = subprocess.run([str(executable), '--builtin-skill-worker', str(job)],
                              cwd=executable.parent, timeout=600, check=False).returncode
    else:
        code = main(job)
    assert before == {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}
    print(f'Worker exit: {code}; sources and template unchanged', flush=True)
    return code


if __name__ == '__main__':
    raise SystemExit(run())
