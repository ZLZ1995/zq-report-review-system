"""Generate the final deduplicated Word review report."""

from __future__ import annotations

import json
import os
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from ..domain.enums import IssueStatus
from ..domain.models import AuditProject, ReviewIssue, utc_now
from ..repositories.issue_repository import IssueRepository

STATUS_TEXT = {
    IssueStatus.NEW: "本轮新发现，尚未确认修改",
    IssueStatus.UNMODIFIED: "未修改",
    IssueStatus.FIXED: "已完成修改",
    IssueStatus.INCORRECT_FIX: "已修改，但修改后仍存在错误",
    IssueStatus.UNCERTAIN: "无法自动判断，需人工复核",
    IssueStatus.IGNORED: "已由用户忽略",
}


class ReportExportService:
    def __init__(self, issue_repository: IssueRepository | None = None) -> None:
        self.issue_repository = issue_repository or IssueRepository()

    def export(
        self,
        project: AuditProject,
        destination: Path,
    ) -> tuple[Path, Path]:
        latest_round = self._latest_completed_round(project)
        issues = self.issue_repository.load_issues(Path(latest_round.issues_path))
        issues = _deduplicate(issues)
        report_path = Path(destination)
        if report_path.suffix.lower() != ".docx":
            report_path = report_path.with_suffix(".docx")
        report_path.parent.mkdir(parents=True, exist_ok=True)
        summary_path = report_path.with_suffix(".summary.json")
        summary = self._summary(project, latest_round.round_number, issues)
        _write_json_atomic(summary_path, summary)
        self._write_docx(report_path, project, summary, issues)
        project.final_report_path = str(report_path)
        return summary_path, report_path

    @staticmethod
    def _latest_completed_round(project: AuditProject):
        completed = [
            item for item in project.rounds if item.status.value == "completed"
        ]
        if not completed:
            raise ValueError("project has no completed audit round")
        return max(completed, key=lambda item: item.round_number)

    @staticmethod
    def _summary(
        project: AuditProject,
        round_number: int,
        issues: list[ReviewIssue],
    ) -> dict[str, Any]:
        counts = Counter(issue.status.value for issue in issues)
        latest_files = [
            item
            for item in project.files
            if item.round_number == project.current_round
        ]
        return {
            "schema_version": "1.0",
            "project_id": project.project_id,
            "project_name": project.name,
            "completed_rounds": round_number,
            "generated_at": utc_now().isoformat(),
            "files": [
                {
                    "file_id": item.file_id,
                    "name": item.original_name,
                    "role": item.role.value,
                }
                for item in latest_files
            ],
            "issue_count": len(issues),
            "status_counts": dict(counts),
            "issues": [issue.model_dump(mode="json") for issue in issues],
        }

    def _write_docx(
        self,
        path: Path,
        project: AuditProject,
        summary: dict[str, Any],
        issues: list[ReviewIssue],
    ) -> None:
        document = Document()
        _configure_styles(document)
        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("资产评估报告审核记录")
        run.bold = True
        run.font.size = Pt(20)

        document.add_heading("一、项目基本信息", level=1)
        document.add_paragraph(f"项目名称：{project.name}")
        document.add_paragraph(f"审核轮次：{summary['completed_rounds']}")
        document.add_paragraph(f"问题总数：{summary['issue_count']}")

        document.add_heading("二、审核文件清单", level=1)
        files = summary["files"]
        if files:
            table = document.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            for cell, value in zip(table.rows[0].cells, ("序号", "文件名", "文件角色")):
                cell.text = value
            for index, item in enumerate(files, start=1):
                cells = table.add_row().cells
                cells[0].text = str(index)
                cells[1].text = str(item["name"])
                cells[2].text = str(item["role"])
        else:
            document.add_paragraph("无当前轮次文件记录。")

        document.add_heading("三、审核总体情况", level=1)
        counts = summary["status_counts"]
        for status, label in STATUS_TEXT.items():
            document.add_paragraph(f"{label}：{counts.get(status.value, 0)}项")

        document.add_heading("四、各文件审核意见", level=1)
        grouped: dict[str, list[ReviewIssue]] = defaultdict(list)
        for issue in issues:
            if issue.status != IssueStatus.IGNORED:
                grouped[issue.source_file_name].append(issue)
        if not grouped:
            document.add_paragraph("无需要列示的审核意见。")
        for file_name, file_issues in grouped.items():
            document.add_heading(file_name, level=2)
            for index, issue in enumerate(file_issues, start=1):
                self._add_issue(document, index, issue)

        unresolved = [
            issue
            for issue in issues
            if issue.status
            in {
                IssueStatus.NEW,
                IssueStatus.UNMODIFIED,
                IssueStatus.INCORRECT_FIX,
                IssueStatus.UNCERTAIN,
            }
        ]
        document.add_heading("五、尚未解决的问题", level=1)
        if unresolved:
            for index, issue in enumerate(unresolved, start=1):
                document.add_paragraph(
                    f"{index}. {issue.source_file_name}：{issue.description}"
                    f"（{STATUS_TEXT[issue.status]}）"
                )
        else:
            document.add_paragraph("无。")

        ignored = [issue for issue in issues if issue.status == IssueStatus.IGNORED]
        document.add_heading("六、已忽略问题汇总", level=1)
        if ignored:
            for index, issue in enumerate(ignored, start=1):
                document.add_paragraph(
                    f"{index}. {issue.source_file_name}：{issue.description}"
                )
        else:
            document.add_paragraph("无。")
        document.save(str(path))

    @staticmethod
    def _add_issue(document: DocxDocument, index: int, issue: ReviewIssue) -> None:
        document.add_paragraph(issue.description, style="List Number")
        document.add_paragraph(f"问题位置：{_location_text(issue)}")
        document.add_paragraph(
            f"问题来源：{issue.origin or 'legacy'}；"
            f"证据状态：{issue.evidence_state or 'unknown'}"
        )
        if issue.original_text:
            document.add_paragraph(f"原文或原值：{issue.original_text}")
        document.add_paragraph(f"最终状态：{STATUS_TEXT[issue.status]}")
        if issue.recommendation:
            document.add_paragraph(f"处理建议：{issue.recommendation}")
        document.add_paragraph(f"处理情况：{_history_summary(issue)}")


def _configure_styles(document: DocxDocument) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def _location_text(issue: ReviewIssue) -> str:
    locations = issue.occurrences or [issue.location]
    return "；".join(
        f"位置{index}：{_single_location_text(location)}"
        for index, location in enumerate(locations, start=1)
    )


def _single_location_text(location) -> str:
    values = []
    if location.chapter:
        values.append(location.chapter)
    if location.page:
        values.append(f"第{location.page}页")
    if location.paragraph:
        values.append(f"第{location.paragraph}段")
    if location.table:
        values.append(f"表格/工作表：{location.table}")
    if location.cell:
        values.append(location.cell)
    return "；".join(values) or "待人工确认"


def _history_summary(issue: ReviewIssue) -> str:
    if not issue.history:
        return STATUS_TEXT[issue.status]
    entries = []
    seen: set[tuple[int, str]] = set()
    for item in issue.history:
        key = (item.round_number, item.status.value)
        if key in seen:
            continue
        seen.add(key)
        entries.append(f"第{item.round_number}轮：{STATUS_TEXT[item.status]}")
    return "；".join(entries)


def _deduplicate(issues: list[ReviewIssue]) -> list[ReviewIssue]:
    by_id: dict[str, ReviewIssue] = {}
    for issue in issues:
        existing = by_id.get(issue.issue_id)
        if existing is None or issue.last_seen_round >= existing.last_seen_round:
            by_id[issue.issue_id] = issue
    return list(by_id.values())


def _write_json_atomic(path: Path, payload: dict[str, Any]) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    os.replace(temporary, path)
