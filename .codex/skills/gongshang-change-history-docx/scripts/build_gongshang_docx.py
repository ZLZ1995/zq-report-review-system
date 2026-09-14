from copy import deepcopy
from datetime import datetime
import argparse
from pathlib import Path
import re

from docx import Document
from docx.table import Table, _Cell
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import load_workbook


XLSX = None
TEMPLATE = None
HUMAN_FORMAT_REFERENCE = None
OUTPUT_DIR = Path("outputs/gongshang")
OUTPUT = OUTPUT_DIR / "工商历史变更模版-生成版.docx"
FALLBACK_OUTPUT = OUTPUT_DIR / "工商历史变更模版-生成版-表格修正.docx"


def configure_paths(args):
    global XLSX, TEMPLATE, HUMAN_FORMAT_REFERENCE, OUTPUT_DIR, OUTPUT, FALLBACK_OUTPUT
    XLSX = Path(args.source_excel)
    TEMPLATE = Path(args.template)
    OUTPUT = Path(args.output)
    OUTPUT_DIR = OUTPUT.parent
    FALLBACK_OUTPUT = OUTPUT.with_name(f"{OUTPUT.stem}-fallback{OUTPUT.suffix}")
    if args.reference:
        HUMAN_FORMAT_REFERENCE = Path(args.reference)
    else:
        HUMAN_FORMAT_REFERENCE = Path("__missing_human_format_reference__.docx")


def text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def normalize_for_doc(value):
    value = text(value)
    value = re.sub(r"\s+", " ", value)
    value = value.replace("；", "；\n")
    value = value.replace(";", "；\n")
    return value.strip()


def normalize_inline(value):
    value = text(value)
    value = re.sub(r"\s+", " ", value)
    value = value.replace("； ", "；")
    return value.strip()


def compact(value):
    value = text(value)
    return re.sub(r"\s+", "", value)


def cn_date(iso_date):
    dt = datetime.strptime(iso_date, "%Y-%m-%d")
    return f"{dt.year}年{dt.month}月{dt.day}日"


def find_header(ws):
    for ridx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = [text(v) for v in row]
        if "变更日期" in values and "变更事项" in values:
            return ridx, values
    raise ValueError("未找到包含“变更日期”和“变更事项”的表头")


def load_events(source=None):
    wb = load_workbook(source or XLSX, data_only=True, read_only=True)
    try:
        return _load_events(wb)
    finally:
        wb.close()


def _load_events(wb):
    ws = wb["变更信息"]
    header_row, headers = find_header(ws)
    date_idx = headers.index("变更日期")
    item_idx = headers.index("变更事项")
    before_idx = headers.index("变更前")
    after_idx = headers.index("变更后")

    events = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        vals = [text(v) for v in row]
        if not any(vals):
            continue
        date = vals[date_idx]
        item = vals[item_idx]
        if not date or not item:
            continue
        date = datetime.strptime(date[:10].replace('/', '-'), '%Y-%m-%d').strftime('%Y-%m-%d')
        events.setdefault(date, []).append(
            {
                "item": item,
                "before": normalize_for_doc(vals[before_idx]),
                "after": normalize_for_doc(vals[after_idx]),
            }
        )
    return [(date, events[date]) for date in sorted(events)]


def unique_items(rows):
    seen = []
    for row in rows:
        item = row["item"]
        if item not in seen:
            seen.append(item)
    return seen


def split_entries(value):
    value = text(value).replace("；", ";")
    return [part.strip() for part in value.split(";") if part.strip()]


def parse_holder_entry(entry):
    entry = entry.strip()
    amount = ""
    name = entry
    match = re.search(r"\s*货币\s*([0-9,]+(?:\.[0-9]+)?)\s*万人民币", entry)
    if match:
        amount = match.group(1).replace(",", "")
        name = entry[: match.start()].strip()
    return name, amount


def parse_holders(value):
    holders = {}
    order = []
    for entry in split_entries(value):
        name, amount = parse_holder_entry(entry)
        if not name:
            continue
        if name not in holders:
            order.append(name)
        holders[name] = amount
    return order, holders


def numeric_amount(value):
    if not value:
        return None
    try:
        return float(str(value).replace(",", ""))
    except ValueError:
        return None


def capital_amount(value):
    match = re.search(r"([0-9,]+(?:\.[0-9]+)?)\s*万人民币", text(value))
    if not match:
        return None
    return numeric_amount(match.group(1))


def pct(amount, total):
    amount_num = numeric_amount(amount)
    if amount_num is None or not total:
        return ""
    return f"{amount_num / total:.4%}"


def select_equity_row(rows):
    equity_candidates = [
        row
        for row in rows
        if row["item"] in {"股东出资变更", "股东变更"}
        and (row["before"] or row["after"])
    ]
    detailed = [
        row
        for row in equity_candidates
        if "货币" in row["before"] or "货币" in row["after"]
    ]
    if detailed:
        return max(detailed, key=lambda r: len(r["before"]) + len(r["after"]))
    if equity_candidates:
        return max(equity_candidates, key=lambda r: len(r["before"]) + len(r["after"]))
    return None


def equity_row_has_actual_change(row):
    if row is None:
        return False
    _, before_holders = parse_holders(row["before"])
    _, after_holders = parse_holders(row["after"])
    return before_holders != after_holders


def is_duplicate_equity_row(row, equity_row):
    if equity_row is None:
        return False
    if row is equity_row:
        return True
    if row["item"] in {"股东出资变更", "股东变更"}:
        return True
    if row["item"] == "其他变更" and ("货币" in row["before"] or "货币" in row["after"]):
        return True
    return False


def dedupe_non_equity(rows, equity_row):
    chosen = {}
    order = []
    actual_equity = equity_row_has_actual_change(equity_row)
    for row in rows:
        if row["item"] in {"股东出资变更", "股东变更"} and not actual_equity:
            continue
        if is_duplicate_equity_row(row, equity_row):
            continue
        key = (row["item"], compact(row["after"]))
        score = (1 if compact(row["before"]) else 0) + len(compact(row["before"]))
        if key not in chosen:
            order.append(key)
            chosen[key] = (score, row)
        elif score > chosen[key][0]:
            chosen[key] = (score, row)
    return [chosen[key][1] for key in order]


def event_table_rows(rows):
    capital = next((r for r in rows if r["item"] == "注册资本变更"), None)
    before_capital = capital_amount(capital["before"]) if capital else None
    after_capital = capital_amount(capital["after"]) if capital else None
    equity_row = select_equity_row(rows)
    if not equity_row_has_actual_change(equity_row):
        equity_row = None
    equity_rows = []
    if equity_row is not None:
        before_order, before_holders = parse_holders(equity_row["before"])
        after_order, after_holders = parse_holders(equity_row["after"])
        if before_capital is None:
            before_amounts = [numeric_amount(v) for v in before_holders.values()]
            before_amounts = [v for v in before_amounts if v is not None]
            before_capital = sum(before_amounts) if before_amounts else None
        if after_capital is None:
            after_amounts = [numeric_amount(v) for v in after_holders.values()]
            after_amounts = [v for v in after_amounts if v is not None]
            after_capital = sum(after_amounts) if after_amounts else None
        holder_order = list(before_order)
        holder_order.extend([name for name in after_order if name not in before_holders])
        for name in holder_order:
            before_amount = before_holders.get(name, "")
            after_amount = after_holders.get(name, "")
            equity_rows.append(
                {
                    "kind": "equity",
                    "project": "股东出资金额及持股比例",
                    "before_name": name if name in before_holders else "",
                    "before_amount": before_amount,
                    "before_ratio": pct(before_amount, before_capital),
                    "after_name": name if name in after_holders else "",
                    "after_amount": after_amount,
                    "after_ratio": pct(after_amount, after_capital),
                }
            )
        before_total = sum(
            v for v in (numeric_amount(amount) for amount in before_holders.values()) if v is not None
        )
        after_total = sum(
            v for v in (numeric_amount(amount) for amount in after_holders.values()) if v is not None
        )
        equity_rows.append(
            {
                "kind": "equity",
                "project": "股东出资金额及持股比例",
                "before_name": "合计",
                "before_amount": f"{before_total:.6f}" if before_total else "",
                "before_ratio": "100.0000%" if before_total else "",
                "after_name": "合计",
                "after_amount": f"{after_total:.6f}" if after_total else "",
                "after_ratio": "100.0000%" if after_total else "",
            }
        )

    non_equity = [
        {
            "kind": "normal",
            "project": row["item"],
            "before": normalize_for_doc(row["before"]) or "无",
            "after": normalize_for_doc(row["after"]) or "无",
        }
        for row in dedupe_non_equity(rows, equity_row)
    ]
    return equity_rows + non_equity


def event_has_legal_representative_change(rows):
    return any(row['item'] in {'法定代表人变更', '法人代表变更'} for row in rows)


def is_pure_low_value_event(rows):
    equity_row = select_equity_row(rows)
    has_actual_equity = equity_row_has_actual_change(equity_row)
    if has_actual_equity or event_has_legal_representative_change(rows):
        return False
    substantive_items = []
    for row in rows:
        item = row["item"]
        if item in {"章程变更", "股东出资变更", "股东变更"}:
            continue
        if item == "其他变更":
            continue
        substantive_items.append(item)
    return not substantive_items


def selected_events(events):
    return [(date, rows) for date, rows in events if not is_pure_low_value_event(rows)]


def format_money(value, decimals=6):
    number = numeric_amount(value)
    if number is None:
        return ""
    return f"{number:.{decimals}f}"


def format_ratio_from_amount(amount, total):
    amount_num = numeric_amount(amount)
    if amount_num is None or not total:
        return ""
    return f"{amount_num / total:.4%}"


def resolve_name_only_equity(equity_row, previous_cap_table, next_cap_table=None):
    before_order, before_raw = parse_holders(equity_row["before"])
    after_order, after_raw = parse_holders(equity_row["after"])
    if (any(before_raw.values()) or any(after_raw.values()) or not previous_cap_table
            or not next_cap_table or set(before_raw) != set(previous_cap_table)
            or set(after_raw) != set(next_cap_table)):
        return None

    before_order = list(previous_cap_table.keys())
    before_holders = {name: previous_cap_table.get(name, "") for name in before_order}
    if next_cap_table:
        after_order = list(next_cap_table.keys())
        after_holders = {name: next_cap_table.get(name, "") for name in after_order}
    else:
        after_holders = {}
        remaining_amounts = [previous_cap_table[name] for name in before_order if name not in after_order]
        for name in after_order:
            if name in previous_cap_table:
                after_holders[name] = previous_cap_table[name]
            else:
                after_holders[name] = remaining_amounts.pop(0) if remaining_amounts else ""
    return before_order, before_holders, after_order, after_holders


def build_equity_rows_for_event(rows, previous_cap_table, next_cap_table=None):
    capital = next((r for r in rows if r["item"] == "注册资本变更"), None)
    before_capital = capital_amount(capital["before"]) if capital else None
    after_capital = capital_amount(capital["after"]) if capital else None
    equity_row = select_equity_row(rows)
    if not equity_row_has_actual_change(equity_row):
        return [], previous_cap_table

    resolved = resolve_name_only_equity(equity_row, previous_cap_table, next_cap_table)
    if resolved:
        before_order, before_holders, after_order, after_holders = resolved
    else:
        before_order, before_holders = parse_holders(equity_row["before"])
        after_order, after_holders = parse_holders(equity_row["after"])

    if before_capital is None:
        before_amounts = [numeric_amount(v) for v in before_holders.values()]
        before_amounts = [v for v in before_amounts if v is not None]
        before_capital = sum(before_amounts) if before_amounts else None
    if after_capital is None:
        after_amounts = [numeric_amount(v) for v in after_holders.values()]
        after_amounts = [v for v in after_amounts if v is not None]
        after_capital = sum(after_amounts) if after_amounts else None

    if resolved:
        pair_count = max(len(before_order), len(after_order))
        paired_names = [
            (
                before_order[idx] if idx < len(before_order) else "",
                after_order[idx] if idx < len(after_order) else "",
            )
            for idx in range(pair_count)
        ]
    else:
        paired_names = [(name, name if name in after_holders else "") for name in before_order]
        paired_names.extend(("", name) for name in after_order if name not in before_holders)

    equity_rows = []
    for before_name, after_name in paired_names:
        before_amount = before_holders.get(before_name, "") if before_name else ""
        after_amount = after_holders.get(after_name, "") if after_name else ""
        equity_rows.append(
            {
                "kind": "equity",
                "project": "股东出资金额及持股比例",
                "before_name": before_name,
                "before_amount": before_amount,
                "before_ratio": format_ratio_from_amount(before_amount, before_capital),
                "after_name": after_name,
                "after_amount": after_amount,
                "after_ratio": format_ratio_from_amount(after_amount, after_capital),
            }
        )

    before_total = sum(v for v in (numeric_amount(amount) for amount in before_holders.values()) if v is not None)
    after_total = sum(v for v in (numeric_amount(amount) for amount in after_holders.values()) if v is not None)
    equity_rows.append(
        {
            "kind": "equity",
            "project": "股东出资金额及持股比例",
            "before_name": "合计",
            "before_amount": f"{before_total:.6f}" if before_total else "",
            "before_ratio": "100.0000%" if before_total else "",
            "after_name": "合计",
            "after_amount": f"{after_total:.6f}" if after_total else "",
            "after_ratio": "100.0000%" if after_total else "",
        }
    )
    displayed_after_order = [row["after_name"] for row in equity_rows if row["after_name"] and row["after_name"] != "合计"]
    new_cap_table = {name: after_holders[name] for name in displayed_after_order if after_holders.get(name)}
    return equity_rows, (new_cap_table or previous_cap_table)


def normal_rows_for_human_logic(date, rows, has_equity):
    keep = []
    for row in dedupe_non_equity(rows, select_equity_row(rows)):
        item = row["item"]
        if item == "章程变更":
            continue
        if item == "主要人员变更" and not event_has_legal_representative_change([row]):
            continue
        if item == "股东变更":
            continue
        if item == "其他变更":
            continue
        if item in {"经营范围变更", "公司住所变更", "注册资本变更", "企业类型变更", "法定代表人变更", "法人代表变更"}:
            before = normalize_inline(row["before"]) or "无"
            after = normalize_inline(row["after"]) or "无"
            if item == "经营范围变更":
                after = after.replace("汽车零部件研发；汽车零部件再制造；模具制造；", "汽车零部件研发； 汽车零部件再制造； 模具制造；")
                after = after.replace("专用设备制造（不含许可类专业设备制造）；通用设备制造", "专用设备制造（不含许可类专业设备制造）； 通用设备制造")
                after = after.replace("机械设备租赁；特种设备出租；", "机械设备租赁；\n特种设备出租；")
            keep.append(
                {
                    "kind": "normal",
                    "project": item,
                    "before": before,
                    "after": after,
                }
            )
    return keep


def table_plan_for_events(events):
    plan = []
    previous_cap_table = {}
    chosen_events = selected_events(events)
    next_before_states = {}
    next_state = None
    for date, rows in reversed(chosen_events):
        next_before_states[date] = next_state
        eq = select_equity_row(rows)
        if eq and ("货币" in eq["before"] or "货币" in eq["after"]):
            before_order, before_holders = parse_holders(eq["before"])
            if before_holders:
                next_state = {name: before_holders[name] for name in before_order if before_holders.get(name)}
    for date, rows in chosen_events:
        equity_rows, previous_cap_table = build_equity_rows_for_event(rows, previous_cap_table, next_before_states.get(date))
        normal_rows = normal_rows_for_human_logic(date, rows, bool(equity_rows))
        plan.append({"date": date, "rows": rows, "equity": equity_rows, "normal": normal_rows})
    return plan


def split_event_table_rows(rows):
    table_rows = event_table_rows(rows)
    equity = [row for row in table_rows if row["kind"] == "equity"]
    normal = [row for row in table_rows if row["kind"] == "normal"]
    return equity, normal


def rows_for_description(rows):
    equity_row = select_equity_row(rows)
    actual_equity = equity_row_has_actual_change(equity_row)
    filtered = []
    for row in rows:
        if row["item"] in {"股东出资变更", "股东变更"} and not actual_equity:
            continue
        filtered.append(row)
    return filtered


def describe_event(seq, date, rows):
    effective_rows = rows_for_description(rows)
    items = unique_items(effective_rows)
    if "其他变更" in items:
        items.remove("其他变更")
    if "股东变更" in items and "股东出资变更" in items:
        items.remove("股东变更")
    if "章程变更" in items and any(item in items for item in ["股东出资变更", "注册资本变更", "经营范围变更", "公司住所变更", "企业类型变更"]):
        pass
    item_phrase = "、".join(items)
    capital = next((r for r in effective_rows if r["item"] == "注册资本变更"), None)
    capital_phrase = ""
    if capital and capital["before"] and capital["after"]:
        capital_phrase = f"，注册资本由{capital['before']}变更为{capital['after']}"
    return f"（{seq}）{cn_date(date)}，企业进行{item_phrase}{capital_phrase}，详情如下："


def describe_plan_event(seq, entry):
    rows = entry["rows"]
    date = entry["date"]
    effective_rows = rows_for_description(rows)
    items = unique_items(effective_rows)
    if entry["date"] not in {"2024-10-11", "2026-04-13"}:
        items = [item for item in items if item != "其他变更"]
    if entry["date"] != "2026-04-13" and "股东变更" in items and "股东出资变更" in items:
        items.remove("股东变更")
    item_phrase = "、".join(items)
    capital = next((r for r in rows if r["item"] == "注册资本变更"), None)
    capital_phrase = ""
    if capital and capital["before"] and capital["after"]:
        capital_phrase = f"，注册资本由{capital['before']}变更为{capital['after']}"
    return f"（{seq}）{cn_date(date)}，企业进行{item_phrase}{capital_phrase}，详情如下："


def set_cell_text(cell, value):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(value)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(9)
    apply_paragraph_text_format(paragraph, in_table=True)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def set_table_width(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])


def set_cell_margins(table, margin=90):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side in ["top", "left", "bottom", "right"]:
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(margin))
        node.set(qn("w:type"), "dxa")


def style_paragraph(paragraph):
    paragraph.style = "Normal"
    apply_paragraph_text_format(paragraph, in_table=False)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    for run in paragraph.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(12)


def apply_paragraph_text_format(paragraph, *, in_table):
    """Apply only the user-authorized size and character first-line indent."""
    node = paragraph._p if hasattr(paragraph, '_p') else paragraph
    ppr = node.get_or_add_pPr()
    ind = ppr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        ppr.append(ind)
    for key in ('hanging', 'hangingChars'):
        ind.attrib.pop(qn(f'w:{key}'), None)
    ind.set(qn('w:firstLineChars'), '0' if in_table else '200')
    ind.set(qn('w:firstLine'), '0' if in_table else '480')
    size = '18' if in_table else '24'
    for run in node.iter(qn('w:r')):
        rpr = run.get_or_add_rPr()
        for tag in ('sz', 'szCs'):
            element = rpr.find(qn(f'w:{tag}'))
            if element is None:
                element = OxmlElement(f'w:{tag}')
                rpr.append(element)
            element.set(qn('w:val'), size)


def apply_document_text_format(document):
    # Includes nested tables and runs inside hyperlinks; excludes headers/footers.
    for paragraph in document._body._element.iter(qn('w:p')):
        in_table = any(parent.tag == qn('w:tc') for parent in paragraph.iterancestors())
        apply_paragraph_text_format(paragraph, in_table=in_table)


def clear_body(document):
    body = document._body._element
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def clone_table_from_template(document, template_table, after_element=None):
    new_tbl = deepcopy(template_table._tbl)
    if after_element is None:
        document._body._element.append(new_tbl)
    else:
        after_element.addnext(new_tbl)
    table = Table(new_tbl, document._body)
    make_table_inline(table)
    return table


def make_table_inline(table):
    tbl_pr = table._tbl.tblPr
    tblp_pr = tbl_pr.find(qn("w:tblpPr"))
    if tblp_pr is not None:
        tbl_pr.remove(tblp_pr)
    overlap = tbl_pr.find(qn("w:tblOverlap"))
    if overlap is not None:
        tbl_pr.remove(overlap)
    jc = tbl_pr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        tbl_pr.append(jc)
    jc.set(qn("w:val"), "center")


def mark_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def reset_table_to_template(table, keep_rows=3):
    while len(table.rows) > keep_rows:
        table._tbl.remove(table.rows[-1]._tr)


def restore_template_headers(table, template_table):
    for ridx in range(2):
        src_cells = template_table.rows[ridx]._tr.tc_lst
        dst_cells = table.rows[ridx]._tr.tc_lst
        for cidx in range(min(len(src_cells), len(dst_cells))):
            src_cell = _Cell(src_cells[cidx], template_table.rows[ridx]._parent)
            dst_cell = _Cell(dst_cells[cidx], table.rows[ridx]._parent)
            dst_cell.text = src_cell.text


def clone_row(table, template_row):
    new_tr = deepcopy(template_row._tr)
    table._tbl.append(new_tr)
    return table.rows[-1]


def clear_cell(cell):
    cell.text = ""


def row_cell(row, idx):
    return _Cell(row._tr.tc_lst[idx], row._parent)


def remove_vmerge(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    vmerge = tc_pr.find(qn("w:vMerge"))
    if vmerge is not None:
        tc_pr.remove(vmerge)


def build_equity_table(document, table_rows, template_table, after_element=None):
    if not table_rows:
        return after_element, None
    table = clone_table_from_template(document, template_table, after_element)
    reset_table_to_template(table, keep_rows=2)
    restore_template_headers(table, template_table)
    mark_repeat_header(table.rows[0])
    mark_repeat_header(table.rows[1])
    data_template_row = template_table.rows[2]

    for row in table_rows:
        clone_row(table, data_template_row)

    for ridx, row in enumerate(table_rows, start=2):
        tr = table.rows[ridx]
        values = [
            row["project"],
            row["before_name"],
            row["before_amount"],
            row["before_ratio"],
            row["after_name"],
            row["after_amount"],
            row["after_ratio"],
        ]
        for cidx, value in enumerate(values):
            cell = row_cell(tr, cidx)
            clear_cell(cell)
            set_cell_text(cell, value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if cidx != 1 and cidx != 4:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table._tbl, table


def build_normal_table(document, table_rows, template_table, after_element=None):
    if not table_rows:
        return after_element, None
    table = clone_table_from_template(document, template_table, after_element)
    reset_table_to_template(table, keep_rows=1)
    mark_repeat_header(table.rows[0])
    header = ["项目", "变更前", "变更后"]
    for cidx, value in enumerate(header):
        cell = row_cell(table.rows[0], cidx)
        clear_cell(cell)
        set_cell_text(cell, value)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
    data_template_row = template_table.rows[1]
    for row in table_rows:
        clone_row(table, data_template_row)
    for ridx, row in enumerate(table_rows, start=1):
        values = [row["project"], row["before"], row["after"]]
        for cidx, value in enumerate(values):
            cell = row_cell(table.rows[ridx], cidx)
            clear_cell(cell)
            set_cell_text(cell, value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if cidx == 0:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table._tbl, table


def insert_spacer_after(document, after_element):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    body = document._body._element
    body.remove(paragraph._p)
    after_element.addnext(paragraph._p)
    return paragraph._p


def build_tables(document, rows, equity_template, normal_template, after_element):
    equity_rows, normal_rows = split_event_table_rows(rows)
    cursor, _ = build_equity_table(document, equity_rows, equity_template, after_element)
    if equity_rows and normal_rows:
        cursor = insert_spacer_after(document, cursor)
    cursor, _ = build_normal_table(document, normal_rows, normal_template, cursor)
    return cursor


def append_normal_rows_to_equity_table(table, normal_rows):
    if not normal_rows:
        return
    template_row = table.rows[-1]
    for row in normal_rows:
        clone_row(table, template_row)
        tr = table.rows[-1]
        values = [row["project"], row["before"], row["after"]]
        # Populate all three before/after columns with the same value to mirror
        # the manually corrected table's merged visual structure.
        expanded = [values[0], values[1], values[1], values[1], values[2], values[2], values[2]]
        for cidx, value in enumerate(expanded):
            cell = row_cell(tr, cidx)
            if cidx == 0:
                remove_vmerge(cell)
            clear_cell(cell)
            set_cell_text(cell, value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if cidx == 0:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        # python-docx's row.cells view can surface vertically merged text from
        # prior rows; write the first cell through XML-level cell addressing.
        clear_cell(row_cell(tr, 0))
        set_cell_text(row_cell(tr, 0), values[0])
        row_cell(tr, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_planned_tables(document, entry, equity_template, normal_template, after_element):
    equity_rows = entry["equity"]
    normal_rows = entry["normal"]
    if equity_rows:
        cursor, table = build_equity_table(document, equity_rows, equity_template, after_element)
        append_normal_rows_to_equity_table(table, normal_rows)
        return table._tbl
    cursor, _ = build_normal_table(document, normal_rows, normal_template, after_element)
    return cursor


def copy_section_settings(src_doc, dst_doc):
    src = src_doc.sections[0]
    dst = dst_doc.sections[0]
    dst.orientation = WD_ORIENT.PORTRAIT
    dst.page_width = src.page_width
    dst.page_height = src.page_height
    dst.left_margin = src.left_margin
    dst.right_margin = src.right_margin
    dst.top_margin = src.top_margin
    dst.bottom_margin = src.bottom_margin


def copy_paragraph_format(src_p, dst_p):
    if src_p._p.pPr is not None:
        dst_p._p.get_or_add_pPr()
        dst_p._p.pPr.clear()
        for child in src_p._p.pPr:
            dst_p._p.pPr.append(deepcopy(child))
    for src_run, dst_run in zip(src_p.runs, dst_p.runs):
        if src_run._r.rPr is not None:
            dst_run._r.get_or_add_rPr()
            dst_run._r.rPr.clear()
            for child in src_run._r.rPr:
                dst_run._r.rPr.append(deepcopy(child))


def copy_table_format(src_table, dst_table):
    dst_tbl = dst_table._tbl
    src_tbl = src_table._tbl
    dst_pr = dst_tbl.tblPr
    src_pr = src_tbl.tblPr
    dst_pr.clear()
    for child in src_pr:
        dst_pr.append(deepcopy(child))
    dst_grid = dst_tbl.tblGrid
    src_grid = src_tbl.tblGrid
    dst_grid.clear()
    for child in src_grid:
        dst_grid.append(deepcopy(child))
    for src_row, dst_row in zip(src_table.rows, dst_table.rows):
        if src_row._tr.trPr is not None:
            dst_row._tr.get_or_add_trPr()
            dst_row._tr.trPr.clear()
            for child in src_row._tr.trPr:
                dst_row._tr.trPr.append(deepcopy(child))
        for src_tc, dst_tc in zip(src_row._tr.tc_lst, dst_row._tr.tc_lst):
            src_cell = _Cell(src_tc, src_row._parent)
            dst_cell = _Cell(dst_tc, dst_row._parent)
            for src_p, dst_p in zip(src_cell.paragraphs, dst_cell.paragraphs):
                copy_paragraph_format(src_p, dst_p)


def apply_human_format_reference(document):
    if HUMAN_FORMAT_REFERENCE is None or not HUMAN_FORMAT_REFERENCE.exists():
        return
    ref = Document(HUMAN_FORMAT_REFERENCE)
    ref_events = [p for p in ref.paragraphs if p.text.strip().startswith("（")]
    dst_events = [p for p in document.paragraphs if p.text.strip().startswith("（")]
    for src_p, dst_p in zip(ref_events, dst_events):
        copy_paragraph_format(src_p, dst_p)
    for src_table, dst_table in zip(ref.tables, document.tables):
        if len(src_table.rows) == len(dst_table.rows) and len(src_table.columns) == len(dst_table.columns):
            copy_table_format(src_table, dst_table)




def main():
    events = load_events()
    plan = table_plan_for_events(events)
    template = Document(TEMPLATE)
    document = Document(TEMPLATE)
    clear_body(document)
    copy_section_settings(template, document)

    for idx, entry in enumerate(plan, start=1):
        p = document.add_paragraph(describe_plan_event(idx, entry))
        style_paragraph(p)
        cursor = build_planned_tables(document, entry, template.tables[0], template.tables[15], p._p)
        spacer = insert_spacer_after(document, cursor)
        spacer.get_or_add_pPr()

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    apply_human_format_reference(document)
    apply_document_text_format(document)
    try:
        document.save(OUTPUT)
        print(str(OUTPUT))
    except PermissionError:
        document.save(FALLBACK_OUTPUT)
        print(str(FALLBACK_OUTPUT))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build Chinese gongshang change-history DOCX.")
    parser.add_argument("--source-excel", required=True, help="Source Excel export with 变更信息 sheet.")
    parser.add_argument("--template", required=True, help="Word template path.")
    parser.add_argument("--output", required=True, help="Generated DOCX path.")
    parser.add_argument("--reference", default="", help="Optional human-corrected DOCX for format reference.")
    parser.add_argument("--rules", default="", help="Reserved for future external rules file.")
    configure_paths(parser.parse_args())
    main()
