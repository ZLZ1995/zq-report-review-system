from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    wb = load_workbook(args.template, data_only=False)
    chain = {
        "classification_summary": {},
        "balance_sheet": {},
    }
    if "分类汇总" in wb.sheetnames:
        ws = wb["分类汇总"]
        for r in range(6, min(ws.max_row, 91) + 1):
            chain["classification_summary"][str(r)] = {
                "name": ws.cell(r, 2).value,
                "book_value_formula": ws.cell(r, 5).value,
                "enterprise_formula": ws.cell(r, 9).value,
                "difference_formula": ws.cell(r, 10).value,
            }
    if "资产负债表" in wb.sheetnames:
        ws = wb["资产负债表"]
        for r in range(7, 39):
            chain["balance_sheet"][str(r)] = {
                "asset_name": ws.cell(r, 1).value,
                "asset_formula": ws.cell(r, 3).value,
                "liability_name": ws.cell(r, 6).value,
                "liability_formula": ws.cell(r, 8).value,
            }
    wb.close()
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(chain, ensure_ascii=False, indent=2), encoding="utf-8")
    print(out.as_posix())


if __name__ == "__main__":
    main()
