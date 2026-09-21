from __future__ import annotations

import argparse
import json
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any

try:
    import xlrd
except ImportError:
    xlrd = None
from openpyxl import load_workbook
from detail_mapping_config import (
    CURRENT_LIABILITY_SHEET_MAPPING,
    NONCURRENT_ASSET_SHEET_MAPPING,
    NONCURRENT_LIABILITY_SHEET_MAPPING,
    OWNER_EQUITY_SHEET_MAPPING,
)

try:
    import xml.etree.ElementTree as _ET

    if not hasattr(_ET.ElementTree, "getiterator"):
        _ET.ElementTree.getiterator = _ET.ElementTree.iter  # type: ignore[attr-defined]
except Exception:
    pass


def clean(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def number(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = clean(value).replace(",", "").replace("，", "")
    if not text or text in {"-", "--"}:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def normalize_bs_labels(values: dict[str, float]) -> dict[str, float]:
    normalized = dict(values)
    rename_pairs = [
        ("预收款项", "预收账款"),
        ("实收资本（或股本）", "实收资本"),
    ]
    for src, dst in rename_pairs:
        if src in normalized and dst not in normalized:
            normalized[dst] = normalized[src]
    return normalized


def parse_trial_balance(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows: list[dict[str, Any]] = []
    is_y71_style = clean(ws.cell(1, 1).value) == "科目汇总试算表" and clean(ws.cell(3, 2).value) == "科目代码"
    is_two_line_tb = (
        clean(ws.cell(1, 1).value) == "科目代码"
        and (
            clean(ws.cell(1, 11).value) == "期末余额"
            or (clean(ws.cell(2, 11).value) == "借方" and clean(ws.cell(2, 12).value) == "贷方")
        )
    )
    start_row = 3 if is_two_line_tb else 2
    for r in range(start_row, ws.max_row + 1):
        if is_y71_style:
            tb_code = clean(ws.cell(r, 2).value)
            aux_code = clean(ws.cell(r, 3).value)
            tb_account_name = clean(ws.cell(r, 4).value)
            aux_name = clean(ws.cell(r, 5).value) or clean(ws.cell(r, 3).value)
            debit_end = number(ws.cell(r, 10).value)
            credit_end = number(ws.cell(r, 11).value)
        elif is_two_line_tb:
            tb_code = clean(ws.cell(r, 1).value)
            aux_code = ""
            tb_account_name = clean(ws.cell(r, 2).value)
            aux_name = ""
            debit_end = number(ws.cell(r, 11).value)
            credit_end = number(ws.cell(r, 12).value)
        else:
            tb_code = clean(ws.cell(r, 1).value)
            aux_code = ""
            tb_account_name = clean(ws.cell(r, 2).value)
            aux_name = clean(ws.cell(r, 3).value)
            debit_end = number(ws.cell(r, 10).value)
            credit_end = number(ws.cell(r, 11).value)
        if not tb_code or not tb_account_name:
            continue
        if aux_name in {"人民币", "综合本位币"} and "." in tb_code:
            aux_name = tb_account_name
        direction = "debit" if abs(debit_end) >= abs(credit_end) else "credit"
        rows.append(
            {
                "tb_code": tb_code,
                "aux_code": aux_code,
                "tb_account_name": tb_account_name,
                "aux_name": aux_name,
                "debit_end": debit_end,
                "credit_end": credit_end,
                "direction": direction,
            }
        )
    wb.close()
    return rows


def parse_balance_sheet(path: Path) -> dict[str, float]:
    if path.suffix.lower() == ".xls":
        if xlrd is None:
            raise RuntimeError("xlrd_required_for_legacy_xls")
        book = xlrd.open_workbook(path.as_posix())
        sheet = book.sheet_by_index(0)
        values: dict[str, float] = {}
        if sheet.ncols >= 4 and clean(sheet.cell_value(0, 0)) == "项目" and clean(sheet.cell_value(0, 1)) == "报表侧":
            for r in range(1, sheet.nrows):
                label = clean(sheet.cell_value(r, 0)).replace("：", "").strip()
                side = clean(sheet.cell_value(r, 1))
                current_value = number(sheet.cell_value(r, 2))
                if not label or label in {"项目", "报表侧", "资产", "负债和所有者权益", "负债和股东权益"}:
                    continue
                if side not in {"资产", "负债", "权益", "负债和权益", "负债及权益"}:
                    continue
                values[label] = current_value
            return normalize_bs_labels(values)
        if sheet.ncols >= 5 and clean(sheet.cell_value(8, 1)) == "期末余额":
            for r in range(sheet.nrows):
                label = clean(sheet.cell_value(r, 0)).replace("帐", "账").strip()
                value = number(sheet.cell_value(r, 1))
                if label and not label.endswith("：") and label not in {"期末余额", "流动资产", "非流动资产", "流动负债", "非流动负债", "所有者权益", "所有者权益（或股东权益）"}:
                    values[label] = value
            return values
        for r in range(sheet.nrows):
            left_label = clean(sheet.cell_value(r, 1)).replace("帐", "账").strip()
            left_value = number(sheet.cell_value(r, 2))
            right_label = clean(sheet.cell_value(r, 4)).replace("帐", "账").strip() if sheet.ncols > 4 else ""
            right_value = number(sheet.cell_value(r, 5)) if sheet.ncols > 5 else 0.0
            if left_label and not left_label.endswith("：") and left_label not in {"报表项名称"}:
                values[left_label] = left_value
            if right_label and not right_label.endswith("：") and right_label not in {"报表项名称"}:
                values[right_label] = right_value
        return normalize_bs_labels(values)
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = wb[wb.sheetnames[0]]
    values: dict[str, float] = {}
    if sheet.max_column >= 4 and clean(sheet.cell(1, 1).value) == "项目" and clean(sheet.cell(1, 2).value) == "报表侧":
        for r in range(2, sheet.max_row + 1):
            label = clean(sheet.cell(r, 1).value).replace("：", "").strip()
            side = clean(sheet.cell(r, 2).value)
            current_value = number(sheet.cell(r, 3).value)
            if not label or label in {"项目", "报表侧", "资产", "负债和所有者权益", "负债和股东权益"}:
                continue
            if side not in {"资产", "负债", "权益", "负债和权益", "负债及权益"}:
                continue
            values[label] = current_value
        wb.close()
        return normalize_bs_labels(values)
    for r in range(1, sheet.max_row + 1):
        left_label = clean(sheet.cell(r, 2).value).replace("帐", "账").strip()
        left_value = number(sheet.cell(r, 4).value) if clean(sheet.cell(6, 3).value) == "行次" else number(sheet.cell(r, 3).value)
        right_label = clean(sheet.cell(r, 5).value).replace("帐", "账").strip()
        right_label = clean(sheet.cell(r, 7).value).replace("帐", "账").strip() if clean(sheet.cell(6, 8).value) == "行次" else right_label
        right_value = number(sheet.cell(r, 9).value) if clean(sheet.cell(6, 8).value) == "行次" else number(sheet.cell(r, 6).value)
        if left_label and not left_label.endswith("：") and left_label not in {"报表项名称"}:
            values[left_label] = left_value
        if right_label and not right_label.endswith("：") and right_label not in {"报表项名称"}:
            values[right_label] = right_value
    wb.close()
    return normalize_bs_labels(values)


def discover_customer_detail_workbook(trial_balance_path: Path) -> Path | None:
    # Kept for caller compatibility; directory proximity is not user consent.
    return None


def load_customer_detail_overrides(path: Path | None) -> dict[tuple[str, str], dict[str, str]]:
    if path is None or not path.exists():
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    if "余额" not in wb.sheetnames:
        wb.close()
        return {}
    ws = wb["余额"]
    overrides: dict[tuple[str, str], dict[str, str]] = {}
    for r in range(3, ws.max_row + 1):
        tb_code = clean(ws.cell(r, 2).value)
        aux_code = clean(ws.cell(r, 3).value)
        aux_name = clean(ws.cell(r, 5).value)
        balance_class = clean(ws.cell(r, 13).value)
        object_name = clean(ws.cell(r, 14).value)
        if not tb_code or not balance_class:
            continue
        payload = {
            "balance_class": balance_class,
            "object_name": object_name,
        }
        if aux_name:
            overrides[(tb_code, aux_name)] = payload
        if aux_code:
            overrides[(tb_code, aux_code)] = payload
    wb.close()
    return overrides


def infer_target(
    tb_code: str,
    tb_account_name: str,
    aux_code: str,
    aux_name: str,
    direction: str,
    bs_values: dict[str, float],
    customer_overrides: dict[tuple[str, str], dict[str, str]],
) -> tuple[str, str, str]:
    if tb_code.startswith(("5", "6")):
        return "", "", "mapping_unresolved"
    if tb_account_name in NONCURRENT_ASSET_SHEET_MAPPING:
        item = NONCURRENT_ASSET_SHEET_MAPPING[tb_account_name]
        return item["sheet"], item["bs_line"], item["policy"]
    if tb_account_name in CURRENT_LIABILITY_SHEET_MAPPING:
        item = CURRENT_LIABILITY_SHEET_MAPPING[tb_account_name]
        return item["sheet"], item["bs_line"], item["policy"]
    if tb_account_name in NONCURRENT_LIABILITY_SHEET_MAPPING:
        item = NONCURRENT_LIABILITY_SHEET_MAPPING[tb_account_name]
        return item["sheet"], item["bs_line"], item["policy"]
    if tb_account_name in OWNER_EQUITY_SHEET_MAPPING:
        item = OWNER_EQUITY_SHEET_MAPPING[tb_account_name]
        return item["sheet"], item["bs_line"], item["policy"]
    name = tb_account_name + " " + aux_name
    override = customer_overrides.get((tb_code, aux_name)) or customer_overrides.get((tb_code, aux_code))
    if override:
        balance_class = override.get("balance_class", "")
        object_name = override.get("object_name", "")
        if balance_class == "其他应付":
            policy = "placeholder_only" if object_name in {"", "见明细", "默认值"} else "detail_fillable"
            return "其他应付款", "其他应付款", policy
        if balance_class == "其他应收":
            policy = "placeholder_only" if object_name in {"", "见明细", "默认值"} else "detail_fillable"
            return "其他应收款", "其他应收款", policy
    if tb_code.startswith("1002") or tb_code.startswith("100201"):
        return "银行存款", "货币资金", "direct_balance_sync"
    if tb_code.startswith("1124") or tb_code.startswith("1125") or tb_code.startswith("1221"):
        return "其他应收款", "其他应收款", "detail_fillable"
    if tb_code.startswith("112403"):
        return "应收账款", "应收账款", "detail_fillable"
    if tb_code.startswith("2202") or tb_code.startswith("220201"):
        return "应付账款", "应付账款", "detail_fillable"
    if tb_code.startswith("2221") or tb_code.startswith("2225") or tb_code.startswith("222108"):
        return "应交税费", "应交税费", "direct_balance_sync"
    if tb_code.startswith("224199"):
        return "其他应付款", "其他应付款", "placeholder_only"
    if tb_code.startswith("2241"):
        return "其他应付款", "其他应付款", "detail_fillable"
    if "应交税费" in tb_account_name or "应交税金" in tb_account_name:
        return "应交税费", "应交税费", "direct_balance_sync"
    if "应付账款" in tb_account_name:
        return "应付账款", "应付账款", "detail_fillable"
    if "预提费用" in tb_account_name:
        return "应付账款", "应付账款", "detail_fillable"
    if "其他应付款" in tb_account_name or ("内部往来" in tb_account_name and direction == "credit"):
        return "其他应付款", "其他应付款", "detail_fillable"
    if "预付" in tb_account_name:
        return "预付账款", "预付款项", "detail_fillable"
    if "应收账款" in tb_account_name:
        return "应收账款", "应收账款", "detail_fillable"
    if "其他应收款" in tb_account_name or ("内部往来" in tb_account_name and direction == "debit"):
        return "其他应收款", "其他应收款", "detail_fillable"
    if "货币资金" in tb_account_name or "银行存款" in tb_account_name:
        return "银行存款", "货币资金", "direct_balance_sync"
    return "", "", "mapping_unresolved"


def parse_trial_balance(path: Path) -> list[dict[str, Any]]:
    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as zf:
        sheet_xml = ET.fromstring(zf.read("xl/worksheets/sheet1.xml"))
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in zf.namelist():
            sst_root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            for si in sst_root.findall(".//x:si", ns):
                text = "".join(t.text or "" for t in si.findall(".//x:t", ns))
                shared_strings.append(text)

    def read_cell_text(cell) -> str:
        cell_type = cell.attrib.get("t", "")
        if cell_type == "inlineStr":
            node = cell.find("x:is/x:t", ns)
            return node.text if node is not None and node.text is not None else ""
        v = cell.find("x:v", ns)
        if v is None or v.text is None:
            return ""
        if cell_type == "s":
            idx = int(v.text)
            return shared_strings[idx] if 0 <= idx < len(shared_strings) else ""
        return v.text

    rows: list[dict[str, Any]] = []
    first_row_map: dict[str, str] = {}
    for row in sheet_xml.findall(".//x:sheetData/x:row", ns):
        if int(row.attrib.get("r", "0")) != 1:
            continue
        for cell in row.findall("x:c", ns):
            ref = cell.attrib.get("r", "")
            col = "".join(ch for ch in ref if ch.isalpha())
            first_row_map[col] = read_cell_text(cell)
        break
    is_y71_style = clean(first_row_map.get("A", "")) == "科目汇总试算表"
    is_two_line_tb = clean(first_row_map.get("A", "")) == "科目代码" and (
        clean(first_row_map.get("K", "")) == "期末余额" or clean(first_row_map.get("K", "")) == "借方"
    )
    for row in sheet_xml.findall(".//x:sheetData/x:row", ns):
        r = int(row.attrib.get("r", "0"))
        if r < (5 if is_y71_style else (3 if is_two_line_tb else 5)):
            continue
        cell_map: dict[str, str] = {}
        for cell in row.findall("x:c", ns):
            ref = cell.attrib.get("r", "")
            col = "".join(ch for ch in ref if ch.isalpha())
            cell_map[col] = read_cell_text(cell)
        if is_y71_style:
            tb_code = clean(cell_map.get("B", ""))
            aux_code = clean(cell_map.get("C", ""))
            tb_account_name = clean(cell_map.get("D", ""))
            aux_name = clean(cell_map.get("E", "")) or aux_code or "默认值"
            debit_end = number(cell_map.get("J", ""))
            credit_end = number(cell_map.get("K", ""))
        elif is_two_line_tb:
            tb_code = clean(cell_map.get("A", ""))
            aux_code = ""
            tb_account_name = clean(cell_map.get("B", ""))
            debit_end = number(cell_map.get("K", ""))
            credit_end = number(cell_map.get("L", ""))
            aux_name = "默认值"
        else:
            tb_code = clean(cell_map.get("B", ""))
            aux_code = clean(cell_map.get("C", ""))
            tb_account_name = clean(cell_map.get("C", ""))
            debit_end = number(cell_map.get("F", ""))
            credit_end = number(cell_map.get("G", ""))
            aux_name = tb_account_name
        if not tb_code or not tb_account_name:
            continue
        direction = "debit" if abs(debit_end) >= abs(credit_end) else "credit"
        rows.append(
            {
                "tb_code": tb_code,
                "aux_code": aux_code,
                "tb_account_name": tb_account_name,
                "aux_name": aux_name,
                "debit_end": debit_end,
                "credit_end": credit_end,
                "direction": direction,
            }
        )
    return rows


def infer_target(
    tb_code: str,
    tb_account_name: str,
    aux_code: str,
    aux_name: str,
    direction: str,
    bs_values: dict[str, float],
    customer_overrides: dict[tuple[str, str], dict[str, str]],
) -> tuple[str, str, str]:
    if tb_code.startswith(("5", "6")):
        return "", "", "mapping_unresolved"
    account_text = clean(tb_account_name)
    override = customer_overrides.get((tb_code, aux_name)) or customer_overrides.get((tb_code, aux_code))
    if override:
        balance_class = override.get("balance_class", "")
        object_name = override.get("object_name", "")
        if balance_class == "其他应付款":
            policy = "placeholder_only" if object_name in {"", "见明细", "默认值"} else "detail_fillable"
            return "其他应付款", "其他应付款", policy
        if balance_class == "其他应收款":
            policy = "placeholder_only" if object_name in {"", "见明细", "默认值"} else "detail_fillable"
            return "其他应收款", "其他应收款", policy

    if account_text in NONCURRENT_ASSET_SHEET_MAPPING:
        item = NONCURRENT_ASSET_SHEET_MAPPING[account_text]
        return item["sheet"], item["bs_line"], item["policy"]
    if account_text in CURRENT_LIABILITY_SHEET_MAPPING:
        item = CURRENT_LIABILITY_SHEET_MAPPING[account_text]
        return item["sheet"], item["bs_line"], item["policy"]
    if account_text in NONCURRENT_LIABILITY_SHEET_MAPPING:
        item = NONCURRENT_LIABILITY_SHEET_MAPPING[account_text]
        return item["sheet"], item["bs_line"], item["policy"]
    if account_text in OWNER_EQUITY_SHEET_MAPPING:
        item = OWNER_EQUITY_SHEET_MAPPING[account_text]
        return item["sheet"], item["bs_line"], item["policy"]

    if "银行存款" in account_text or "货币资金" in account_text:
        return "银行存款", "货币资金", "direct_balance_sync"
    if "内部往来" in account_text and direction == "credit":
        return "其他应付款", "其他应付款", "detail_fillable"
    if "内部往来" in account_text and direction == "debit":
        return "其他应收款", "其他应收款", "detail_fillable"
    if "应收账款" in account_text and "坏账准备" not in account_text:
        return "应收账款", "应收账款", "detail_fillable"
    if "预付账款" in account_text or "预付款项" in account_text:
        return "预付账款", "预付款项", "detail_fillable"
    if any(token in account_text for token in ["其他应收款", "押金", "保证金", "房屋押金"]):
        return "其他应收款", "其他应收款", "detail_fillable"
    if any(token in account_text for token in ["存货", "原材料", "库存商品", "发出商品", "委托加工物资", "半成品", "在产品"]):
        return "存货汇总", "存货", "placeholder_only"
    if "应付账款" in account_text:
        return "应付账款", "应付账款", "detail_fillable"
    if "预收账款" in account_text or "预收款项" in account_text or "短期预收账款" in account_text:
        return "预收账款", "预收账款", "detail_fillable"
    if "应付职工薪酬" in account_text or "工资" in account_text:
        return "职工薪酬", "应付职工薪酬", "detail_fillable"
    if any(token in account_text for token in ["应交税费", "应交税金", "所得税", "增值税", "印花税", "城建税", "教育费附加", "代扣代缴税金", "个人所得税"]):
        return "应交税费", "应交税费", "direct_balance_sync"
    if any(token in account_text for token in ["其他应付款", "代扣代缴", "工会经费"]):
        return "其他应付款", "其他应付款", "detail_fillable"

    if tb_code.startswith("1002"):
        return "银行存款", "货币资金", "direct_balance_sync"
    if tb_code.startswith("1122"):
        return "应收账款", "应收账款", "detail_fillable"
    if tb_code.startswith("1123"):
        return "预付账款", "预付款项", "detail_fillable"
    if tb_code.startswith("122102"):
        return "其他应收款", "其他应收款", "detail_fillable"
    if tb_code.startswith("2202"):
        return "应付账款", "应付账款", "detail_fillable"
    if tb_code.startswith("2203"):
        return "预收账款", "预收账款", "detail_fillable"
    if tb_code.startswith("2221") or tb_code.startswith("2225"):
        return "应交税费", "应交税费", "direct_balance_sync"
    if tb_code.startswith("2241"):
        return "其他应付款", "其他应付款", "detail_fillable"
    return "", "", "mapping_unresolved"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--trial-balance")
    parser.add_argument("--balance-sheet", required=True)
    parser.add_argument("--journal")
    parser.add_argument("--customer-detail", help="Explicitly selected customer detail workbook only")
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    bs_values = parse_balance_sheet(Path(args.balance_sheet))
    tb_rows = parse_trial_balance(Path(args.trial_balance)) if args.trial_balance else []
    customer_detail_workbook = Path(args.customer_detail) if args.customer_detail else None
    customer_overrides = load_customer_detail_overrides(customer_detail_workbook)
    account_mappings = []
    for row in tb_rows:
        if "." in row["tb_code"]:
            top_name = row["tb_code"].split(".", 1)[0]
            if top_name in {"1002", "1131", "1133", "1151", "2121", "2131", "2151", "2161", "2171", "2181", "2191", "2312"}:
                base_name_map = {
                    "1002": "银行存款",
                    "1131": "应收账款",
                    "1133": "其他应收款",
                    "1151": "预付账款",
                    "2121": "应付账款",
                    "2131": "预收账款",
                    "2151": "应付工资",
                    "2161": "应付股利",
                    "2171": "应交税金",
                    "2181": "其他应付款",
                    "2191": "预提费用",
                    "2312": "应付利息",
                }
                row = {**row, "tb_account_name": base_name_map[top_name]}
        target_sheet, bs_line, detail_policy = infer_target(
            row["tb_code"],
            row["tb_account_name"],
            row.get("aux_code", ""),
            row["aux_name"],
            row["direction"],
            bs_values,
            customer_overrides,
        )
        if not target_sheet:
            continue
        account_mappings.append(
            {
                "tb_code": row["tb_code"],
                "aux_code": row.get("aux_code", ""),
                "tb_account_name": row["tb_account_name"],
                "aux_name": row["aux_name"],
                "direction": row["direction"],
                "bs_line": bs_line,
                "target_sheet": target_sheet,
                "detail_policy": detail_policy,
                "balance_class_override": (customer_overrides.get((row["tb_code"], row["aux_name"])) or customer_overrides.get((row["tb_code"], row.get("aux_code", ""))) or {}).get("balance_class", ""),
                "object_name_override": (customer_overrides.get((row["tb_code"], row["aux_name"])) or customer_overrides.get((row["tb_code"], row.get("aux_code", ""))) or {}).get("object_name", ""),
                "evidence": ["trial_balance"],
            }
        )

    payload = {
        "project_id": Path(args.trial_balance or args.balance_sheet).stem,
        "sources": {
            "trial_balance": args.trial_balance,
            "balance_sheet": args.balance_sheet,
            "journal": args.journal or "",
        },
        "account_mappings": account_mappings,
        "bs_lines": bs_values,
    }
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(out.as_posix())


if __name__ == "__main__":
    main()
