from __future__ import annotations

import argparse
import json
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


TOTAL_LABELS = {"\u5408\u8ba1", "åˆè®¡"}
AMOUNT_TOLERANCE = Decimal("0.0001")


def set_run_font(run, font_name: str = "仿宋", size_pt: float = 12.0, bold: bool = False) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size_pt)
    run.bold = bold


def add_paragraph(doc: Document, text: str, bold: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(24)
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)


def set_cell_text(cell, text: str, font_size: float = 9.0, align: str = "center") -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    set_run_font(run, "宋体", font_size)
    if align == "center":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def mark_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def add_normal_table(doc: Document, rows: list[dict[str, str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=3)
    table.autofit = True
    set_table_borders(table)
    mark_repeat_header(table.rows[0])
    for idx, value in enumerate(["项目", "变更前", "变更后"]):
        set_cell_text(table.cell(0, idx), value, 9)
    for row_idx, row in enumerate(rows, start=1):
        values = [row.get("project", ""), row.get("before", ""), row.get("after", "")]
        for col_idx, value in enumerate(values):
            set_cell_text(table.cell(row_idx, col_idx), value, 8.5, "center" if col_idx == 0 else "left")


def add_founding_table(doc: Document, rows: list[dict[str, str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=3)
    table.autofit = True
    set_table_borders(table)
    mark_repeat_header(table.rows[0])
    for idx, value in enumerate(["出资人", "出资金额（万元人民币）", "出资比例"]):
        set_cell_text(table.cell(0, idx), value, 9)
    for row_idx, row in enumerate(rows, start=1):
        values = [row.get("name", ""), row.get("amount", ""), row.get("ratio", "")]
        for col_idx, value in enumerate(values):
            set_cell_text(table.cell(row_idx, col_idx), value, 8.5, "left" if col_idx == 0 else "center")


def add_equity_table(doc: Document, rows: list[dict[str, str]]) -> None:
    table = doc.add_table(rows=2 + len(rows), cols=7)
    table.autofit = True
    set_table_borders(table)
    mark_repeat_header(table.rows[0])
    mark_repeat_header(table.rows[1])

    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 1).merge(table.cell(0, 3))
    table.cell(0, 4).merge(table.cell(0, 6))
    set_cell_text(table.cell(0, 0), "项目", 9)
    set_cell_text(table.cell(0, 1), "变更前", 9)
    set_cell_text(table.cell(0, 4), "变更后", 9)
    for idx, value in enumerate(["股东名称", "出资金额（万元）", "持股比例", "股东名称", "出资金额（万元）", "持股比例"], start=1):
        set_cell_text(table.cell(1, idx), value, 8.5)

    for row_idx, row in enumerate(rows, start=2):
        values = [
            row.get("project", "") if row_idx == 2 else "",
            row.get("before_name", ""),
            row.get("before_amount", ""),
            row.get("before_ratio", ""),
            row.get("after_name", ""),
            row.get("after_amount", ""),
            row.get("after_ratio", ""),
        ]
        for col_idx, value in enumerate(values):
            align = "left" if col_idx in {1, 4} else "center"
            set_cell_text(table.cell(row_idx, col_idx), value, 8.2, align)


def is_founding_equity_rows(rows: list[dict[str, str]]) -> bool:
    if not rows:
        return False
    has_after_holder = any(str(row.get("after_name", "")).strip() for row in rows)
    has_before_holder = any(str(row.get("before_name", "")).strip() not in {"", "合计"} for row in rows)
    has_before_amount = any(str(row.get("before_amount", "")).strip() for row in rows)
    has_before_ratio = any(str(row.get("before_ratio", "")).strip() for row in rows)
    return has_after_holder and not (has_before_holder or has_before_amount or has_before_ratio)


def founding_rows_from_equity(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    return [
        {
            "name": str(row.get("after_name") or row.get("before_name") or ""),
            "amount": str(row.get("after_amount") or row.get("before_amount") or ""),
            "ratio": str(row.get("after_ratio") or row.get("before_ratio") or ""),
        }
        for row in rows
    ]


def normalize_events(content: dict[str, Any]) -> list[dict[str, Any]]:
    history = content.get("history_section", content)
    events = []
    visible_seq = 1
    for event in history.get("events", []):
        if event.get("omitted"):
            continue
        desc = event.get("description", "")
        if "seq" in event:
            desc = desc.replace(f"（{event['seq']}）", f"（{visible_seq}）", 1)
        events.append({**event, "visible_seq": visible_seq, "description": desc})
        visible_seq += 1
    return events


def _normalize_text(value: Any) -> str:
    return str(value or "").strip()


def _is_total_label(value: Any) -> bool:
    return _normalize_text(value) in TOTAL_LABELS


def _decimal_value(value: Any) -> Decimal | None:
    text = _normalize_text(value).replace(",", "").replace("%", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _decimal_close(left: Decimal | None, right: Decimal | None) -> bool:
    if left is None or right is None:
        return False
    return abs(left - right) <= AMOUNT_TOLERANCE


def _side_totals(rows: list[dict[str, Any]], side: str) -> dict[str, Any]:
    name_key = f"{side}_name"
    amount_key = f"{side}_amount"
    ratio_key = f"{side}_ratio"
    amount_sum = Decimal("0")
    ratio_sum = Decimal("0")
    total_amount = None
    total_ratio = None
    names = []
    missing_amount_or_ratio = []
    duplicate_names = []
    seen = set()

    for index, row in enumerate(rows):
        name = _normalize_text(row.get(name_key))
        amount = _decimal_value(row.get(amount_key))
        ratio = _decimal_value(row.get(ratio_key))

        if _is_total_label(name):
            total_amount = amount
            total_ratio = ratio
            continue

        if not name and amount is None and ratio is None:
            continue

        if name:
            if name in seen:
                duplicate_names.append(name)
            seen.add(name)
            names.append(name)

        if amount is None or ratio is None:
            missing_amount_or_ratio.append({"row_index": index, "name": name})

        if amount is not None:
            amount_sum += amount
        if ratio is not None:
            ratio_sum += ratio

    return {
        "amount_sum": str(amount_sum),
        "total_amount": str(total_amount) if total_amount is not None else None,
        "ratio_sum": str(ratio_sum),
        "total_ratio": str(total_ratio) if total_ratio is not None else None,
        "amount_ok": _decimal_close(amount_sum, total_amount),
        "ratio_ok": _decimal_close(ratio_sum, total_ratio),
        "complete": not missing_amount_or_ratio,
        "missing_amount_or_ratio": missing_amount_or_ratio,
        "unique_names": not duplicate_names,
        "duplicate_names": duplicate_names,
        "names": names,
    }


def _side_holder_amounts(rows: list[dict[str, Any]], side: str) -> dict[str, str]:
    result = {}
    for row in rows:
        name = _normalize_text(row.get(f"{side}_name"))
        if not name or _is_total_label(name):
            continue
        amount = _decimal_value(row.get(f"{side}_amount"))
        if amount is not None:
            result[name] = str(amount)
    return result


def _parse_fact_holder_amounts(value: Any) -> dict[str, str]:
    result = {}
    text = _normalize_text(value)
    if not text:
        return result
    parts = re.split(r"[\u3001,;；]", text)
    for part in parts:
        match = re.search(
            r"(.+?)(?:\uff0c?\s*\u51fa\u8d44\u989d)?\s*([0-9]+(?:\.[0-9]+)?)\s*\u4e07",
            part.strip(),
        )
        if not match:
            continue
        name = match.group(1).strip()
        amount = _decimal_value(match.group(2))
        if name and amount is not None:
            result[name] = str(amount)
    return result


def _equity_fact_maps(content: dict[str, Any] | None) -> dict[tuple[str, str], dict[str, str]]:
    if not content:
        return {}
    fact_sheet = content.get("fact_sheet", {})
    facts = fact_sheet.get("confirmed_facts", [])
    result = {}
    for fact in facts:
        field = _normalize_text(fact.get("field"))
        value = fact.get("value")
        match = re.match(r"(\d{4}-\d{2}-\d{2})_equity_change_(before|after)$", field)
        if not match:
            continue
        holder_amounts = _parse_fact_holder_amounts(value)
        if holder_amounts:
            result[(match.group(1), match.group(2))] = holder_amounts
    return result


def _latest_shareholders_after(rows: list[dict[str, Any]]) -> list[str]:
    names = []
    for row in rows:
        name = _normalize_text(row.get("after_name"))
        if name and name != "合计":
            names.append(name)
    return names


def _find_latest_event(events: list[dict[str, Any]]) -> dict[str, Any] | None:
    dated = [event for event in events if event.get("date")]
    if not dated:
        return None
    return max(dated, key=lambda item: _normalize_text(item.get("date")))


def validate_events(
    events: list[dict[str, Any]], fragment_path: Path, content: dict[str, Any] | None = None
) -> dict[str, Any]:
    checks = []

    def check(check_id: str, ok: bool, actual: Any = None) -> None:
        checks.append({"id": check_id, "ok": ok, "actual": actual})

    fact_maps = _equity_fact_maps(content)
    check("fragment_docx_exists", fragment_path.exists(), str(fragment_path))
    for event in events:
        rows = event.get("equity_rows") or []
        if not rows:
            continue
        if is_founding_equity_rows(rows):
            founding_rows = founding_rows_from_equity(rows)
            has_total = any(row.get("name") == "合计" for row in founding_rows)
            total_row = next((row for row in founding_rows if row.get("name") == "合计"), {})
            check("founding_table_has_3_columns", True)
            check("founding_table_has_total_row", has_total, total_row if total_row else None)
            check("founding_after_ratio_total_100", total_row.get("ratio") == "100.0000%", total_row.get("ratio"))
            continue
        check("equity_table_has_7_columns", True)
        check("equity_table_has_two_level_header", True)
        has_total = any(row.get("before_name") == "合计" and row.get("after_name") == "合计" for row in rows)
        check("equity_table_has_total_row", has_total, rows[-1] if rows else None)
        total_row = next((row for row in rows if row.get("before_name") == "合计" and row.get("after_name") == "合计"), {})
        founding_event = any(not str(row.get("before_name", "")).strip() for row in rows if str(row.get("after_name", "")).strip())
        before_ratio_ok = total_row.get("before_ratio") == "100.0000%" or (founding_event and not str(total_row.get("before_ratio", "")).strip())
        check("equity_before_ratio_total_100", before_ratio_ok, total_row.get("before_ratio"))
        check("equity_after_ratio_total_100", total_row.get("after_ratio") == "100.0000%", total_row.get("after_ratio"))
        before_totals = _side_totals(rows, "before")
        after_totals = _side_totals(rows, "after")
        check("equity_before_amount_sum_matches_total", before_totals["amount_ok"], before_totals)
        check("equity_after_amount_sum_matches_total", after_totals["amount_ok"], after_totals)
        check("equity_before_ratio_sum_matches_total", before_totals["ratio_ok"], before_totals)
        check("equity_after_ratio_sum_matches_total", after_totals["ratio_ok"], after_totals)
        check(
            "equity_before_shareholder_amount_ratio_complete",
            before_totals["complete"],
            before_totals,
        )
        check(
            "equity_after_shareholder_amount_ratio_complete",
            after_totals["complete"],
            after_totals,
        )
        check("equity_before_shareholder_names_unique", before_totals["unique_names"], before_totals)
        check("equity_after_shareholder_names_unique", after_totals["unique_names"], after_totals)
        for side in ("before", "after"):
            expected = fact_maps.get((event.get("date"), side))
            if not expected:
                continue
            actual = _side_holder_amounts(rows, side)
            check(
                f"equity_{side}_matches_confirmed_fact_sheet",
                actual == expected,
                {"expected": expected, "actual": actual, "date": event.get("date")},
            )

    latest_event = _find_latest_event(events)
    if latest_event:
        latest_normal_rows = latest_event.get("normal_rows", [])
        enterprise_type_row = next((row for row in latest_normal_rows if _normalize_text(row.get("project")) == "企业类型变更"), None)
        if enterprise_type_row:
            after_type = _normalize_text(enterprise_type_row.get("after"))
            shareholder_count = len(_latest_shareholders_after(latest_event.get("equity_rows", [])))
            if "法人独资" in after_type:
                check("latest_corporate_sole_owner_has_one_shareholder", shareholder_count == 1, {"after_type": after_type, "shareholder_count": shareholder_count})
            if "自然人独资" in after_type:
                check("latest_natural_sole_owner_has_one_shareholder", shareholder_count == 1, {"after_type": after_type, "shareholder_count": shareholder_count})

        description = _normalize_text(latest_event.get("description"))
        has_legal_rep_row = any(_normalize_text(row.get("project")) == "法定代表人变更" for row in latest_normal_rows)
        has_personnel_row = any(_normalize_text(row.get("project")) == "主要人员变更" for row in latest_normal_rows)
        miscast_legal_rep = "法定代表人变更" in description and has_personnel_row and not has_legal_rep_row
        check("personnel_change_not_miscast_as_legal_rep_change", not miscast_legal_rep, {"description": description, "normal_rows": latest_normal_rows})
    return {"ok": all(item["ok"] for item in checks), "checks": checks, "outputs": {"history_fragment_docx": str(fragment_path)}}


def build(content: dict[str, Any], out_dir: Path) -> dict[str, Any]:
    out_dir.mkdir(parents=True, exist_ok=True)
    events = normalize_events(content)
    doc = Document()
    history = content.get("history_section", content)
    add_paragraph(doc, history.get("title", "公司设立及历史沿革"), bold=True)
    for event in events:
        add_paragraph(doc, event["description"])
        if event.get("founding_rows"):
            add_founding_table(doc, event["founding_rows"])
        equity_rows = event.get("equity_rows") or []
        if equity_rows:
            if is_founding_equity_rows(equity_rows):
                add_founding_table(doc, founding_rows_from_equity(equity_rows))
            else:
                add_equity_table(doc, equity_rows)
        normal_rows = event.get("normal_rows", [])
        if normal_rows:
            add_normal_table(doc, normal_rows)
    if history.get("ending_text"):
        add_paragraph(doc, history["ending_text"])
    fragment_path = out_dir / "history_fragment.docx"
    doc.save(fragment_path)
    events_path = out_dir / "history_events.json"
    validation_path = out_dir / "history_validation.json"
    events_path.write_text(json.dumps({"events": events, "ending_text": history.get("ending_text", "")}, ensure_ascii=False, indent=2), encoding="utf-8")
    validation = validate_events(events, fragment_path, content)
    validation_path.write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"history_events": str(events_path), "history_fragment": str(fragment_path), "history_validation": str(validation_path), "validation": validation}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--content-json", type=Path, required=True)
    parser.add_argument("--out-dir", type=Path, required=True)
    args = parser.parse_args()
    content = json.loads(args.content_json.read_text(encoding="utf-8"))
    result = build(content, args.out_dir)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["validation"]["ok"] else 2)


if __name__ == "__main__":
    main()
