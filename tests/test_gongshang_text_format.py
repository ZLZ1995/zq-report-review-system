from io import BytesIO
from pathlib import Path
import runpy

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


SCRIPT = Path(__file__).resolve().parents[1] / '.codex/skills/gongshang-change-history-docx/scripts/build_gongshang_docx.py'


def test_format_survives_save_and_preserves_unrelated_properties():
    ns = runpy.run_path(str(SCRIPT))
    doc = Document()
    p = doc.add_paragraph('正文123')
    p.runs[0].bold = True
    p.runs[0].font.name = '仿宋'
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Pt(-12)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = '表头'
    table.cell(1, 0).text = '名称'
    table.cell(1, 1).text = '100.00'
    nested = table.cell(1, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = '嵌套文字'
    header = doc.sections[0].header.paragraphs[0]
    header.add_run('原页眉').font.size = Pt(15)
    ns['apply_document_text_format'](doc)
    first_xml = doc._element.xml
    ns['apply_document_text_format'](doc)
    assert doc._element.xml == first_xml
    stream = BytesIO()
    doc.save(stream)
    stream.seek(0)
    saved = Document(stream)
    for node in saved._body._element.iter(qn('w:p')):
        in_table = any(a.tag == qn('w:tc') for a in node.iterancestors())
        ind = node.pPr.find(qn('w:ind'))
        assert ind.get(qn('w:firstLineChars')) == ('0' if in_table else '200')
        assert ind.get(qn('w:hanging')) is None
        assert ind.get(qn('w:hangingChars')) is None
        for run in node.iter(qn('w:r')):
            assert run.rPr.find(qn('w:sz')).get(qn('w:val')) == ('18' if in_table else '24')
            assert run.rPr.find(qn('w:szCs')).get(qn('w:val')) == ('18' if in_table else '24')
    assert saved.paragraphs[0].text == '正文123'
    assert saved.paragraphs[0].runs[0].bold is True
    assert saved.paragraphs[0].runs[0].font.name == '仿宋'
    assert saved.paragraphs[0].paragraph_format.space_after.pt == 6
    assert saved.sections[0].header.paragraphs[0].runs[0].font.size.pt == 15
    assert saved.tables[0].cell(0, 0)._tc is saved.tables[0].cell(0, 1)._tc


def test_generation_helpers_use_confirmed_sizes_and_character_indent():
    ns = runpy.run_path(str(SCRIPT))
    doc = Document()
    p = doc.add_paragraph('事件')
    ns['style_paragraph'](p)
    assert p.runs[0].font.size.pt == 12
    assert p._p.pPr.find(qn('w:ind')).get(qn('w:firstLineChars')) == '200'
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    ns['set_cell_text'](cell, '项目')
    assert cell.paragraphs[0].runs[0].font.size.pt == 9
    assert cell.paragraphs[0]._p.pPr.find(qn('w:ind')).get(qn('w:firstLineChars')) == '0'
