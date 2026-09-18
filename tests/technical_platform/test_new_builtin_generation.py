import json
from pathlib import Path

from openpyxl import Workbook


def financial_source(path: Path, company: str, date_text: str):
    book = Workbook()
    balance = book.active
    balance.title = '资产负债表'
    profit = book.create_sheet('利润表')
    for sheet in (balance, profit):
        sheet['A1'] = company
        sheet['A2'] = sheet.title
        sheet['A3'] = date_text
        sheet['A4'] = '项目'
    balance['B4'] = '期末余额'
    for row, (label, value) in enumerate((('资产总计', 300000), ('负债合计', 100000),
                                          ('所有者权益（或股东权益）合计', 200000)), 5):
        balance.cell(row, 1, label)
        balance.cell(row, 2, value)
    profit['B4'] = '本年金额'
    for row, (label, value) in enumerate((('一、营业收入', 500000), ('四、利润总额', 50000),
                                          ('五、净利润', 40000)), 5):
        profit.cell(row, 1, label)
        profit.cell(row, 2, value)
    book.save(path)


def test_financial_metadata_requires_and_reads_company_and_full_date(tmp_path):
    from asset_based_agent.technical_platform.generation_worker import (
        _financial_metadata,
    )

    path = tmp_path / '2024.xlsx'
    financial_source(path, '合成测试有限公司', '2024年12月31日')
    company, report_date = _financial_metadata(path)
    assert company == '合成测试有限公司'
    assert report_date.isoformat() == '2024-12-31'


def test_confirmed_office_workflow_contract_runs_real_validator(tmp_path):
    from asset_based_agent.technical_platform.generation import bundle_directory
    from asset_based_agent.technical_platform.generation_worker import workflow_contract
    from asset_based_agent.technical_platform.skills import WORKFLOW_TO_SKILL

    contract = tmp_path / 'workflow.json'
    contract.write_text(json.dumps({
        'schema_version': '1', 'proposed_skill_name': 'example-office-flow',
        'purpose': '生成一份合成办公文件', 'triggers': ['用户要求生成'],
        'inputs': ['合成输入'], 'steps': ['读取', '生成', '校验'],
        'exceptions': ['缺失输入时停止'], 'outputs': ['Word副本'],
        'acceptance_criteria': ['输出存在且来源未变化'],
        'office_routes': [{'skill': 'documents', 'responsibility': '生成并校验Word'}],
        'open_questions': [],
        'confirmation': {'workflow_confirmed': True, 'ready_to_build': True},
        'proof_run': {'sample_inputs': [], 'synthetic_sample_allowed': True,
                      'expected_output': '一份合成Word副本'},
    }, ensure_ascii=False), encoding='utf-8')
    output = tmp_path / 'output'
    output.mkdir()
    assert workflow_contract(bundle_directory(WORKFLOW_TO_SKILL.id) / 'scripts',
                             {'workflow_contract': str(contract)}, output)
    report = json.loads((output / 'office_workflow_contract_validation.json').read_text('utf-8'))
    assert report['ok'] is True
    assert report['office_skills'] == ['documents']
