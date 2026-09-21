"""Dialog-free local generation under a standing permission grant (issue: 二次确认弹窗)."""
import json
import os

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

from PySide6.QtWidgets import QApplication, QDialog

from asset_based_agent.technical_platform import generation_dialog
from asset_based_agent.technical_platform.app import PlatformWindow, TaskWorker
from asset_based_agent.technical_platform.skills import (
    DETAIL,
    FINANCIAL_BRIEF,
)
from asset_based_agent.technical_platform.store import PlatformStore
from asset_based_agent.technical_platform.task_events import TaskDestination


def make_report(path, date_text):
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = '资产负债表'
    sheet['A1'] = '资产负债表'
    sheet['A2'] = f'编制单位：测试公司 {date_text} 单位：元'
    book.create_sheet('利润表')
    book.save(path)


def make_plain_xlsx(path):
    from openpyxl import Workbook
    book = Workbook()
    book.active.title = '流水'
    book.save(path)


def forbid_dialog(monkeypatch):
    def boom(self, *args, **kwargs):
        raise AssertionError('完全访问/帮我批准模式下不得弹出二次确认对话框')
    monkeypatch.setattr(generation_dialog.GenerationDialog, 'exec', boom)


def no_worker_start(monkeypatch):
    monkeypatch.setattr(TaskWorker, 'start', lambda self: None)


def build_window(tmp_path, monkeypatch, *, mode='full', reports=3, extras=0, client=None, models=None):
    assert QApplication.instance() or QApplication([])
    store = PlatformStore(tmp_path / 'state.db', 'test')
    project = store.create_project('直接生成')
    store.create_session(project)
    window = PlatformWindow(store, client=client, models=models)
    window.reload_projects(project)
    window.set_agent_permission_mode(mode)
    paths = []
    dates = ('2024年12月31日', '2025年12月31日', '2026年6月30日')
    for index in range(reports):
        path = tmp_path / f'report{index}.xlsx'
        make_report(path, dates[index])
        paths.append(path)
    for index in range(extras):
        path = tmp_path / f'extra{index}.xlsx'
        make_plain_xlsx(path)
        paths.append(path)
    window.import_files(paths)
    return window, store


def snapshot_of(store, run_id):
    return json.loads(store.run(run_id)['snapshot'])


def test_run_feedback_success_without_worker_note_does_not_misreport(tmp_path):
    from asset_based_agent.technical_platform.generation import run_feedback
    (tmp_path / 'output').mkdir()
    assert '生成完成' in run_feedback(tmp_path, True)
    assert '生成未完成' in run_feedback(tmp_path, False)
    (tmp_path / 'output' / 'user_feedback.md').write_text('工作簿已生成。', encoding='utf-8')
    assert run_feedback(tmp_path, True) == '工作簿已生成。'


def test_full_permission_generates_brief_without_dialog_and_infers_roles(tmp_path, monkeypatch):
    forbid_dialog(monkeypatch)
    no_worker_start(monkeypatch)
    window, store = build_window(tmp_path, monkeypatch, mode='full')
    try:
        window.execute_plan('根据三期报表生成财务状况简表', window.registry.get(FINANCIAL_BRIEF.id))
        assert window.run_id is not None
        snapshot = snapshot_of(store, window.run_id)
        assert snapshot['skill_id'] == FINANCIAL_BRIEF.id
        assert snapshot['mode'] == 'local_generation'
        assert len(snapshot['files']) == 3
        roles = snapshot['input_roles']
        names = {item['id']: item['name'] for item in snapshot['files']}
        assert names[roles['period_one']] == 'report0.xlsx'
        assert names[roles['period_two']] == 'report1.xlsx'
        assert names[roles['basis_date']] == 'report2.xlsx'
        messages = [m['text'] for m in store.messages(window.session_id) if m['role'] == 'event']
        assert any('直接执行' in text for text in messages)
    finally:
        window.close()


def test_risk_permission_also_skips_dialog(tmp_path, monkeypatch):
    forbid_dialog(monkeypatch)
    no_worker_start(monkeypatch)
    window, store = build_window(tmp_path, monkeypatch, mode='risk')
    try:
        window.execute_plan('生成财务简报', window.registry.get(FINANCIAL_BRIEF.id))
        assert window.run_id is not None
        assert snapshot_of(store, window.run_id)['skill_id'] == FINANCIAL_BRIEF.id
    finally:
        window.close()


def test_request_permission_keeps_explicit_dialog(tmp_path, monkeypatch):
    no_worker_start(monkeypatch)
    calls = []

    def fake_exec(self):
        calls.append(self.skill.id)
        return QDialog.DialogCode.Rejected

    monkeypatch.setattr(generation_dialog.GenerationDialog, 'exec', fake_exec)
    window, _store = build_window(tmp_path, monkeypatch, mode='request')
    try:
        window.execute_plan('生成财务简报', window.registry.get(FINANCIAL_BRIEF.id))
        assert calls == [FINANCIAL_BRIEF.id]
        assert window.run_id is None
        assert window.composer.toPlainText() == '生成财务简报'
    finally:
        window.close()


def test_brief_narrows_extra_files_to_three_reports(tmp_path, monkeypatch):
    forbid_dialog(monkeypatch)
    no_worker_start(monkeypatch)
    window, store = build_window(tmp_path, monkeypatch, mode='full', reports=3, extras=2)
    try:
        window.execute_plan('生成财务简报', window.registry.get(FINANCIAL_BRIEF.id))
        assert window.run_id is not None
        snapshot = snapshot_of(store, window.run_id)
        assert sorted(item['name'] for item in snapshot['files']) == [
            'report0.xlsx', 'report1.xlsx', 'report2.xlsx']
    finally:
        window.close()


def test_brief_without_three_reports_explains_in_conversation(tmp_path, monkeypatch):
    forbid_dialog(monkeypatch)
    no_worker_start(monkeypatch)
    # 两份报表 + 一份非报表，无法凑齐两年一期
    window, store = build_window(tmp_path, monkeypatch, mode='full', reports=2, extras=1)
    try:
        window.execute_plan('生成财务简报', window.registry.get(FINANCIAL_BRIEF.id))
        assert window.run_id is None
        last = store.messages(window.session_id)[-1]
        assert last['role'] == 'assistant'
        assert '三个期间' in last['text'] or '三份' in last['text']
        assert window.composer.toPlainText() == '生成财务简报'
    finally:
        window.close()


def test_compound_request_chains_detail_then_brief(tmp_path, monkeypatch):
    forbid_dialog(monkeypatch)
    no_worker_start(monkeypatch)
    window, store = build_window(tmp_path, monkeypatch, mode='full', reports=3, extras=2,
                                 client=object(), models=[{'model_id': 'm', 'display_name': '模型'}])
    try:
        window.composer.setPlainText('根据上传的文件生成评估明细表和两年一期财务简报')
        window.submit()
        first_run = window.run_id
        assert first_run is not None
        assert snapshot_of(store, first_run)['skill_id'] == DETAIL.id
        assert window._local_chain is not None
        assert [spec.id for spec in window._local_chain['specs']] == [FINANCIAL_BRIEF.id]
        events = [m['text'] for m in store.messages(window.session_id) if m['role'] == 'event']
        assert any('依次执行' in text for text in events)
        # 第一步成功完成后自动接续第二步，无需用户再次确认
        store.transition(first_run, 'running', 'test')
        store.transition(first_run, 'validating', 'test')
        store.transition(first_run, 'succeeded', 'test')
        worker = window.worker
        destination = TaskDestination.resolve(store, first_run)
        window.finished(worker, destination)
        assert window.run_id != first_run
        snapshot = snapshot_of(store, window.run_id)
        assert snapshot['skill_id'] == FINANCIAL_BRIEF.id
        assert sorted(item['name'] for item in snapshot['files']) == [
            'report0.xlsx', 'report1.xlsx', 'report2.xlsx']
        roles = snapshot['input_roles']
        names = {item['id']: item['name'] for item in snapshot['files']}
        assert names[roles['basis_date']] == 'report2.xlsx'
    finally:
        window.close()


def test_chain_continues_when_first_step_failed(tmp_path, monkeypatch):
    """链式步骤互相独立：首步失败不连坐后续交付，仅取消才停止。"""
    forbid_dialog(monkeypatch)
    no_worker_start(monkeypatch)
    window, store = build_window(tmp_path, monkeypatch, mode='full', reports=3, extras=2,
                                 client=object(), models=[{'model_id': 'm', 'display_name': '模型'}])
    try:
        window.composer.setPlainText('生成评估明细表和财务简报')
        window.submit()
        first_run = window.run_id
        assert first_run is not None
        store.transition(first_run, 'running', 'test')
        store.transition(first_run, 'failed', '资料识别不完整')
        worker = window.worker
        destination = TaskDestination.resolve(store, first_run)
        window.finished(worker, destination)
        assert window.run_id != first_run, '首步失败后应继续执行简报步骤'
        snapshot = snapshot_of(store, window.run_id)
        assert snapshot['skill_id'] == FINANCIAL_BRIEF.id
        notes = [m['text'] for m in store.messages(window.session_id) if m['role'] == 'assistant']
        assert any('继续执行后续技能' in text for text in notes)
    finally:
        window.close()


def test_chain_stops_when_first_step_cancelled(tmp_path, monkeypatch):
    forbid_dialog(monkeypatch)
    no_worker_start(monkeypatch)
    window, store = build_window(tmp_path, monkeypatch, mode='full', reports=3, extras=2,
                                 client=object(), models=[{'model_id': 'm', 'display_name': '模型'}])
    try:
        window.composer.setPlainText('生成评估明细表和财务简报')
        window.submit()
        first_run = window.run_id
        assert first_run is not None
        store.transition(first_run, 'running', 'test')
        store.transition(first_run, 'cancelled', 'test')
        worker = window.worker
        destination = TaskDestination.resolve(store, first_run)
        window.finished(worker, destination)
        assert window.run_id == first_run
        assert window._local_chain is None
        assert window.worker is None
    finally:
        window.close()
