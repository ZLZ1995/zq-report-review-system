"""Read-only comparison of all template formulas and explicit OOXML hyperlinks."""
import argparse
import json
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook


def snapshot(path):
    wb = load_workbook(path, data_only=False)
    result = {'sheets': wb.sheetnames, 'formulas': {}}
    for ws in wb:
        for row in ws:
            for cell in row:
                if cell.data_type == 'f':
                    result['formulas'][f'{ws.title}!{cell.coordinate}'] = str(cell.value)
    wb.close()
    result['hyperlinks'] = {}
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                root = etree.fromstring(archive.read(name))
                links = root.xpath('//*[local-name()="hyperlink"]')
                rel_name = 'xl/worksheets/_rels/' + name.rsplit('/', 1)[-1] + '.rels'
                relations = {}
                if rel_name in archive.namelist():
                    relations = {rel.get('Id'): dict(rel.attrib)
                                 for rel in etree.fromstring(archive.read(rel_name))}
                normalized = []
                for link in links:
                    attributes = dict(link.attrib)
                    rid = attributes.pop('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id', None)
                    if rid:
                        relation = dict(relations[rid])
                        relation.pop('Id')
                        attributes['relationship'] = relation
                    normalized.append(attributes)
                result['hyperlinks'][name] = sorted(normalized, key=lambda item: item['ref'])
    return result


def compare(template, generated):
    before, after = snapshot(template), snapshot(generated)
    changed = [key for key, value in before['formulas'].items()
               if after['formulas'].get(key) != value]
    links = [key for key in before['hyperlinks']
             if before['hyperlinks'][key] != after['hyperlinks'].get(key, [])]
    return {'ok': not changed and not links and before['sheets'] == after['sheets'],
            'template_formula_count': len(before['formulas']),
            'changed_formula_locations': changed, 'changed_link_parts': links,
            'sheet_order_preserved': before['sheets'] == after['sheets']}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('template', type=Path)
    parser.add_argument('generated', type=Path)
    args = parser.parse_args()
    report = compare(args.template, args.generated)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    raise SystemExit(0 if report['ok'] else 2)
