from datetime import datetime
import importlib

import pytest
from openpyxl import Workbook, load_workbook

from test_detail_workbook_pipeline_guards import load_pipeline_module


def statement(path, when, company='上海禾念信息科技有限公司'):
    wb = Workbook()
    wb.active.title = '其他页'
    ws = wb.create_sheet('资产负债表')
    ws['B2'] = '资产负债表'
    ws['A1'] = 11997
    ws['B3'] = when
    ws['B5'] = '编制单位：' + company
    ws['B6'] = '资产'
    wb.save(path)
    return path


def test_latest_statement_uses_content_not_filename(tmp_path):
    load_pipeline_module()
    m = importlib.import_module('cover_metadata')
    old = statement(tmp_path / '2099.xlsx', '2024年11月30日')
    new = statement(tmp_path / 'old.xlsx', '2026年3月31日')
    result = m.select_latest_statement([old, new])
    assert result['source'] == str(new)
    assert result['company'] == '上海禾念信息科技有限公司'
    assert result['report_date'] == datetime(2026, 3, 31)
    assert result['company_cell'] == 'B5'
    assert result['date_cell'] == 'B3'


@pytest.mark.parametrize('name', ['封面', '封面页'])
def test_cover_writes_only_confirmed_cells(name, tmp_path):
    m = load_pipeline_module()
    wb = Workbook()
    ws = wb.active
    ws.title = name
    ws.merge_cells('F7:M7')
    ws['D7'] = '被评估单位：'
    ws['D9'] = '评估基准日：'
    ws['G9'], ws['I9'], ws['K9'] = '年', '月', '日'
    ws['F13'], ws['H13'], ws['J13'] = 2025, 1, 9
    linked = wb.create_sheet('分类汇总表')
    linked['A3'] = f"='{name}'!F9"
    payload = {'company': '上海禾念信息科技有限公司', 'report_date': datetime(2026, 3, 31)}
    m.fill_cover_and_balance_sheet(wb, payload['company'], payload, {}, {})
    path = tmp_path / 'filled.xlsx'
    wb.save(path)
    out = load_workbook(path)
    assert [out[name][c].value for c in ['F7','F9','H9','J9']] == [payload['company'],2026,3,31]
    assert [out[name][c].value for c in ['F13','H13','J13']] == [2025,1,9]
    assert [out[name][c].value for c in ['G9','I9','K9']] == ['年','月','日']
    assert str(out[name].merged_cells) == 'F7:M7'
    assert out['分类汇总表']['A3'].data_type == 'f'


def test_missing_date_and_mixed_companies_fail(tmp_path):
    load_pipeline_module()
    m = importlib.import_module('cover_metadata')
    missing = statement(tmp_path / 'missing.xlsx', None)
    with pytest.raises(ValueError, match='metadata_missing'):
        m.select_latest_statement([missing])
    a = statement(tmp_path / 'a.xlsx', '2026/03/31', '甲公司')
    b = statement(tmp_path / 'b.xlsx', '2026/04/30', '乙公司')
    with pytest.raises(ValueError, match='companies_conflict'):
        m.select_latest_statement([a, b])


def test_date_accepts_excel_date_not_unformatted_numbers():
    load_pipeline_module()
    m = importlib.import_module('cover_metadata')
    assert m.statement_date(datetime(2026, 3, 31)) == datetime(2026, 3, 31)
    assert m.statement_date(11997) is None
    assert m.statement_date('11997') is None


def test_parse_balance_sheet_keeps_real_header_metadata(tmp_path):
    m = load_pipeline_module()
    p = statement(tmp_path / 'source.xlsx', '2026年3月31日')
    result = m.parse_balance_sheet(p)
    assert result['report_date'] == datetime(2026, 3, 31)
    assert result['company'] == '上海禾念信息科技有限公司'
