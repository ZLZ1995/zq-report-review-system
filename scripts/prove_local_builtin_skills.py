"""Run synthetic local-only proof cases for the newly bundled skills."""
from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook

from asset_based_agent.technical_platform.generation import bundle_directory
from asset_based_agent.technical_platform.generation_worker import (
    financial_brief,
    workflow_contract,
)
from asset_based_agent.technical_platform.skills import (
    FINANCIAL_BRIEF,
    WORKFLOW_TO_SKILL,
)

ROOT = Path(__file__).resolve().parents[1]


def financial_source(path: Path, company: str, date_text: str, scale: int) -> None:
    book = Workbook()
    balance = book.active
    balance.title = '资产负债表'
    profit = book.create_sheet('利润表')
    for sheet in (balance, profit):
        sheet['A1'] = company
        sheet['A2'] = sheet.title
        sheet['A3'] = date_text
        sheet['A4'] = '项目'
    balance['B4'] = '期末余额'
    for row, (label, value) in enumerate((('资产总计', 300000 * scale),
                                          ('负债合计', 100000 * scale),
                                          ('所有者权益（或股东权益）合计', 200000 * scale)), 5):
        balance.cell(row, 1, label)
        balance.cell(row, 2, value)
    profit['B4'] = '本年金额'
    for row, (label, value) in enumerate((('一、营业收入', 500000 * scale),
                                          ('四、利润总额', 50000 * scale),
                                          ('五、净利润', 40000 * scale)), 5):
        profit.cell(row, 1, label)
        profit.cell(row, 2, value)
    book.save(path)


def main() -> int:
    root = ROOT / 'build' / (
        'local-skill-proof-' + datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S')
    )
    root.mkdir(parents=True)
    inputs = root / 'inputs'
    inputs.mkdir()
    periods = [('period_one', '2023年12月31日'), ('period_two', '2024年12月31日'),
               ('basis_date', '2025年6月30日')]
    selected = {}
    for index, (role, date_text) in enumerate(periods, 1):
        path = inputs / f'{index}.xlsx'
        financial_source(path, '合成测试有限公司', date_text, index)
        selected[role] = str(path)
    selected['template'] = str(bundle_directory(FINANCIAL_BRIEF.id) /
                               'assets' / 'financial_table.docx')
    financial_output = root / 'financial'
    financial_output.mkdir()
    financial_ok = financial_brief(bundle_directory(FINANCIAL_BRIEF.id) / 'scripts',
                                   selected, financial_output)

    contract = inputs / 'workflow.json'
    contract.write_text(json.dumps({
        'schema_version': '1', 'proposed_skill_name': 'example-office-flow',
        'purpose': '生成一份合成办公文件', 'triggers': ['用户要求生成'],
        'inputs': ['合成输入'], 'steps': ['读取', '生成', '校验'],
        'exceptions': ['缺失输入时停止'], 'outputs': ['Word副本'],
        'acceptance_criteria': ['输出存在且来源未变化'],
        'office_routes': [{'skill': 'documents', 'responsibility': '生成并校验Word'}],
        'open_questions': [],
        'confirmation': {'workflow_confirmed': True, 'ready_to_build': True},
        'proof_run': {'sample_inputs': [], 'synthetic_sample_allowed': True,
                      'expected_output': '一份合成Word副本'},
    }, ensure_ascii=False), encoding='utf-8')
    workflow_output = root / 'workflow'
    workflow_output.mkdir()
    workflow_ok = workflow_contract(bundle_directory(WORKFLOW_TO_SKILL.id) / 'scripts',
                                    {'workflow_contract': str(contract)}, workflow_output)
    result = {'root': str(root), 'financial_ok': financial_ok, 'workflow_ok': workflow_ok,
              'financial_files': sorted(path.name for path in financial_output.iterdir()),
              'workflow_files': sorted(path.name for path in workflow_output.iterdir())}
    (root / 'proof.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False))
    return 0 if financial_ok and workflow_ok else 2


if __name__ == '__main__':
    raise SystemExit(main())
