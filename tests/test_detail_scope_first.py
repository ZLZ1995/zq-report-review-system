from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from openpyxl import Workbook
import pytest


def load_pipeline_module():
    path = Path(".codex/skills/valuation-detail-workbook-fill/scripts/run_detail_workbook_pipeline.py")
    sys.path.insert(0, str(path.parent.resolve()))
    spec = importlib.util.spec_from_file_location("detail_pipeline_scope_test", path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def bank_only_balance_sheet() -> dict[str, float]:
    return {
        "货币资金": 47499.02,
        "资产总计": 47499.02,
        "负债合计": 0.0,
        "实收资本": 260000.0,
        "未分配利润": -212500.98,
        "所有者权益合计": 47499.02,
        "负债和所有者权益合计": 47499.02,
    }


def test_bank_only_with_evidence_selects_lightweight_scope() -> None:
    module = load_pipeline_module()
    sheet_plan = [("银行存款", [{"debit_end": 47499.02, "credit_end": 0.0}], "detail_fillable")]

    result = module.select_execution_scope(
        bank_only_balance_sheet(), sheet_plan, requested_mode="auto", bank_evidence_available=True
    )

    assert result["selected_mode"] == "single_asset_lightweight"
    assert result["active_detail_sheets"] == ["银行存款"]
    assert result["stage1_excel_recalc_required"] is False
    assert "银行存款" in result["required_dependency_sheets"]


def test_total_label_alias_does_not_activate_asset_family():
    module = load_pipeline_module()
    values = bank_only_balance_sheet()
    values['负债和所有者权益总计'] = values.pop('负债和所有者权益合计')
    result = module.select_execution_scope(values,
        [('银行存款', [{'book_value': 47499.02}], 'detail_fillable')], bank_evidence_available=True)
    assert result['selected_mode'] == 'single_asset_lightweight'


def test_second_nonzero_asset_selects_scoped_standard() -> None:
    module = load_pipeline_module()
    values = bank_only_balance_sheet() | {"应收账款": 100.0, "资产总计": 47599.02, "负债和所有者权益合计": 47599.02}
    sheet_plan = [
        ("银行存款", [{"book_value": 47499.02}], "detail_fillable"),
        ("应收账款", [{"book_value": 100.0}], "detail_fillable"),
    ]

    result = module.select_execution_scope(values, sheet_plan, bank_evidence_available=True)

    assert result["selected_mode"] == "scoped_standard"
    assert result["active_detail_sheets"] == ["应收账款", "银行存款"]
    assert result["stage1_excel_recalc_required"] is True


def test_explicit_full_mode_is_preserved_as_fallback() -> None:
    module = load_pipeline_module()

    result = module.select_execution_scope(
        bank_only_balance_sheet(), [], requested_mode="full", bank_evidence_available=True
    )

    assert result["selected_mode"] == "full_template"


def test_lightweight_scope_skips_out_of_scope_fixed_asset_cleanup() -> None:
    module = load_pipeline_module()
    wb = Workbook()
    wb.active.title = "银行存款"
    fixed = wb.create_sheet("机器设备")
    fixed["A6"] = "旧项目设备"
    fixed["B6"] = "=1/0"

    _, meta = module.stage3_enrich_detail_pages_from_journals(
        wb,
        {"values": bank_only_balance_sheet()},
        [],
        {
            "selected_mode": "single_asset_lightweight",
            "active_detail_sheets": ["银行存款"],
            "required_dependency_sheets": ["银行存款", "流动汇总", "分类汇总"],
        },
    )

    assert fixed["A6"].value == "旧项目设备"
    assert fixed["B6"].value == "=1/0"
    assert meta["fixed_asset_zero_balance_cleanup"]["skipped"] == "out_of_scope"


def test_scoped_stage1_validation_checks_direct_balance_sheet_inputs(tmp_path: Path) -> None:
    module = load_pipeline_module()
    workbook_path = tmp_path / "stage1.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "资产负债表"
    ws["D7"] = 47499.02
    ws["H31"] = 260000.0
    ws["I31"] = "=H31"
    ws["I35"] = -212500.98
    wb.save(workbook_path)

    result = module.validate_saved_stage1_scoped(
        workbook_path, {"values_current": bank_only_balance_sheet()}
    )

    assert result["assets_equal_liabilities_equity"] is True
    assert result["validation_mode"] == "scoped_direct_inputs"


def test_scoped_plan_excludes_inactive_pages():
    module = load_pipeline_module()
    plan = [('银行存款', [{'book_value': 100}], 'detail_fillable'),
            ('应收账款', [], 'zero_balance_cleanup')]
    scope = {'selected_mode': 'scoped_standard', 'active_detail_sheets': ['银行存款']}
    assert module.restrict_plan_to_scope(plan, scope) == plan[:1]


def test_bank_statement_difference_cannot_be_replaced_with_balance_sheet_amount(tmp_path):
    module = load_pipeline_module()
    wb = Workbook()
    wb.active.append(['账号', '账户名称', '开户行', '交易日期', '账户余额'])
    wb.active.append(['1234567890', '测试有限公司', '测试银行', '2026/03/31', 100])
    source = tmp_path / 'bank.xlsx'
    wb.save(source)
    with pytest.raises(ValueError, match='bank_statement_balance_mismatch'):
        module.load_bank_statement_evidence([str(source)], {'货币资金': 500})
def test_bank_evidence_activates_bank_page_without_trial_balance():
    import importlib.util
    import sys
    from pathlib import Path
    scripts = Path(__file__).resolve().parents[1] / '.codex/skills/valuation-detail-workbook-fill/scripts'
    sys.path.insert(0, str(scripts))
    spec = importlib.util.spec_from_file_location('no_tb_scope_probe', scripts / 'run_detail_workbook_pipeline.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    result = module.select_execution_scope({'货币资金': 1000, '实收资本': 1000},
                                           [('银行存款', [], 'cash')], bank_evidence_available=True)
    assert result['active_detail_sheets'] == ['银行存款']


def test_parse_date_text_accepts_datetime_with_time() -> None:
    module = load_pipeline_module()
    from datetime import datetime
    assert module.parse_date_text('2026-06-21 02:09:18') == datetime(2026, 6, 21)
    assert module.parse_date_text('2026/6/5 07:33:15') == datetime(2026, 6, 5)
