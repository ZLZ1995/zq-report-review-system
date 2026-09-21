"""Structure checks for formal Chinese valuation report DOCX files."""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from docx import Document


REQUIRED_TOC_HEADINGS = [
    "声明",
    "资产评估报告摘要",
    "资产评估报告",
    "一、委托人、被评估单位及其他资产评估报告使用人",
    "二、评估目的",
    "三、评估对象和评估范围",
    "四、价值类型",
    "五、评估基准日",
    "六、评估依据",
    "七、评估方法",
    "八、评估程序实施过程和情况",
    "九、评估假设",
    "十、评估结论",
    "十一、特别事项说明",
    "十二、资产评估报告使用限制说明",
    "十三、资产评估报告日",
    "十四、资产评估专业人员签名和资产评估机构印章",
    "资产评估报告附件",
]

REQUIRED_BODY_HEADINGS = REQUIRED_TOC_HEADINGS[3:-1]

_BODY_HEADING_RE = re.compile(r"^([一二三四五六七八九十]+)、(.+)$")
_TOC_PAGE_RE = re.compile(r"\s*\d+\s*$")


@dataclass
class ReportStructureSnapshot:
    path: str
    paragraph_count: int
    nonempty_paragraph_count: int
    table_count: int
    section_count: int
    toc_headings: list[str]
    body_major_headings: list[str]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def normalize_heading(text: str) -> str:
    text = text.replace("\u3000", " ").strip()
    text = re.sub(r"\s+", "", text)
    text = text.replace("，", "、", 1) if text.startswith("一，") else text
    return text


def extract_structure_snapshot(path: str | Path) -> ReportStructureSnapshot:
    docx_path = Path(path)
    doc = Document(str(docx_path))
    paragraphs = [_paragraph_text(paragraph).strip() for paragraph in doc.paragraphs]
    nonempty = [text for text in paragraphs if text]
    toc_headings = _extract_toc_headings(nonempty)
    body_headings = _extract_body_major_headings(nonempty)
    return ReportStructureSnapshot(
        path=str(docx_path),
        paragraph_count=len(paragraphs),
        nonempty_paragraph_count=len(nonempty),
        table_count=len(doc.tables),
        section_count=len(doc.sections),
        toc_headings=toc_headings,
        body_major_headings=body_headings,
    )


def compare_structure(
    before: ReportStructureSnapshot | dict[str, Any],
    after: ReportStructureSnapshot | dict[str, Any],
    *,
    allowed_deleted_paragraphs: int = 0,
    allowed_deleted_tables: int = 0,
) -> dict[str, Any]:
    before_data = before.to_dict() if isinstance(before, ReportStructureSnapshot) else before
    after_data = after.to_dict() if isinstance(after, ReportStructureSnapshot) else after
    before_body = [normalize_heading(item) for item in before_data.get("body_major_headings", [])]
    after_body = [normalize_heading(item) for item in after_data.get("body_major_headings", [])]
    required_body = [normalize_heading(item) for item in REQUIRED_BODY_HEADINGS]
    missing_required = [item for item in REQUIRED_BODY_HEADINGS if normalize_heading(item) not in after_body]
    baseline_missing_after = [item for item in before_data.get("body_major_headings", []) if normalize_heading(item) not in after_body]
    order_ok = _is_subsequence(required_body, after_body)
    deleted_nonempty = int(before_data.get("nonempty_paragraph_count", 0)) - int(
        after_data.get("nonempty_paragraph_count", 0)
    )
    deleted_tables = int(before_data.get("table_count", 0)) - int(after_data.get("table_count", 0))
    checks = {
        "toc_major_headings_complete": _contains_all(after_data.get("toc_headings", []), REQUIRED_TOC_HEADINGS),
        "body_major_headings_complete": not missing_required,
        "major_heading_order_unchanged": order_ok,
        "major_heading_names_unchanged": not baseline_missing_after,
        "missing_major_headings": missing_required,
        "baseline_headings_missing_after": baseline_missing_after,
        "unauthorized_deleted_paragraphs": max(0, deleted_nonempty - allowed_deleted_paragraphs),
        "unauthorized_deleted_tables": max(0, deleted_tables - allowed_deleted_tables),
        "paragraph_count_before": before_data.get("nonempty_paragraph_count", 0),
        "paragraph_count_after": after_data.get("nonempty_paragraph_count", 0),
        "table_count_before": before_data.get("table_count", 0),
        "table_count_after": after_data.get("table_count", 0),
    }
    checks["ok"] = (
        checks["body_major_headings_complete"]
        and checks["major_heading_order_unchanged"]
        and checks["major_heading_names_unchanged"]
        and checks["unauthorized_deleted_paragraphs"] == 0
        and checks["unauthorized_deleted_tables"] == 0
    )
    return checks


def _paragraph_text(paragraph: Any) -> str:
    return "".join(run.text for run in paragraph.runs) or paragraph.text


def _extract_toc_headings(nonempty: list[str]) -> list[str]:
    headings: list[str] = []
    in_toc = False
    for text in nonempty:
        compact = normalize_heading(_TOC_PAGE_RE.sub("", text).replace(".", ""))
        if compact == "目录":
            in_toc = True
            continue
        if not in_toc:
            continue
        if compact.startswith("资产评估报告附件"):
            headings.append("资产评估报告附件")
            break
        for heading in REQUIRED_TOC_HEADINGS:
            if normalize_heading(heading) in compact:
                headings.append(heading)
                break
    return _dedupe_preserve_order(headings)


def _extract_body_major_headings(nonempty: list[str]) -> list[str]:
    headings: list[str] = []
    body_seen = False
    for text in nonempty:
        normalized = normalize_heading(text)
        if normalized == "资产评估报告":
            body_seen = True
            continue
        if not body_seen:
            continue
        match = _BODY_HEADING_RE.match(normalized)
        if not match:
            continue
        candidate = f"{match.group(1)}、{match.group(2)}"
        for required in REQUIRED_BODY_HEADINGS:
            if normalize_heading(required) == candidate:
                headings.append(required)
                break
    return _dedupe_preserve_order(headings)


def _contains_all(actual: list[str], required: list[str]) -> bool:
    actual_norm = {normalize_heading(item) for item in actual}
    return all(normalize_heading(item) in actual_norm for item in required)


def _is_subsequence(required: list[str], actual: list[str]) -> bool:
    if not required:
        return True
    pos = 0
    for item in actual:
        if item == required[pos]:
            pos += 1
            if pos == len(required):
                return True
    return False


def _dedupe_preserve_order(items: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for item in items:
        key = normalize_heading(item)
        if key in seen:
            continue
        seen.add(key)
        result.append(item)
    return result
