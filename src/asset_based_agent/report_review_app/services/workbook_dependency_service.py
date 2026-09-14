"""Cell-level formula dependency tracing for workbook review."""

from __future__ import annotations

from collections.abc import Iterable

from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import get_column_letter, range_boundaries

CellKey = tuple[str, str]


def summary_impact_cells(
    workbook: object,
    summary_sheets: set[str],
) -> set[CellKey]:
    """Return cells reachable backwards from final summary cells."""
    dependencies: dict[CellKey, set[CellKey]] = {}
    roots: set[CellKey] = set()
    sheet_names = {worksheet.title for worksheet in workbook.worksheets}
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value in (None, ""):
                    continue
                key = (worksheet.title, cell.coordinate)
                if worksheet.title in summary_sheets:
                    roots.add(key)
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    dependencies[key] = formula_references(
                        cell.value,
                        worksheet.title,
                        sheet_names,
                    )

    relevant = set(roots)
    pending = list(roots)
    while pending:
        key = pending.pop()
        for dependency in dependencies.get(key, set()):
            if dependency in relevant:
                continue
            relevant.add(dependency)
            pending.append(dependency)
    return relevant


def formula_references(
    formula: str,
    current_sheet: str,
    sheet_names: set[str],
) -> set[CellKey]:
    references: set[CellKey] = set()
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return references
    for token in tokens:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        references.update(
            _range_references(token.value, current_sheet, sheet_names)
        )
    return references


def _range_references(
    value: str,
    current_sheet: str,
    sheet_names: set[str],
) -> Iterable[CellKey]:
    if "[" in value or "]" in value:
        return ()
    if "!" in value:
        raw_sheet, address = value.rsplit("!", 1)
        sheet = raw_sheet.strip("'").replace("''", "'")
    else:
        sheet = current_sheet
        address = value
    if sheet not in sheet_names:
        return ()
    address = address.replace("$", "")
    try:
        min_col, min_row, max_col, max_row = range_boundaries(address)
    except ValueError:
        return ()
    return (
        (sheet, f"{get_column_letter(column)}{row}")
        for row in range(min_row, max_row + 1)
        for column in range(min_col, max_col + 1)
    )
