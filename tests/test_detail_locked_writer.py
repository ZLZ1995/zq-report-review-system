import pytest
from openpyxl import Workbook
from test_detail_workbook_pipeline_guards import load_pipeline_module


def fixture():
    wb = Workbook()
    ws = wb.active
    ws.title = '应付账款'
    ws._locked_template_layout = True
    ws['A8'] = '合计'
    ws['G8'] = '=SUM(G6:G7)'
    ws['H6'], ws['H7'], ws['H8'] = '=G6', '=G7', '=SUM(H6:H7)'
    ws['A9'] = '企业负责人：'
    registry = {'selected_sheets': {ws.title: {'confirmed_input_cells': [
        f'{col}{row}' for row in (6, 7) for col in 'ABCDGI'],
        'detail_body_cells': [f'{col}{row}' for row in (6, 7) for col in 'ABCDGI']}}}
    return ws, registry


def test_locked_writer_keeps_formulas_footer_and_capacity():
    m = load_pipeline_module()
    ws, registry = fixture()
    before = m.capture_template_formulas(ws.parent)
    m.write_simple_detail_sheet(ws, [{'counterparty': '甲有限公司', 'book_value': 100,
        'fill_mode': 'placeholder_only'}], 'payable', {}, registry, [])
    m.assert_template_formulas_preserved(ws.parent, before)
    assert ws['G6'].value == 100
    assert ws['A9'].value == '企业负责人：'
    assert ws.max_row == 9


def test_locked_writer_rejects_overflow_before_changing_cells():
    m = load_pipeline_module()
    ws, registry = fixture()
    ws['G6'] = 987
    with pytest.raises(m.ProtectionViolation, match='locked_template_capacity_exceeded'):
        m.write_simple_detail_sheet(ws, [{'book_value': 100}] * 3, 'payable', {}, registry, [])
    assert ws['G6'].value == 987


def test_locked_footer_normalization_does_not_rewrite_fixed_template():
    m = load_pipeline_module()
    ws, _ = fixture()
    m.normalize_all_detail_sheet_footers(ws.parent)
    assert ws['A8'].value == '合计'
    assert ws['G8'].value == '=SUM(G6:G7)'
    assert ws['A9'].value == '企业负责人：'


def test_age_bucket_uses_current_project_date():
    from datetime import datetime
    m = load_pipeline_module()
    # Excel stores local calendar dates without timezone information.
    assert m.derive_age_bucket(datetime(2024, 3, 1), base_date=datetime(2024, 3, 31)) == ('1年以内', 'G')  # noqa: DTZ001
    with pytest.raises(ValueError, match='after_report_date'):
        m.derive_age_bucket(datetime(2024, 4, 1), base_date=datetime(2024, 3, 31))  # noqa: DTZ001


def test_journal_entity_does_not_concatenate_duplicate_source_fields():
    m = load_pipeline_module()
    rows = [{'vendor_name': '甲有限公司', 'counterparty_desc': '甲有限公司',
        'account_name': '应付账款', 'summary': '采购服务', 'tb_code': '2202', 'debit': 0, 'credit': 100}]
    assert m.build_journal_entity_index(rows)['2202'][0]['entity'] == '甲有限公司'


def test_conflicting_journal_entities_are_not_silently_selected():
    m = load_pipeline_module()
    rows = [{'vendor_name': '甲有限公司', 'counterparty_desc': '乙有限公司',
        'account_name': '应付账款', 'summary': '采购服务', 'tb_code': '2202', 'debit': 0, 'credit': 100}]
    with pytest.raises(ValueError, match='conflicting_journal_counterparties'):
        m.build_journal_entity_index(rows)
